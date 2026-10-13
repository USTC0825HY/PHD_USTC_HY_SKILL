#!/usr/bin/env python3
"""Audit Laser result-summary DOCX structure and formatting."""

from __future__ import annotations

import argparse
import json
import re
import zipfile
from pathlib import Path
from collections import Counter

from docx import Document
from docx.oxml.ns import qn

from result_summary_style_contract import BODY_STYLE, HEADING_STYLES, PAGE_STYLE, PLACEHOLDERS, SPECIAL_STYLES, TABLE_STYLE, TITLE_STYLE

from docx_safe_edit import table_checks, surgical_diff, reference_checks

TOL = 0.04


def _cm(length) -> float:
    return length.cm


def _attr(node, name, default=""):
    return node.get(qn(name), default) if node is not None else default


def audit(path: Path, mode: str, final_content: bool = False, table_profiles=None, before=None) -> dict:
    doc = Document(path)
    findings: list[dict] = []

    def add(code: str, message: str, severity: str = "error"):
        if mode == "surgical":
            severity = "warning"  # global inspection cannot establish the authorized edit scope
        findings.append({"severity": severity, "code": code, "message": message})

    def advise(code: str, message: str):
        findings.append({"severity": "warning", "code": code, "message": message})

    if len(doc.sections) != 1:
        add("PAGE.SECTIONS", f"应为单节，实际 {len(doc.sections)} 节")
    for index, section in enumerate(doc.sections, 1):
        expected = {
            "page_width": PAGE_STYLE["page_width_cm"], "page_height": PAGE_STYLE["page_height_cm"],
            "top_margin": PAGE_STYLE["top_margin_cm"], "bottom_margin": PAGE_STYLE["bottom_margin_cm"],
            "left_margin": PAGE_STYLE["left_margin_cm"], "right_margin": PAGE_STYLE["right_margin_cm"],
        }
        for prop, value in expected.items():
            if abs(_cm(getattr(section, prop)) - value) > TOL:
                add("PAGE.GEOMETRY", f"第{index}节 {prop}={_cm(getattr(section, prop)):.2f} cm，要求 {value:.2f} cm")

    style_contracts = {"Normal": BODY_STYLE, TITLE_STYLE["name"]: TITLE_STYLE}
    style_contracts.update({v["name"]: v for v in HEADING_STYLES.values()})
    style_contracts.update(SPECIAL_STYLES)
    for name, contract in style_contracts.items():
        try:
            style = doc.styles[name]
        except KeyError:
            add("STYLE.MISSING", f"缺少样式：{name}")
            continue
        rpr = style.element.rPr
        fonts = rpr.find(qn("w:rFonts")) if rpr is not None else None
        if _attr(fonts, "w:eastAsia") != contract["east_asia"] or _attr(fonts, "w:ascii") != contract["latin"]:
            add("STYLE.FONT", f"{name} 字体不符合约定")
        size = rpr.find(qn("w:sz")) if rpr is not None else None
        if _attr(size, "w:val") != str(int(contract["size_pt"] * 2)):
            add("STYLE.SIZE", f"{name} 字号不符合约定")
        color = rpr.find(qn("w:color")) if rpr is not None else None
        if _attr(color, "w:val").upper() != contract.get("color", "000000"):
            add("STYLE.COLOR", f"{name} 颜色不符合约定")

    levels = {v["name"]: k for k, v in HEADING_STYLES.items()}
    previous = 0
    for index, paragraph in enumerate(doc.paragraphs, 1):
        level = levels.get((paragraph.style.name if paragraph.style is not None else ""))
        if level:
            if previous and level > previous + 1:
                add("HEADING.LEVEL_JUMP", f"第{index}段标题级别从{previous}跳到{level}")
            previous = level
            if re.match(r"^\s*\d+\.", paragraph.text):
                add("HEADING.MANUAL_NUMBER", f"第{index}段疑似手工编号：{paragraph.text[:40]}")
            ppr = paragraph._p.pPr
            numpr = ppr.find(qn("w:numPr")) if ppr is not None else None
            style_numpr = paragraph.style.element.pPr.find(qn("w:numPr")) if paragraph.style.element.pPr is not None else None
            if numpr is None and style_numpr is None and not paragraph.text.startswith("附录："):
                add("HEADING.NO_NUMBERING", f"第{index}段标题未绑定自动编号")
        if (paragraph.style.name if paragraph.style is not None else "") in ("图片", "图表题注"):
            ppr = paragraph._p.pPr
            ind = ppr.find(qn("w:ind")) if ppr is not None else None
            if ind is not None and any(_attr(ind, key) not in ("", "0") for key in ("w:firstLine", "w:firstLineChars", "w:left", "w:leftChars")):
                add("OBJECT.INDENT", f"第{index}段图片/题注存在缩进")
            if (paragraph.style.name if paragraph.style is not None else "") == "图表题注" and not re.match(r"^(图|表)\d+\s+\S", paragraph.text):
                add("CAPTION.FORMAT", f"第{index}段题注缺少可见编号或内容")

    for ti, table in enumerate(doc.tables, 1):
        tblpr = table._tbl.tblPr
        layout = tblpr.find(qn("w:tblLayout"))
        if _attr(layout, "w:type") != "fixed": add("TABLE.LAYOUT", f"表{ti}不是固定布局")
        ind = tblpr.find(qn("w:tblInd"))
        if ind is not None and _attr(ind, "w:w") not in ("", "0"): add("TABLE.INDENT", f"表{ti}存在表格缩进")
        if not table.rows: continue
        trpr = table.rows[0]._tr.trPr
        if trpr is None or trpr.find(qn("w:tblHeader")) is None: add("TABLE.REPEAT_HEADER", f"表{ti}未设置重复表头")
        for ri, row in enumerate(table.rows, 1):
            trpr = row._tr.trPr
            if trpr is None or trpr.find(qn("w:cantSplit")) is None: add("TABLE.CANT_SPLIT", f"表{ti}第{ri}行允许跨页拆分")
            for ci, cell in enumerate(row.cells, 1):
                tcpr = cell._tc.tcPr
                shd = tcpr.find(qn("w:shd")) if tcpr is not None else None
                fill = _attr(shd, "w:fill").upper()
                if ri == 1 and fill != TABLE_STYLE["header_fill"]: add("TABLE.HEADER_FILL", f"表{ti}表头填充错误")
                for paragraph in cell.paragraphs:
                    ppr = paragraph._p.pPr
                    pind = ppr.find(qn("w:ind")) if ppr is not None else None
                    if pind is not None and any(_attr(pind, key) not in ("", "0") for key in ("w:firstLine", "w:firstLineChars", "w:left")):
                        add("TABLE.PARAGRAPH_INDENT", f"表{ti}第{ri}行第{ci}列存在段落缩进")
                    if ri == 1:
                        for run in paragraph.runs:
                            color = run._r.rPr.find(qn("w:color")) if run._r.rPr is not None else None
                            if _attr(color, "w:val").upper() != TABLE_STYLE["header_font_color"]:
                                add("TABLE.HEADER_FONT", f"表{ti}表头文字不是白色")
                                break

    with zipfile.ZipFile(path) as archive:
        image_count = sum(name.startswith("word/media/") for name in archive.namelist())
    if final_content:
        joined = "\n".join(p.text for p in doc.paragraphs) + "\n" + "\n".join(cell.text for t in doc.tables for row in t.rows for cell in row.cells)
        for token in PLACEHOLDERS:
            if token in joined:
                add("CONTENT.PLACEHOLDER", f"残留占位内容：{token}")
    if mode == "new" and not any(p.text.strip() == "附录：数据溯源表" for p in doc.paragraphs):
        add("TRACE.APPENDIX", "缺少文末数据溯源附录")
    rework_pre = "改焊前"
    rework_post = "改焊后"
    for t_index, table in enumerate(doc.tables, 1):
        row_texts = []
        for row in table.rows:
            texts = [c.text.strip() for c in row.cells]
            row_texts.append(texts)
            if not any(k in texts for k in (rework_pre, rework_post)):
                continue
        status_idx = [i for i, tx in enumerate(row_texts) if rework_pre in tx or rework_post in tx]
        if not status_idx:
            continue
        post_names = {row_texts[i][0] for i in status_idx
                      if rework_post in row_texts[i] and row_texts[i] and row_texts[i][0]}
        for n, i in enumerate(status_idx):
            texts = row_texts[i]
            first = texts[0] if texts else ""
            is_post = rework_post in texts
            is_pre = rework_pre in texts
            # 规则A（孤行改焊前）：对比组改焊前行（接口格空/全角后缀/与改焊后行同名）
            # 其紧邻上一状态行必须是改焊后行。
            if is_pre and (not first or first.endswith("（改焊前）") or first in post_names):
                prev = row_texts[status_idx[n - 1]] if n > 0 else None
                if not (prev is not None and rework_post in prev):
                    advise("REWORK.ORDER", f"第{t_index}张表：改焊前行未紧跟其改焊后行，应改焊后在前")
            # 规则B（前后互换）：改焊后行紧邻上一行是改焊前行且同接口（或接口名为空）。
            if is_post and n > 0:
                prev = row_texts[status_idx[n - 1]]
                prev_first = prev[0] if prev else ""
                if rework_pre in prev and (not first or first == prev_first):
                    advise("REWORK.ORDER", f"第{t_index}张表：改焊前行出现在改焊后行之前，应改焊后在前")
            # 规则C（灰字）：对比组改焊前状态格（同规则A范围）的可见文字应为灰字 595959；
            # 未复测接口单行（独有接口名+改焊前、黑字）是合法样式，不检查。
            if is_pre and (not first or first.endswith("（改焊前）") or first in post_names):
                for cell in table.rows[i].cells:
                    if cell.text.strip() != rework_pre:
                        continue
                    for paragraph in cell.paragraphs:
                        for run in paragraph.runs:
                            if not run.text.strip():
                                continue
                            rpr = run._r.rPr
                            if rpr is None:
                                continue
                            node = rpr.find(qn("w:color"))
                            if node is not None and node.get(qn("w:val")) != "595959":
                                advise("REWORK.GRAY", f"第{t_index}张表：改焊前状态格非灰字(595959)")

    for t_index, table in enumerate(doc.tables, 1):
        info = []
        for tr in table._tbl.tr_lst:
            tcs = tr.tc_lst
            if not tcs:
                info.append(("none", "", ""))
                continue
            tc0 = tcs[0]
            vm = tc0.tcPr.find(qn("w:vMerge")) if tc0.tcPr is not None else None
            if vm is None:
                state = "none"
            elif vm.get(qn("w:val")) == "restart":
                state = "restart"
            else:
                state = "cont"
            text = "".join(n.text or "" for n in tc0.iter(qn("w:t"))).strip()
            stext = "".join(n.text or "" for n in tcs[1].iter(qn("w:t"))).strip() if len(tcs) > 1 else ""
            info.append((state, text, stext))
        for i, (state, text, stext) in enumerate(info):
            if stext != "改焊前":
                continue
            prev = info[i - 1] if i > 0 else None
            in_pair = (not text) or text.endswith("（改焊前）") or (
                prev is not None and prev[2] == "改焊后" and (prev[1] == text or not text))
            if not in_pair:
                continue  # 未复测接口单行，不要求合并
            if state != "cont":
                advise("REWORK.MERGE", f"第{t_index}张表：改焊前行接口格未纵向合并（应为 vMerge 延续格）")

    previous = Counter()
    baseline_doc = Document(before) if mode == "surgical" and before else None
    def unchanged_area(f):
        if baseline_doc is None: return False
        def area(document):
            node=document.tables[f["table"]-1]._tbl
            if f["row"]: node=node.tr_lst[f["row"]-1]
            if f.get("cell"): node=node.tc_lst[f["cell"]-1]
            return node.xml
        try: return area(baseline_doc)==area(doc)
        except IndexError: return False
    def identity(f): return (f['code'],f['table'],f['row'],f.get('cell',0),f['message'])
    if mode == "surgical" and before:
        previous = Counter(identity(f) for f in table_checks(Document(before), table_profiles))
    for finding in table_checks(doc, table_profiles):
        key = identity(finding)
        historical = mode == "surgical" and before and previous[key] > 0 and unchanged_area(finding)
        if historical: previous[key] -= 1
        findings.append(dict(finding, severity="warning" if historical else "error",
                             baseline="existing" if historical else "new_or_unverified"))
    for f in reference_checks(doc):
        findings.append(dict(severity="warning" if mode=="surgical" else "error",code="CAPTION.REFERENCE",message=f["location"]+": "+f["text"]))
    # Dangling figure references and changed caption defects are also incremental failures.
    if mode == "surgical" and before:
        old = audit(Path(before), "surgical", final_content, table_profiles)
        counts = Counter((f['code'],f['message']) for f in old['findings'])
        for f in findings:
            if f['code'].startswith(("CAPTION.", "OBJECT.", "CONTENT.")):
                k=(f['code'],f['message'])
                if counts[k]: counts[k]-=1
                else: f['severity']='error'

    return {
        "path": str(path.resolve()), "mode": mode, "final_content": final_content,
        "errors": sum(x["severity"] == "error" for x in findings),
        "warnings": sum(x["severity"] == "warning" for x in findings),
        "tables": len(doc.tables), "images": image_count, "findings": findings,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("docx", type=Path)
    parser.add_argument("--mode", choices=("new", "surgical"), required=True)
    parser.add_argument("--final-content", action="store_true")
    parser.add_argument("--json", action="store_true")
    parser.add_argument("--strict", action="store_true")
    parser.add_argument("--before", type=Path)
    parser.add_argument("--scope-json", type=Path)
    parser.add_argument("--table-profiles", type=Path)
    args = parser.parse_args()
    if bool(args.before) != bool(args.scope_json):
        parser.error("--before and --scope-json must be supplied together")
    profiles = json.loads(args.table_profiles.read_text(encoding="utf-8-sig")) if args.table_profiles else None
    report = audit(args.docx, args.mode, args.final_content, profiles, args.before)
    if args.before:
        scope = json.loads(args.scope_json.read_text(encoding="utf-8-sig"))
        issues = surgical_diff(args.before, args.docx, scope)
        report["scope_verified"] = not issues
        report["findings"].extend({"severity":"error","code":"SURGICAL.OUTSIDE_SCOPE","message":issue} for issue in issues)
        report["errors"] += len(issues)
    elif args.mode == "surgical":
        report["scope_verified"] = False
        report["findings"].append({"severity":"warning","code":"SURGICAL.UNVERIFIED","message":"global checks only; supply --before and --scope-json for acceptance"})
        report["warnings"] += 1
    if args.strict and args.mode == "surgical" and not args.before:
        report["errors"] += 1
        report["findings"].append(dict(severity="error",code="SURGICAL.BASELINE_REQUIRED",message="strict surgical acceptance requires baseline and scope"))
    report["acceptance"] = "declared_checks_passed" if not report["errors"] and (args.mode != "surgical" or report.get("scope_verified")) else "not_accepted"
    print(json.dumps(report, ensure_ascii=False, indent=2) if args.json else f"errors={report['errors']} warnings={report['warnings']}")
    raise SystemExit(1 if report["errors"] else 0)


if __name__ == "__main__":
    main()

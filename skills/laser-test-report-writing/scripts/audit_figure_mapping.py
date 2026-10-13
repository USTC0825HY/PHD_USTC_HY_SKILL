#!/usr/bin/env python3
"""Audit DOCX image, caption, evidence-ID, reference, and traceability mapping."""

from __future__ import annotations

import argparse
import csv
import json
import re
import sys
import zipfile
from dataclasses import asdict, dataclass
from pathlib import Path

from docx import Document
from lxml import etree


W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
NS = {"w": W, "a": A}
REQUIRED_COLUMNS = [
    "sequence",
    "evidence_id",
    "object",
    "interface",
    "content_type",
    "caption",
    "source_filename",
    "word_section",
]
EVIDENCE_ID_RE = re.compile(r"^[A-Z0-9]+(?:-[A-Z0-9]+){3,}$")
VISIBLE_CAPTION_RE = re.compile(r"^图\s*(\d+)\s*(.*)$")
FIGURE_REFERENCE_RE = re.compile(r"图\s*(\d+)(?:\s*[～~-]\s*(?:图\s*)?(\d+))?")
FILENAME_RE = re.compile(r"(?i)\.(?:png|jpe?g|bmp|gif|tiff?)$")


@dataclass
class FigureAudit:
    path: str
    manifest: str
    image_count: int
    caption_count: int
    manifest_count: int
    figure_mapping_issues: list[dict[str, str]]
    duplicate_evidence_id_issues: list[dict[str, str]]
    cross_reference_issues: list[dict[str, str]]
    traceability_table_issues: list[dict[str, str]]
    warnings: list[dict[str, str]]

    @property
    def passed(self) -> bool:
        return not (
            self.figure_mapping_issues
            or self.duplicate_evidence_id_issues
            or self.cross_reference_issues
            or self.traceability_table_issues
        )


def text_of(element: etree._Element) -> str:
    return "".join(element.xpath(".//w:t/text()", namespaces=NS)).strip()


def has_image(paragraph: etree._Element) -> bool:
    return bool(paragraph.xpath(".//a:blip | .//w:pict", namespaces=NS))


def has_seq_figure(paragraph: etree._Element) -> bool:
    instruction = " ".join(paragraph.xpath(".//w:instrText/text()", namespaces=NS))
    return bool(re.search(r"\bSEQ\s+图\b", instruction))


def document_containers(root: etree._Element) -> list[tuple[str, etree._Element]]:
    containers: list[tuple[str, etree._Element]] = [("body", root.find(f"{{{W}}}body"))]
    for index, cell in enumerate(root.xpath(".//w:tc", namespaces=NS), 1):
        containers.append((f"table-cell-{index}", cell))
    return [(name, element) for name, element in containers if element is not None]


def image_caption_pairs(root: etree._Element, caption_styles=()) -> tuple[int, list[dict[str, object]], list[dict[str, str]]]:
    image_count = len(root.xpath(".//a:blip", namespaces=NS))
    captions: list[dict[str, object]] = []
    issues: list[dict[str, str]] = []
    paired_count = 0

    for container_name, container in document_containers(root):
        children = list(container)
        for index, child in enumerate(children):
            if child.tag != f"{{{W}}}p":
                continue
            previous = next((children[pos] for pos in range(index - 1, -1, -1) if text_of(children[pos]) or has_image(children[pos])), None)
            style_nodes = child.xpath("./w:pPr/w:pStyle/@w:val", namespaces=NS)
            plain_caption = VISIBLE_CAPTION_RE.match(text_of(child)) and (
                (previous is not None and has_image(previous)) or
                any(name in caption_styles for name in style_nodes))
            if not (has_seq_figure(child) or plain_caption):
                continue
            visible = text_of(child)
            match = VISIBLE_CAPTION_RE.match(visible)
            number = int(match.group(1)) if match else -1
            description = match.group(2).strip() if match else visible
            captions.append(
                {
                    "number": number,
                    "description": description,
                    "visible": visible,
                    "container": container_name,
                    "relationship_ids": previous.xpath(".//a:blip/@r:embed", namespaces={**NS,"r":"http://schemas.openxmlformats.org/officeDocument/2006/relationships"}) if previous is not None else [],
                }
            )
            if re.match(r"^图\s+\d", visible):
                issues.append({"location": container_name, "issue": "space between 图 and figure number", "text": visible})
            if re.search(r"\[[A-Z0-9-]+\]", visible):
                issues.append({"location": container_name, "issue": "visible caption exposes evidence ID", "text": visible})
            previous = next((children[pos] for pos in range(index - 1, -1, -1) if text_of(children[pos]) or has_image(children[pos])), None)
            if previous is None or previous.tag != f"{{{W}}}p" or not has_image(previous):
                issues.append({"location": container_name, "issue": "figure caption is not immediately below an image", "text": visible})
            else:
                paired_count += 1

    if paired_count != image_count:
        issues.append(
            {
                "location": "document",
                "issue": "inserted images and adjacent captions are not one-to-one",
                "text": f"images={image_count}, paired={paired_count}",
            }
        )
    return image_count, captions, issues


def read_manifest(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as stream:
        reader = csv.DictReader(stream)
        missing = [name for name in REQUIRED_COLUMNS if name not in (reader.fieldnames or [])]
        if missing:
            raise ValueError(f"manifest missing columns: {', '.join(missing)}")
        return [{key: (value or "").strip() for key, value in row.items()} for row in reader]


def normalized_caption(text: str) -> str:
    return re.sub(r"\s+", "", text).replace("：", ":")


def traceability_rows(doc: Document) -> tuple[str | None, list[dict[str, str]]]:
    for table in doc.tables:
        if not table.rows:
            continue
        headers = [cell.text.strip().replace("\n", "") for cell in table.rows[0].cells]
        if headers[:5] == ["图号", "报告内编号", "对象/接口", "图像内容", "原始文件名（仅追溯）"]:
            rows = []
            for row in table.rows[1:]:
                cells = [cell.text.strip() for cell in row.cells]
                match = re.search(r"图\s*(\d+)", cells[0])
                rows.append(
                    {
                        "sequence": match.group(1) if match else "",
                        "evidence_id": cells[1],
                        "object_interface": cells[2],
                        "content": cells[3],
                        "source_filename": cells[4],
                    }
                )
            return "five-column", rows
        if headers[:4] == ["图号及报告内编号", "器件/接口", "图像内容", "原始文件名（仅追溯）"]:
            rows = []
            for row in table.rows[1:]:
                cells = [cell.text.strip() for cell in row.cells]
                number = re.search(r"图\s*(\d+)", cells[0])
                evidence = re.search(r"([A-Z0-9]+(?:-[A-Z0-9]+)+)", cells[0])
                rows.append(
                    {
                        "sequence": number.group(1) if number else "",
                        "evidence_id": evidence.group(1) if evidence else "",
                        "object_interface": cells[1],
                        "content": cells[2],
                        "source_filename": cells[3],
                    }
                )
            return "legacy-combined", rows
    return None, []


def audit_document(docx: Path, manifest_path: Path, mode: str = "new", strict: bool = False, before: Path | None = None) -> FigureAudit:
    with zipfile.ZipFile(docx) as archive:
        root = etree.fromstring(archive.read("word/document.xml"))
    doc = Document(docx)
    manifest = read_manifest(manifest_path)
    caption_styles = {style.style_id for style in doc.styles if style.name in ("图表题注", "Caption", "题注")}
    image_count, captions, figure_issues = image_caption_pairs(root, caption_styles)
    duplicate_issues: list[dict[str, str]] = []
    reference_issues: list[dict[str, str]] = []
    trace_issues: list[dict[str, str]] = []
    warnings: list[dict[str, str]] = []

    sequences: list[int] = []
    evidence_ids: list[str] = []
    for row_index, row in enumerate(manifest, 2):
        try:
            sequences.append(int(row["sequence"]))
        except ValueError:
            figure_issues.append({"location": f"manifest row {row_index}", "issue": "sequence is not an integer", "text": row["sequence"]})
        evidence_id = row["evidence_id"]
        evidence_ids.append(evidence_id)
        if not EVIDENCE_ID_RE.fullmatch(evidence_id):
            warnings.append({"location": f"manifest row {row_index}", "issue": "legacy or noncanonical evidence ID", "text": evidence_id})

    for evidence_id in sorted(set(evidence_ids)):
        count = evidence_ids.count(evidence_id)
        if count > 1:
            duplicate_issues.append({"location": "manifest", "issue": "duplicate evidence ID", "text": f"{evidence_id} x{count}"})

    expected = list(range(1, len(manifest) + 1))
    caption_numbers = [int(item["number"]) for item in captions]
    if image_count != len(manifest) or len(captions) != len(manifest):
        figure_issues.append({"location": "document", "issue": "image/caption/manifest counts differ", "text": f"images={image_count}, captions={len(captions)}, manifest={len(manifest)}"})
    if mode != "surgical" and sequences != expected:
        figure_issues.append({"location": "manifest", "issue": "manifest sequence is not continuous", "text": str(sequences)})
    if mode != "surgical" and caption_numbers != expected:
        figure_issues.append({"location": "captions", "issue": "visible figure numbers are not continuous", "text": str(caption_numbers)})

    if len(set(caption_numbers)) != len(caption_numbers) or caption_numbers != sequences:
        figure_issues.append({"location":"captions","issue":"duplicate numbers or caption/manifest number mismatch","text":str(caption_numbers)})
    for index, (caption, row) in enumerate(zip(captions, manifest), 1):
        if normalized_caption(str(caption["description"])) != normalized_caption(row["caption"]):
            figure_issues.append({"location": f"figure {index}", "issue": "visible caption differs from manifest caption", "text": f"docx={caption['description']} | manifest={row['caption']}"})

    all_text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
    all_text += "\n" + "\n".join(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
    max_figure = max(caption_numbers, default=0)
    for match in FIGURE_REFERENCE_RE.finditer(all_text):
        start, end = int(match.group(1)), int(match.group(2) or match.group(1))
        if start < 1 or end < start or end > max_figure or any(number not in caption_numbers for number in range(start, min(end, max_figure) + 1)):
            reference_issues.append({"location": "document", "issue": "figure reference is out of range", "text": match.group(0)})

    for table_index, table in enumerate(doc.tables, 1):
        if not table.rows or not table.rows[0].cells:
            continue
        first_header = table.rows[0].cells[0].text.strip().replace("\n", "")
        if first_header == "图号/报告内编号":
            trace_issues.append(
                {
                    "location": f"table {table_index}",
                    "issue": "main result table exposes evidence IDs",
                    "text": "use 图号 only; keep evidence IDs in the traceability table",
                }
            )

    trace_format, trace_rows = traceability_rows(doc)
    # Optional figure-only index is distinct from the required seven-column data appendix.
    if trace_format is not None:
        if trace_format == "legacy-combined":
            warnings.append({"location": "traceability table", "issue": "legacy combined figure-number/evidence-ID column", "text": "use five columns for new reports"})
        if len(trace_rows) != len(manifest):
            trace_issues.append({"location": "traceability table", "issue": "traceability row count differs from manifest", "text": f"rows={len(trace_rows)}, manifest={len(manifest)}"})
        for index, (trace, row) in enumerate(zip(trace_rows, manifest), 1):
            if trace["sequence"] != row["sequence"] or trace["evidence_id"] != row["evidence_id"]:
                trace_issues.append({"location": f"traceability row {index + 1}", "issue": "figure number or evidence ID differs from manifest", "text": f"trace={trace['sequence']}/{trace['evidence_id']} manifest={row['sequence']}/{row['evidence_id']}"})
            if trace["source_filename"] != row["source_filename"]:
                trace_issues.append({"location": f"traceability row {index + 1}", "issue": "source filename differs from manifest", "text": f"trace={trace['source_filename']} manifest={row['source_filename']}"})

    trace_source_names = {row["source_filename"] for row in manifest if row["source_filename"]}
    trace_table_text = "\n".join(row["source_filename"] for row in trace_rows)
    trace_table_text += "\n" + "\n".join(cell.text for table in doc.tables if table.rows and "源数据路径" in " ".join(c.text for c in table.rows[0].cells) for row in table.rows for cell in row.cells)
    for source_name in trace_source_names:
        if FILENAME_RE.search(source_name) and all_text.count(source_name) > trace_table_text.count(source_name):
            trace_issues.append({"location": "main report", "issue": "source filename appears outside traceability table", "text": source_name})

    if strict and mode == "surgical" and before:
        from collections import Counter
        baseline = audit_document(before, manifest_path, "surgical")
        for current, previous in ((figure_issues,baseline.figure_mapping_issues),(duplicate_issues,baseline.duplicate_evidence_id_issues),(reference_issues,baseline.cross_reference_issues),(trace_issues,baseline.traceability_table_issues)):
            counts=Counter(json.dumps(x,sort_keys=True,ensure_ascii=False) for x in previous)
            new=[]
            for issue in current:
                key=json.dumps(issue,sort_keys=True,ensure_ascii=False)
                if counts[key]:
                    counts[key]-=1
                    warnings.append(dict(issue,issue="existing baseline defect: "+issue['issue']))
                else: new.append(issue)
            current[:]=new
    if strict:
        from figure_identity import identity_checks
        identity_issues, identity_notes = identity_checks(docx, manifest, manifest_path, mode, before)
        figure_issues.extend(identity_issues)
        warnings.extend(identity_notes)
    else:
        warnings.append(dict(location="document", issue="structure_only; image identity not verified", text="use --strict"))

    return FigureAudit(
        path=str(docx),
        manifest=str(manifest_path),
        image_count=image_count,
        caption_count=len(captions),
        manifest_count=len(manifest),
        figure_mapping_issues=figure_issues,
        duplicate_evidence_id_issues=duplicate_issues,
        cross_reference_issues=reference_issues,
        traceability_table_issues=trace_issues,
        warnings=warnings,
    )


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("docx", type=Path)
    parser.add_argument("--manifest", required=True, type=Path)
    parser.add_argument("--json-out", type=Path)
    parser.add_argument("--mode", choices=("new","surgical"), default="new")
    parser.add_argument("--strict", action="store_true")
    parser.add_argument("--before", type=Path)
    args = parser.parse_args()
    result = audit_document(args.docx.resolve(), args.manifest.resolve(), args.mode, args.strict, args.before)
    payload = asdict(result) | {"passed": result.passed, "verification_level": "strict_identity" if args.strict else "structure_only"}
    text = json.dumps(payload, ensure_ascii=False, indent=2)
    if args.json_out:
        args.json_out.parent.mkdir(parents=True, exist_ok=True)
        args.json_out.write_text(text + "\n", encoding="utf-8")
    print(text)
    return 0 if result.passed else 1


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    raise SystemExit(main())

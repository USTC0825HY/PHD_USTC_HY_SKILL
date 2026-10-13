#!/usr/bin/env python3
"""Create the canonical compact Laser test-result summary DOCX template."""

from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from result_summary_style_contract import (
    BODY_STYLE,
    HEADING_STYLES,
    NUMBERING_LEVELS,
    PAGE_STYLE,
    SPECIAL_STYLES,
    TABLE_STYLE,
    TITLE_STYLE,
)


def _set_attr(element, name: str, value: str) -> None:
    element.set(qn(name), str(value))


def _remove_children(parent, tag: str) -> None:
    for child in list(parent.findall(qn(tag))):
        parent.remove(child)


def _set_run_properties(rpr, *, east_asia: str, latin: str, size_pt: float, bold: bool, color: str) -> None:
    fonts = rpr.get_or_add_rFonts()
    for key in list(fonts.attrib):
        if "theme" in key.lower():
            del fonts.attrib[key]
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        _set_attr(fonts, attr, latin)
    _set_attr(fonts, "w:eastAsia", east_asia)

    for tag in ("w:sz", "w:szCs"):
        node = rpr.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            rpr.append(node)
        _set_attr(node, "w:val", str(int(round(size_pt * 2))))

    for tag in ("w:b", "w:bCs"):
        node = rpr.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            rpr.append(node)
        _set_attr(node, "w:val", "1" if bold else "0")

    color_node = rpr.find(qn("w:color"))
    if color_node is None:
        color_node = OxmlElement("w:color")
        rpr.append(color_node)
    _set_attr(color_node, "w:val", color)


def _set_zero_indent(ppr) -> None:
    ind = ppr.find(qn("w:ind"))
    if ind is None:
        ind = OxmlElement("w:ind")
        ppr.append(ind)
    for attr in ("w:firstLine", "w:firstLineChars", "w:left", "w:leftChars", "w:right", "w:rightChars"):
        _set_attr(ind, attr, "0")


def _set_alignment(ppr, value: str) -> None:
    jc = ppr.find(qn("w:jc"))
    if jc is None:
        jc = OxmlElement("w:jc")
        ppr.append(jc)
    _set_attr(jc, "w:val", value)


def _set_spacing(ppr, before_pt: float, after_pt: float) -> None:
    spacing = ppr.find(qn("w:spacing"))
    if spacing is None:
        spacing = OxmlElement("w:spacing")
        ppr.append(spacing)
    _set_attr(spacing, "w:before", str(int(round(before_pt * 20))))
    _set_attr(spacing, "w:after", str(int(round(after_pt * 20))))


def _set_toggle(ppr, tag: str, enabled: bool = True) -> None:
    node = ppr.find(qn(tag))
    if enabled and node is None:
        ppr.append(OxmlElement(tag))
    elif not enabled and node is not None:
        ppr.remove(node)


def _style(doc: Document, name: str, style_type=WD_STYLE_TYPE.PARAGRAPH):
    try:
        return doc.styles[name]
    except KeyError:
        return doc.styles.add_style(name, style_type)


def configure_page(doc: Document) -> None:
    if len(doc.sections) != 1:
        raise ValueError("Result-summary documents must contain exactly one section.")
    section = doc.sections[0]
    section.page_width = Cm(PAGE_STYLE["page_width_cm"])
    section.page_height = Cm(PAGE_STYLE["page_height_cm"])
    section.top_margin = Cm(PAGE_STYLE["top_margin_cm"])
    section.bottom_margin = Cm(PAGE_STYLE["bottom_margin_cm"])
    section.left_margin = Cm(PAGE_STYLE["left_margin_cm"])
    section.right_margin = Cm(PAGE_STYLE["right_margin_cm"])
    section.header_distance = Cm(PAGE_STYLE["header_distance_cm"])
    section.footer_distance = Cm(PAGE_STYLE["footer_distance_cm"])
    for part in (section.header, section.footer):
        for paragraph in part.paragraphs:
            paragraph.clear()


def _configure_paragraph_style(style, contract: dict, *, zero_indent: bool) -> None:
    rpr = style.element.get_or_add_rPr()
    _set_run_properties(
        rpr,
        east_asia=contract["east_asia"],
        latin=contract["latin"],
        size_pt=contract["size_pt"],
        bold=bool(contract.get("bold", False)),
        color=contract.get("color", "000000"),
    )
    ppr = style.element.get_or_add_pPr()
    if zero_indent:
        _set_zero_indent(ppr)
    if contract.get("alignment"):
        _set_alignment(ppr, contract["alignment"])
    _set_toggle(ppr, "w:keepNext", bool(contract.get("keep_with_next", False)))


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    _configure_paragraph_style(normal, BODY_STYLE | {"bold": False}, zero_indent=False)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    ppr = normal.element.get_or_add_pPr()
    ind = ppr.find(qn("w:ind"))
    if ind is None:
        ind = OxmlElement("w:ind")
        ppr.append(ind)
    for attr in ("w:left", "w:leftChars", "w:right", "w:rightChars"):
        _set_attr(ind, attr, "0")
    _set_attr(ind, "w:firstLineChars", str(BODY_STYLE["first_line_chars"]))
    _set_attr(ind, "w:firstLine", str(int(round(BODY_STYLE["size_pt"] * 2 * 20))))

    title = _style(doc, TITLE_STYLE["name"])
    title.base_style = normal
    _configure_paragraph_style(title, TITLE_STYLE, zero_indent=True)

    for level, contract in HEADING_STYLES.items():
        style = _style(doc, contract["name"])
        style.base_style = normal
        _configure_paragraph_style(style, contract, zero_indent=True)
        ppr = style.element.get_or_add_pPr()
        _set_spacing(ppr, contract["space_before_pt"], contract["space_after_pt"])
        _set_toggle(ppr, "w:keepNext", True)
        _set_toggle(ppr, "w:keepLines", True)
        outline = ppr.find(qn("w:outlineLvl"))
        if outline is None:
            outline = OxmlElement("w:outlineLvl")
            ppr.append(outline)
        _set_attr(outline, "w:val", str(contract["outline_level"]))
        qformat = style.element.find(qn("w:qFormat"))
        if qformat is None:
            style.element.append(OxmlElement("w:qFormat"))

    for name, contract in SPECIAL_STYLES.items():
        style = _style(doc, name)
        style.base_style = normal
        _configure_paragraph_style(style, contract, zero_indent=True)
        style.paragraph_format.line_spacing = 1.0
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(0)

    _style(doc, TABLE_STYLE["name"], WD_STYLE_TYPE.TABLE)


def _next_numeric_id(elements, attribute: str) -> int:
    values = []
    for element in elements:
        raw = element.get(qn(attribute))
        if raw is not None and raw.isdigit():
            values.append(int(raw))
    return max(values, default=0) + 1


def bind_heading_numbering(doc: Document) -> int:
    numbering = doc.part.numbering_part.element

    for style_contract in HEADING_STYLES.values():
        ppr = doc.styles[style_contract["name"]].element.get_or_add_pPr()
        _remove_children(ppr, "w:numPr")

    abstract_id = _next_numeric_id(numbering.findall(qn("w:abstractNum")), "w:abstractNumId")
    num_id = _next_numeric_id(numbering.findall(qn("w:num")), "w:numId")

    abstract = OxmlElement("w:abstractNum")
    _set_attr(abstract, "w:abstractNumId", str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    _set_attr(multi, "w:val", "multilevel")
    abstract.append(multi)

    for level, text, left, hanging in NUMBERING_LEVELS:
        lvl = OxmlElement("w:lvl")
        _set_attr(lvl, "w:ilvl", str(level))
        start = OxmlElement("w:start")
        _set_attr(start, "w:val", "1")
        lvl.append(start)
        num_fmt = OxmlElement("w:numFmt")
        _set_attr(num_fmt, "w:val", "decimal")
        lvl.append(num_fmt)
        pstyle = OxmlElement("w:pStyle")
        _set_attr(pstyle, "w:val", doc.styles[HEADING_STYLES[level + 1]["name"]].style_id)
        lvl.append(pstyle)
        lvl_text = OxmlElement("w:lvlText")
        _set_attr(lvl_text, "w:val", text)
        lvl.append(lvl_text)
        suff = OxmlElement("w:suff")
        _set_attr(suff, "w:val", "nothing")
        lvl.append(suff)
        ppr = OxmlElement("w:pPr")
        ind = OxmlElement("w:ind")
        _set_attr(ind, "w:left", str(left))
        _set_attr(ind, "w:hanging", str(hanging))
        ppr.append(ind)
        lvl.append(ppr)
        abstract.append(lvl)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    _set_attr(num, "w:numId", str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    _set_attr(abstract_ref, "w:val", str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)

    for level, contract in HEADING_STYLES.items():
        ppr = doc.styles[contract["name"]].element.get_or_add_pPr()
        num_pr = OxmlElement("w:numPr")
        ilvl = OxmlElement("w:ilvl")
        _set_attr(ilvl, "w:val", str(level - 1))
        num_id_node = OxmlElement("w:numId")
        _set_attr(num_id_node, "w:val", str(num_id))
        num_pr.append(ilvl)
        num_pr.append(num_id_node)
        ppr.insert(0, num_pr)
    return num_id


def _set_paragraph_zero_center(paragraph, style_name: str) -> None:
    paragraph.style = style_name
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    ppr = paragraph._p.get_or_add_pPr()
    _set_zero_indent(ppr)


def _set_cell_fill(cell, fill: str) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcpr.append(shd)
    _set_attr(shd, "w:val", "clear")
    _set_attr(shd, "w:color", "auto")
    _set_attr(shd, "w:fill", fill)


def _set_cell_margins(cell) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    margins = tcpr.find(qn("w:tcMar"))
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tcpr.append(margins)
    values = {
        "top": TABLE_STYLE["cell_margin_top_twips"],
        "left": TABLE_STYLE["cell_margin_left_twips"],
        "bottom": TABLE_STYLE["cell_margin_bottom_twips"],
        "right": TABLE_STYLE["cell_margin_right_twips"],
    }
    for side, value in values.items():
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        _set_attr(node, "w:w", value)
        _set_attr(node, "w:type", "dxa")


def _set_run_direct(run, contract: dict, *, color: str | None = None) -> None:
    rpr = run._r.get_or_add_rPr()
    _set_run_properties(
        rpr,
        east_asia=contract["east_asia"],
        latin=contract["latin"],
        size_pt=contract["size_pt"],
        bold=bool(contract.get("bold", False)),
        color=color or contract.get("color", "000000"),
    )


def format_table(table, *, preserve_widths: bool = True) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    try:
        table.style = TABLE_STYLE["name"]
    except KeyError:
        pass
    tblpr = table._tbl.tblPr
    layout = tblpr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tblpr.append(layout)
    _set_attr(layout, "w:type", "fixed")
    ind = tblpr.find(qn("w:tblInd"))
    if ind is None:
        ind = OxmlElement("w:tblInd")
        tblpr.append(ind)
    _set_attr(ind, "w:w", "0")
    _set_attr(ind, "w:type", "dxa")

    borders = tblpr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tblpr.append(borders)
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            borders.append(node)
        _set_attr(node, "w:val", "single")
        _set_attr(node, "w:sz", TABLE_STYLE["border_size_eighth_pt"])
        _set_attr(node, "w:space", "0")
        _set_attr(node, "w:color", TABLE_STYLE["border_color"])

    for row_index, row in enumerate(table.rows):
        trpr = row._tr.get_or_add_trPr()
        if trpr.find(qn("w:cantSplit")) is None:
            trpr.append(OxmlElement("w:cantSplit"))
        header = trpr.find(qn("w:tblHeader"))
        if row_index == 0 and header is None:
            trpr.append(OxmlElement("w:tblHeader"))
        elif row_index != 0 and header is not None:
            trpr.remove(header)

        for cell in row.cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            _set_cell_margins(cell)
            value = cell.text.strip()
            existing_shd = cell._tc.get_or_add_tcPr().find(qn("w:shd"))
            existing_fill = existing_shd.get(qn("w:fill"), "").upper() if existing_shd is not None else ""
            if row_index == 0:
                fill = TABLE_STYLE["header_fill"]
                font_color = TABLE_STYLE["header_font_color"]
                style_name = "表格表头"
            else:
                fill = TABLE_STYLE["status_fills"].get(value, TABLE_STYLE["body_fill"])
                if existing_fill == TABLE_STYLE["exception_fill"]:
                    fill = TABLE_STYLE["exception_fill"]
                font_color = TABLE_STYLE["status_font_colors"].get(value, "000000")
                if fill == TABLE_STYLE["exception_fill"]:
                    font_color = TABLE_STYLE["exception_font_color"]
                style_name = "表格文字"
            _set_cell_fill(cell, fill)
            for paragraph in cell.paragraphs:
                _set_paragraph_zero_center(paragraph, style_name)
                for run in paragraph.runs:
                    _set_run_direct(run, SPECIAL_STYLES[style_name], color=font_color)


def format_rework_table(table, status_column: int = 1) -> None:
    """Apply rework before/after styling on top of the base table format.

    Post-rework rows (状态=改焊后) get the light-blue fill on status+data cells;
    pre-rework rows (状态=改焊前) get the light-red fill and gray font; the
    interface column (column 0) stays white. Rows with other status values are
    left as plain body rows.
    """
    format_table(table)
    keywords = TABLE_STYLE["rework_status_keywords"]
    for row_index, row in enumerate(table.rows):
        if row_index == 0:
            continue
        cells = row.cells
        if status_column >= len(cells):
            continue
        status = cells[status_column].text.strip()
        if status not in keywords:
            continue
        post = status == "改焊后"
        fill = TABLE_STYLE["rework_post_fill"] if post else TABLE_STYLE["rework_pre_fill"]
        font_color = "000000" if post else TABLE_STYLE["rework_pre_font_color"]
        for idx, cell in enumerate(cells):
            if idx == 0:
                _set_cell_fill(cell, "FFFFFF")
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        _set_run_direct(run, SPECIAL_STYLES["表格文字"], color=font_color)
                continue
            _set_cell_fill(cell, fill)
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    _set_run_direct(run, SPECIAL_STYLES["表格文字"], color=font_color)
    # Interface-column vertical merge for adjacent post/pre pairs (blueprint style)
    table_rows = table.rows
    i = 1
    while i + 1 < len(table_rows):
        cells_i = table_rows[i].cells
        cells_j = table_rows[i + 1].cells
        post_status = cells_i[status_column].text.strip() if status_column < len(cells_i) else ""
        pre_status = cells_j[status_column].text.strip() if status_column < len(cells_j) else ""
        if post_status == "改焊后" and pre_status == "改焊前":
            for cell in cells_i:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.keep_with_next = True
            post_tc = cells_i[0]._tc
            pre_tc = cells_j[0]._tc
            for tc, restart in ((post_tc, True), (pre_tc, False)):
                tcpr = tc.get_or_add_tcPr()
                for old in tcpr.findall(qn("w:vMerge")):
                    tcpr.remove(old)
                vm = OxmlElement("w:vMerge")
                if restart:
                    vm.set(qn("w:val"), "restart")
                tcpr.append(vm)
            for paragraph in cells_j[0].paragraphs:
                for run in list(paragraph.runs):
                    run._r.getparent().remove(run._r)
            i += 2
            continue
        i += 1


def add_rework_example_table(doc: Document) -> None:
    table = doc.add_table(rows=4, cols=4)
    headers = ("接口", "状态", "测试条件", "测试结果")
    values = (
        ("填写接口", "改焊后", "填写改焊后条件", "填写改焊后实测值"),
        ("", "改焊前", "填写改焊前条件", "填写改焊前实测值"),
        ("填写接口", "—", "填写未复测接口条件", "填写实测值"),
    )
    for cell, value in zip(table.rows[0].cells, headers):
        cell.text = value
    for row, value_row in zip(table.rows[1:], values):
        for cell, value in zip(row.cells, value_row):
            cell.text = value
    _set_table_widths(table, [1560, 1000, 2800, 4000])
    format_rework_table(table)


def _set_table_widths(table, widths: list[int]) -> None:
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        _set_attr(grid_col, "w:w", str(width))
        grid.append(grid_col)
    tblw = table._tbl.tblPr.find(qn("w:tblW"))
    if tblw is None:
        tblw = OxmlElement("w:tblW")
        table._tbl.tblPr.insert(0, tblw)
    _set_attr(tblw, "w:w", str(sum(widths)))
    _set_attr(tblw, "w:type", "dxa")
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            tcw = cell._tc.get_or_add_tcPr().get_or_add_tcW()
            _set_attr(tcw, "w:w", str(width))
            _set_attr(tcw, "w:type", "dxa")


def add_result_table(doc: Document) -> None:
    table = doc.add_table(rows=2, cols=4)
    headers = ("通道", "指标要求", "测试结果", "结论")
    values = ("填写通道", "填写指标要求", "填写测试结果", "填写结论")
    for cell, value in zip(table.rows[0].cells, headers):
        cell.text = value
    for cell, value in zip(table.rows[1].cells, values):
        cell.text = value
    _set_table_widths(table, [1560, 2700, 3300, 1800])
    format_table(table)


def add_pending_table(doc: Document) -> None:
    table = doc.add_table(rows=2, cols=5)
    headers = ("序号", "单板/细则来源", "指标/接口", "细则要求", "状态")
    values = ("1", "填写来源", "填写未完成测试项目", "填写指标要求", "未测试")
    for cell, value in zip(table.rows[0].cells, headers):
        cell.text = value
    for cell, value in zip(table.rows[1].cells, values):
        cell.text = value
    _set_table_widths(table, [720, 1800, 2160, 3600, 1080])
    format_table(table)


def add_seq_figure_caption(paragraph, label: str, number: int = 1) -> None:
    """Compatibility name: current compact profile uses a visible plain-text number."""
    _set_paragraph_zero_center(paragraph, "图表题注")
    run = paragraph.add_run(f"图{number} {label}")
    _set_run_direct(run, SPECIAL_STYLES["图表题注"])


def create_template(output: Path) -> None:
    doc = Document()
    configure_page(doc)
    configure_styles(doc)
    bind_heading_numbering(doc)

    title = doc.add_paragraph(style=TITLE_STYLE["name"])
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("测试结果说明")

    doc.add_paragraph("测试对象：填写测试对象；更新日期：填写日期。")
    doc.add_paragraph("器件名称", style=HEADING_STYLES[1]["name"])
    doc.add_paragraph("测试项目名称", style=HEADING_STYLES[2]["name"])
    doc.add_paragraph("测试/分析方法：填写实际测试条件和分析方法。")
    doc.add_paragraph("指标要求：填写指标要求；缺少要求时写未提供。")
    add_result_table(doc)

    doc.add_paragraph("结果说明：填写实测值、判定及未覆盖内容。")
    image = doc.add_paragraph("［插入结果图片］", style="图片")
    _set_paragraph_zero_center(image, "图片")
    add_seq_figure_caption(doc.add_paragraph(), "测试结果图")

    doc.add_paragraph("测试项目名称（改焊对比示例）", style=HEADING_STYLES[2]["name"])
    doc.add_paragraph("测试/分析方法：填写改焊前后同一参考面、同一条件的测试方法。")
    doc.add_paragraph("指标要求：填写指标要求；缺少要求时写未提供。")
    add_rework_example_table(doc)
    doc.add_paragraph(
        "结果说明：改焊后行在前（浅蓝）、改焊前行在后（浅红灰字）；"
        "改焊前后对比须同基准；未复测接口单行标改焊前或—。"
    )

    appendix = doc.add_paragraph("附录：数据溯源表", style=HEADING_STYLES[1]["name"])
    num = OxmlElement("w:numPr")
    numid = OxmlElement("w:numId")
    numid.set(qn("w:val"), "0")
    num.append(numid)
    appendix._p.get_or_add_pPr().append(num)
    doc.add_paragraph("路径相对所选主库根；库外证据保留已核验绝对路径并说明原因。").paragraph_format.keep_with_next = True
    trace = doc.add_table(rows=2, cols=7)
    for cell, text in zip(trace.rows[0].cells, ("器件","指标","接口","改焊状态","源数据路径","分析结果路径","汇总文件/备注")):
        cell.text = text
    for cell in trace.rows[1].cells:
        cell.text = "待填"
    _set_table_widths(trace, [800,1100,900,900,2000,2000,1660])
    format_table(trace)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()
    create_template(args.output.resolve())
    print(args.output.resolve())


if __name__ == "__main__":
    main()

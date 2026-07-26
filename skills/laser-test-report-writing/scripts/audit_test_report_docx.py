#!/usr/bin/env python3
"""Audit a generated laser electrical-test DOCX report."""

from __future__ import annotations

import argparse
import json
import re
import sys
import zipfile
from dataclasses import asdict, dataclass
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph


DEFAULT_PROBLEM_TERMS = [
    "本报告汇总",
    "正文保留",
    "定量观察：",
    "P0 阻塞项",
    "P0阻塞项",
    "具体流程同",
    "本次记录满足",
    "暂不能验收",
    "理论满足",
    "裸芯片本底",
    "无法判断是否是否",
    "JG12 满足 ASD 满足指标",
    "JG12满足ASD满足指标",
    "本流程用于",
    "本节列出",
    "下表用于",
    "以下三张图用于",
    "每幅图仅",
    "图像不在其他章节重复",
    "结果读法",
    "证据链",
    "正式判定",
    "可形成正式通过结论",
    "历史复算",
    "当前方案",
    "分析边界",
    "PASS",
    "FAIL",
]
SOFT_LANGUAGE_TERMS = [
    "当前证据",
    "现有证据",
    "证据不足",
    "证据边界",
    "证据覆盖",
    "结果分析：",
    "当前测试状态",
    "核心结果",
    "判定与问题",
    "证据状态",
    "值得注意的是",
    "可以清楚地看到",
    "不难发现",
    "显而易见",
    "综合来看",
    "综上所述",
]

CAPTION_RE = re.compile(r"^\s*(图|表)\s*\d+\s*.*$")
CAPTION_CROSS_REFERENCE_RE = re.compile(r"^\s*(图|表)\s*\d+\s*(?:至|~|～|-|—)\s*(图|表)\s*\d+")
FULLWIDTH_PAREN_RE = re.compile(r"（([^（）]{1,200})）")
LATIN_OR_GREEK_RE = re.compile(r"[A-Za-z_]|[\u0370-\u03ff]|[µΩ℃]")
CJK_RE = re.compile(r"[\u3400-\u9fff]")
FORMULA_RE = re.compile(r"[0-9]\s*[+\-−×÷/*=<>≤≥^~～√]|[+\-−×÷/*=<>≤≥^~～√]\s*[0-9]")
ABSOLUTE_PATH_RE = re.compile(r"(?i)(?:[A-Z]:\\|//tsclient/|\\\\tsclient\\)")
ALLOWED_FORMAL_STATES = {"满足", "不满足", "暂不能判定", "未测试"}
LEGACY_FORMAL_STATE_TOKENS = {
    "通过",
    "未通过",
    "部分通过",
    "待补证",
    "待补测",
    "正式通过",
    "数值通过",
    "整项待补测",
    "未开展",
    "理论满足",
    "本次记录满足",
    "数值满足",
    "略低于指标要求",
    "暂不能验收",
    "初步判定",
    "满足指标需求",
}

VALUE_UNIT_NO_SPACE_RE = re.compile(
    r"(?<![A-Za-z0-9_])"
    r"(?P<value>[+−-]?(?:\d+(?:\.\d+)?|\.\d+))"
    r"(?P<unit>GHz|MHz|kHz|Hz|GSPS|MSPS|kSPS|SPS|ns|µs|ms|s|"
    r"dBc/Hz|dBm|dBc|dB|µVpp/codepp|Vpp/codepp|µV/√Hz|nV/√Hz|V/√Hz|"
    r"µVrms|mVrms|Vrms|µVpp|mVpp|Vpp|µV|mV|V|Ω|bit)"
    r"(?![A-Za-z0-9_./\\])"
)
ASCII_MINUS_WITH_UNIT_RE = re.compile(
    r"(?<![A-Za-z0-9_./\\])-(?P<value>\d+(?:\.\d+)?)\s+"
    r"(?P<unit>GHz|MHz|kHz|Hz|dBc/Hz|dBm|dBc|dB|µV|mV|V|Ω)"
)
AT_FORMAT_RE = re.compile(r"\b(?:ASD|PSD|ADEV|MDEV)@\s*\d", re.IGNORECASE)
LEGACY_UV_RE = re.compile(r"(?<![A-Za-z0-9_])\d+(?:\.\d+)?uV(?![A-Za-z0-9_./\\])")
ASCII_SCIENTIFIC_RE = re.compile(r"(?<![A-Za-z0-9_])\d+(?:\.\d+)?[eE][+\-]?\d+(?![A-Za-z0-9_])")
DAMAGED_SUPERSCRIPT_RE = re.compile(
    r"10\s+(?:[⁰¹²³⁴⁵⁶⁷⁸⁹]+\s*[⁻⁺]|[⁻⁺]\s+[⁰¹²³⁴⁵⁶⁷⁸⁹]+|[⁰¹²³⁴⁵⁶⁷⁸⁹]*[⁻⁺]\s+[⁰¹²³⁴⁵⁶⁷⁸⁹]+)"
)


@dataclass
class AuditResult:
    path: str
    mode: str
    paragraph_count: int
    table_count: int
    image_count: int
    office_math_count: int
    office_math_structured_count: int
    problem_hits: list[dict[str, str]]
    language_warnings: list[dict[str, str]]
    formal_status_issues: list[dict[str, str]]
    unit_format_issues: list[dict[str, str]]
    scientific_notation_issues: list[dict[str, str]]
    revision_artifact_issues: list[dict[str, str]]
    parenthesis_issues: list[dict[str, str]]
    table_indent_issues: list[dict[str, str]]
    table_alignment_issues: list[dict[str, str]]
    table_object_issues: list[dict[str, str]]
    table_pagination_issues: list[dict[str, str]]
    table_header_issues: list[dict[str, str]]
    caption_issues: list[dict[str, str]]
    structure_issues: list[dict[str, str]]
    engineering_profile_issues: list[dict[str, str]]
    ad2208_da9726_profile_issues: list[dict[str, str]]
    dds_profile_issues: list[dict[str, str]]
    formula_structure_warnings: list[str]
    calibration_warnings: list[str]
    has_next_plan_heading: bool


def context(text: str, term: str, span: int = 35) -> str:
    idx = text.find(term)
    if idx < 0:
        return ""
    start = max(0, idx - span)
    end = min(len(text), idx + len(term) + span)
    return text[start:end].replace("\n", " ")


def collect_problem_hits(doc: Document, terms: list[str]) -> list[dict[str, str]]:
    hits: list[dict[str, str]] = []
    blocks: list[tuple[str, str]] = []
    for i, para in enumerate(doc.paragraphs, 1):
        if para.text.strip():
            blocks.append((f"paragraph {i}", para.text))
    for ti, table in enumerate(doc.tables, 1):
        for ri, row in enumerate(table.rows, 1):
            for ci, cell in enumerate(row.cells, 1):
                text = "\n".join(p.text for p in cell.paragraphs)
                if text.strip():
                    blocks.append((f"table {ti} row {ri} col {ci}", text))

    for location, text in blocks:
        for term in terms:
            if term in text:
                hits.append({"term": term, "location": location, "context": context(text, term)})
    return hits


def text_blocks(doc: Document) -> list[tuple[str, str]]:
    blocks: list[tuple[str, str]] = []
    for i, para in enumerate(doc.paragraphs, 1):
        if para.text.strip():
            blocks.append((f"paragraph {i}", para.text))
    for ti, table in enumerate(doc.tables, 1):
        for ri, row in enumerate(table.rows, 1):
            for ci, cell in enumerate(row.cells, 1):
                text = "\n".join(p.text for p in cell.paragraphs)
                if text.strip():
                    blocks.append((f"table {ti} row {ri} col {ci}", text))
    return blocks


def unit_format_issues(doc: Document) -> list[dict[str, str]]:
    issues: list[dict[str, str]] = []
    for location, text in text_blocks(doc):
        for match in VALUE_UNIT_NO_SPACE_RE.finditer(text):
            issues.append(
                {
                    "location": location,
                    "issue": "missing space between value and unit",
                    "text": match.group(0),
                    "context": context(text, match.group(0)),
                }
            )
        for match in ASCII_MINUS_WITH_UNIT_RE.finditer(text):
            issues.append(
                {
                    "location": location,
                    "issue": "negative quantity uses ASCII hyphen instead of mathematical minus",
                    "text": match.group(0),
                    "context": context(text, match.group(0)),
                }
            )
        for match in AT_FORMAT_RE.finditer(text):
            issues.append(
                {
                    "location": location,
                    "issue": "ASD/PSD/ADEV/MDEV uses @ without surrounding spacing",
                    "text": match.group(0),
                    "context": context(text, match.group(0)),
                }
            )
        for match in LEGACY_UV_RE.finditer(text):
            issues.append(
                {
                    "location": location,
                    "issue": "legacy ASCII microvolt unit",
                    "text": match.group(0),
                    "context": context(text, match.group(0)),
                }
            )
        for old_unit in ("µV/code", "欧 load", "欧load"):
            if old_unit in text:
                issues.append(
                    {
                        "location": location,
                        "issue": "legacy or ambiguous unit expression",
                        "text": old_unit,
                        "context": context(text, old_unit),
                    }
                )
    return issues


def scientific_notation_issues(doc: Document) -> list[dict[str, str]]:
    issues: list[dict[str, str]] = []
    for location, text in text_blocks(doc):
        for pattern, issue in (
            (DAMAGED_SUPERSCRIPT_RE, "damaged or reordered superscript exponent"),
            (ASCII_SCIENTIFIC_RE, "ASCII e-notation remains in formal report text"),
        ):
            for match in pattern.finditer(text):
                issues.append(
                    {
                        "location": location,
                        "issue": issue,
                        "text": match.group(0),
                        "context": context(text, match.group(0)),
                    }
                )
    return issues


def revision_artifact_issues(path: Path) -> list[dict[str, str]]:
    issues: list[dict[str, str]] = []
    with zipfile.ZipFile(path) as archive:
        names = set(archive.namelist())
        if "word/comments.xml" in names:
            issues.append({"location": "DOCX package", "issue": "comments part remains", "text": "word/comments.xml"})
        if "word/commentsExtended.xml" in names:
            issues.append({"location": "DOCX package", "issue": "extended comments part remains", "text": "word/commentsExtended.xml"})

        xml_names = [name for name in names if name.startswith("word/") and name.endswith(".xml")]
        xml_text = "\n".join(archive.read(name).decode("utf-8", errors="ignore") for name in xml_names)
        checks = (
            (r"<w:ins(?:\s|>)", "tracked insertions remain"),
            (r"<w:del(?:\s|>)", "tracked deletions remain"),
            (r"<w:moveFrom(?:\s|>)", "tracked move-from records remain"),
            (r"<w:moveTo(?:\s|>)", "tracked move-to records remain"),
            (r"<w:trackRevisions(?:\s|/|>)", "track-revisions setting is enabled"),
            (r"<w:commentRangeStart(?:\s|>)|<w:commentReference(?:\s|>)", "comment anchors remain"),
            (r"<w:highlight\s+w:val=\"(?!none\")", "text highlighting remains"),
        )
        for pattern, issue in checks:
            count = len(re.findall(pattern, xml_text))
            if count:
                issues.append({"location": "DOCX XML", "issue": issue, "text": str(count)})
    return issues


def is_technical_parenthetical(inner: str) -> bool:
    """Return True when Chinese full-width parentheses enclose technical/Latin content."""
    if LATIN_OR_GREEK_RE.search(inner):
        return True
    return not CJK_RE.search(inner) and bool(FORMULA_RE.search(inner))


def collect_parenthesis_issues(doc: Document) -> list[dict[str, str]]:
    issues: list[dict[str, str]] = []
    blocks: list[tuple[str, str]] = []
    for i, para in enumerate(doc.paragraphs, 1):
        if para.text.strip():
            blocks.append((f"paragraph {i}", para.text))
    for ti, table in enumerate(doc.tables, 1):
        for ri, row in enumerate(table.rows, 1):
            for ci, cell in enumerate(row.cells, 1):
                text = "\n".join(p.text for p in cell.paragraphs)
                if text.strip():
                    blocks.append((f"table {ti} row {ri} col {ci}", text))

    for location, text in blocks:
        for match in FULLWIDTH_PAREN_RE.finditer(text):
            if not is_technical_parenthetical(match.group(1)):
                continue
            issues.append(
                {
                    "location": location,
                    "issue": "technical/Latin content uses full-width Chinese parentheses",
                    "text": match.group(0),
                    "context": context(text, match.group(0)),
                }
            )
    return issues


def effective_paragraph_value(paragraph: Paragraph, attribute: str):
    value = getattr(paragraph.paragraph_format, attribute)
    if value is not None:
        return value
    style = paragraph.style
    while style is not None:
        value = getattr(style.paragraph_format, attribute)
        if value is not None:
            return value
        style = style.base_style
    return None


def table_indent_issues(doc: Document) -> list[dict[str, str]]:
    issues: list[dict[str, str]] = []
    for ti, table in enumerate(doc.tables, 1):
        for ri, row in enumerate(table.rows, 1):
            for ci, cell in enumerate(row.cells, 1):
                for pi, para in enumerate(cell.paragraphs, 1):
                    if not para.text.strip() and not para.runs:
                        continue
                    first = effective_paragraph_value(para, "first_line_indent")
                    left = effective_paragraph_value(para, "left_indent")
                    right = effective_paragraph_value(para, "right_indent")
                    first_pt = None if first is None else first.pt
                    left_pt = None if left is None else left.pt
                    right_pt = None if right is None else right.pt
                    if first_pt != 0 or left_pt != 0 or right_pt != 0:
                        issues.append(
                            {
                                "location": f"table {ti} row {ri} col {ci} paragraph {pi}",
                                "firstLine": str(first_pt),
                                "left": str(left_pt),
                                "right": str(right_pt),
                                "text": para.text[:80],
                            }
                        )
    return issues


def table_alignment_issues(doc: Document) -> list[dict[str, str]]:
    issues: list[dict[str, str]] = []
    for ti, table in enumerate(doc.tables, 1):
        for ri, row in enumerate(table.rows, 1):
            for ci, cell in enumerate(row.cells, 1):
                vertical_ok = cell.vertical_alignment == WD_ALIGN_VERTICAL.CENTER
                for pi, para in enumerate(cell.paragraphs, 1):
                    if not para.text.strip() and not para.runs:
                        continue
                    horizontal = effective_paragraph_value(para, "alignment")
                    horizontal_ok = horizontal == WD_ALIGN_PARAGRAPH.CENTER
                    if not horizontal_ok or not vertical_ok:
                        issues.append(
                            {
                                "location": f"table {ti} row {ri} col {ci} paragraph {pi}",
                                "horizontal": str(horizontal),
                                "vertical": str(cell.vertical_alignment),
                                "text": para.text[:80],
                            }
                        )
    return issues


def table_object_issues(doc: Document) -> list[dict[str, str]]:
    issues: list[dict[str, str]] = []
    for ti, table in enumerate(doc.tables, 1):
        if table.alignment != WD_TABLE_ALIGNMENT.CENTER:
            issues.append({"location": f"table {ti}", "issue": "table object is not centered"})
        indent = table._tbl.tblPr.find(qn("w:tblInd"))
        if indent is not None and indent.get(qn("w:w"), "0") != "0":
            issues.append(
                {
                    "location": f"table {ti}",
                    "issue": f"table object has nonzero indent {indent.get(qn('w:w'))}",
                }
            )
    return issues


def table_pagination_issues(doc: Document) -> list[dict[str, str]]:
    issues: list[dict[str, str]] = []
    for ti, table in enumerate(doc.tables, 1):
        if not table.rows:
            continue
        header_pr = table.rows[0]._tr.get_or_add_trPr()
        if header_pr.find(qn("w:tblHeader")) is None:
            issues.append({"location": f"table {ti} row 1", "issue": "header row is not repeated"})
        for ri, row in enumerate(table.rows, 1):
            row_pr = row._tr.get_or_add_trPr()
            if row_pr.find(qn("w:cantSplit")) is None:
                issues.append({"location": f"table {ti} row {ri}", "issue": "row can split across pages"})
    return issues


def table_header_issues(doc: Document) -> list[dict[str, str]]:
    issues: list[dict[str, str]] = []
    for ti, table in enumerate(doc.tables, 1):
        if not table.rows:
            continue
        for ci, cell in enumerate(table.rows[0].cells, 1):
            text = cell.text.strip()
            if text.endswith(("。", ".", "；", ";", "：", ":")):
                issues.append(
                    {
                        "location": f"table {ti} row 1 col {ci}",
                        "issue": "table header ends with punctuation",
                        "text": text,
                    }
                )
    return issues


def body_blocks(doc: Document) -> list[tuple[str, Paragraph | Table]]:
    blocks: list[tuple[str, Paragraph | Table]] = []
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            paragraph = Paragraph(child, doc)
            has_drawing = bool(child.xpath(".//w:drawing | .//w:pict"))
            if paragraph.text.strip() or has_drawing:
                blocks.append(("paragraph", paragraph))
        elif child.tag == qn("w:tbl"):
            blocks.append(("table", Table(child, doc)))
    return blocks


def block_has_drawing(kind: str, block: Paragraph | Table) -> bool:
    element = block._p if kind == "paragraph" else block._tbl
    return bool(element.xpath(".//w:drawing | .//w:pict"))


def caption_issues(doc: Document) -> list[dict[str, str]]:
    issues: list[dict[str, str]] = []
    blocks = body_blocks(doc)
    for index, (kind, block) in enumerate(blocks):
        if kind != "paragraph":
            continue
        text = block.text.strip()
        if CAPTION_CROSS_REFERENCE_RE.match(text):
            continue
        match = CAPTION_RE.match(text)
        if not match:
            continue
        label = match.group(1)
        location = f"body block {index + 1}"
        if text.endswith(("。", ".")):
            issues.append({"location": location, "issue": f"{label} caption ends with a period", "text": text})
        instruction = "".join(node.text or "" for node in block._p.xpath(".//w:instrText"))
        if not re.search(fr"\bSEQ\s*{label}\b", instruction):
            issues.append({"location": location, "issue": f"{label} caption lacks an automatic SEQ field", "text": text})
        if label == "图":
            previous_is_template_prompt = (
                index > 0
                and blocks[index - 1][0] == "paragraph"
                and blocks[index - 1][1].style.name == "模板提示"
                and any(token in blocks[index - 1][1].text for token in ("在此插入", "待补充", "占位"))
            )
            is_explicit_placeholder = (
                "待补充" in text and any(token in text for token in ("位置", "占位"))
            ) or previous_is_template_prompt
            if not is_explicit_placeholder and (index == 0 or not block_has_drawing(*blocks[index - 1])):
                issues.append({"location": location, "issue": "figure caption is not immediately below a figure", "text": text})
        elif index + 1 >= len(blocks) or blocks[index + 1][0] != "table":
            issues.append({"location": location, "issue": "table caption is not immediately above a table", "text": text})
    return issues


def structure_issues(doc: Document) -> list[dict[str, str]]:
    issues: list[dict[str, str]] = []
    for index, paragraph in enumerate(doc.paragraphs, 1):
        text = paragraph.text.strip()
        if text in {"图像结果检查", "图像结果与谱形检查"}:
            issues.append(
                {
                    "location": f"paragraph {index}",
                    "issue": "standalone image-inspection chapter is not allowed",
                    "text": text,
                }
            )
    return issues


def normalized_text(text: str) -> str:
    return re.sub(r"\s+", "", text or "")


def heading_level(paragraph: Paragraph) -> int | None:
    name = paragraph.style.name if paragraph.style else ""
    lower = name.lower()
    if name in {"一级标题", "附录标题"} or lower in {"heading 1", "heading1"}:
        return 1
    if name in {"二级标题", "附录小标题"} or lower in {"heading 2", "heading2"}:
        return 2
    if name == "三级标题" or lower in {"heading 3", "heading3"}:
        return 3
    return None


def formal_status_issues(doc: Document) -> list[dict[str, str]]:
    issues: list[dict[str, str]] = []
    for ti, table in enumerate(doc.tables, 1):
        if not table.rows:
            continue
        headers = [normalized_text(cell.text) for cell in table.rows[0].cells]
        status_columns = [
            index
            for index, header in enumerate(headers)
            if header in {"结论", "判定", "正式状态"}
        ]
        for ci in status_columns:
            for ri, row in enumerate(table.rows[1:], 2):
                if ci >= len(row.cells):
                    continue
                value = normalized_text(row.cells[ci].text).strip("。.;；")
                if not value or value in ALLOWED_FORMAL_STATES:
                    continue
                has_no_independent_requirement = any(
                    "无独立指标" in normalized_text(cell.text) for cell in row.cells
                )
                if value == "仅作分析" and has_no_independent_requirement:
                    continue
                if value in LEGACY_FORMAL_STATE_TOKENS or any(
                    token in value for token in LEGACY_FORMAL_STATE_TOKENS
                ):
                    issues.append(
                        {
                            "location": f"table {ti} row {ri} col {ci + 1}",
                            "issue": "formal conclusion uses a legacy or mixed status",
                            "text": row.cells[ci].text.strip(),
                        }
                    )
                else:
                    issues.append(
                        {
                            "location": f"table {ti} row {ri} col {ci + 1}",
                            "issue": "formal conclusion is not one of 满足/不满足/暂不能判定/未测试",
                            "text": row.cells[ci].text.strip(),
                        }
                    )
    return issues


def engineering_profile_issues(doc: Document) -> list[dict[str, str]]:
    issues: list[dict[str, str]] = []
    paragraphs = list(doc.paragraphs)
    top = [
        (index, paragraph.text.strip())
        for index, paragraph in enumerate(paragraphs)
        if heading_level(paragraph) == 1 and paragraph.text.strip()
    ]
    if not top:
        return [{"location": "document", "issue": "no top-level engineering report headings found"}]

    first_text = top[0][1]
    if "指标需求" in first_text or "判定依据" in first_text:
        issues.append(
            {
                "location": f"paragraph {top[0][0] + 1}",
                "issue": "management-style requirement matrix is the first body chapter",
                "text": first_text,
            }
        )
    if "测试目的" not in first_text or not any(token in first_text for token in ("结论", "结果摘要")):
        issues.append(
            {
                "location": f"paragraph {top[0][0] + 1}",
                "issue": "first body chapter must combine test purpose and conclusion summary",
                "text": first_text,
            }
        )

    required_groups = [
        ("purpose and summary", lambda text: "测试目的" in text and any(token in text for token in ("结论", "结果摘要"))),
        (
            "object, connection, and conditions",
            lambda text: "测试对象" in text and "条件" in text and any(token in text for token in ("连接", "系统", "接口")),
        ),
        ("device-oriented results", lambda text: "测试结果" in text and "目的" not in text),
        ("anomalies and limitations", lambda text: "异常" in text and ("局限" in text or "限制" in text)),
        (
            "summary and follow-up",
            lambda text: ("总结" in text or "实测对比" in text) and any(token in text for token in ("后续", "下一步")),
        ),
        ("appendix", lambda text: text.startswith("附录")),
    ]
    positions: list[tuple[str, int]] = []
    for label, predicate in required_groups:
        match = next(((index, text) for index, text in top if predicate(text)), None)
        if match is None:
            issues.append({"location": "document", "issue": f"missing engineering section: {label}"})
        else:
            positions.append((label, match[0]))
    if any(positions[index][1] >= positions[index + 1][1] for index in range(len(positions) - 1)):
        issues.append(
            {
                "location": "document",
                "issue": "engineering top-level sections are out of order",
                "order": " -> ".join(label for label, _ in positions),
            }
        )

    for index, text in top:
        if text in {"测试方法与数据处理", "测试仪器、测试方法与数据处理", "测试仪器、测试方法与分析方法"}:
            issues.append(
                {
                    "location": f"paragraph {index + 1}",
                    "issue": "all item-specific methods are centralized instead of being colocated with results",
                    "text": text,
                }
            )

    summary_table_count = 0
    accepted_summary_headers = (
        {"对象", "测试项目", "指标", "测试条件", "实测结果", "结论", "备注"},
        {"对象", "已完成项目", "主要异常或未测试项", "总体结论", "关键限制"},
    )
    for table in doc.tables:
        if not table.rows:
            continue
        headers = {normalized_text(cell.text) for cell in table.rows[0].cells}
        if any(required.issubset(headers) for required in accepted_summary_headers):
            summary_table_count += 1
    if summary_table_count != 1:
        issues.append(
            {
                "location": "测试目的与结论摘要",
                "issue": "the report must contain exactly one approved core result-summary table",
                "found": str(summary_table_count),
            }
        )

    appendix_position = next((index for index, text in top if text.startswith("附录")), len(paragraphs))
    for index, paragraph in enumerate(paragraphs[:appendix_position], 1):
        text = paragraph.text.strip()
        if text and ABSOLUTE_PATH_RE.search(text):
            issues.append(
                {
                    "location": f"paragraph {index}",
                    "issue": "absolute local path appears in the main report instead of the data appendix",
                    "text": text,
                }
            )
    return issues


def ad2208_da9726_profile_issues(doc: Document) -> list[dict[str, str]]:
    issues: list[dict[str, str]] = []
    paragraphs = [paragraph.text.strip() for paragraph in doc.paragraphs if paragraph.text.strip()]
    table_text = [cell.text for table in doc.tables for row in table.rows for cell in row.cells]
    all_text = "\n".join(paragraphs + table_text)
    compact = normalized_text(all_text)
    title = paragraphs[0] if paragraphs else ""

    if "数字锁定版" in title:
        issues.append({"location": "title", "issue": "hardware report title uses 数字锁定版 instead of 数字锁定板", "text": title})
    if "_" in title:
        issues.append({"location": "title", "issue": "formal visible title contains underscores", "text": title})

    for index, text in enumerate(paragraphs, 1):
        compact_text = normalized_text(text)
        if "AD2208输入端等效噪声" in compact_text and "总链路" not in compact_text:
            issues.append(
                {
                    "location": f"paragraph {index}",
                    "issue": "equivalent-input noise is named as AD2208-only noise without total-chain qualification",
                    "text": text,
                }
            )

    if "总链路" in compact and "等效噪声" in compact:
        for label, tokens in [
            ("ADC calibration slope", ("kADC", "k_ADC", "kAD", "k_AD", "LADC", "L_ADC", "AD2208Vpp–Codepp")),
            ("DAC calibration slope", ("kDAC", "k_DAC", "kDA", "k_DA", "Kout", "K_out", "DA9726输出刻度")),
            (
                "FPGA gain or transfer",
                ("FPGA增益", "数字增益", "HFPGA", "H_FPGA", "G/HFPGA", "G/H_FPGA", "G=", "G＝"),
            ),
        ]:
            if not any(token in compact for token in tokens):
                issues.append({"location": "total-chain noise", "issue": f"missing {label}"})
        if "截距" in compact and not any(token in compact for token in ("不参与", "不计入", "去均值")):
            issues.append(
                {
                    "location": "total-chain noise",
                    "issue": "intercept is mentioned without stating that it does not participate in mean-removed PSD/ASD scaling",
                }
            )

    has_10_25 = any(token in compact for token in ("10～25MHz", "10-25MHz", "10—25MHz"))
    if has_10_25 and "10MSPS" in compact and not any(
        token in compact
        for token in (
            "不能用于评价10～25MHz",
            "不能用于10～25MHz",
            "不能评价10～25MHz",
            "不参与JG12/JG15的10～25MHz指标判定",
            "奈奎斯特",
            "Nyquist",
            "仅覆盖0～5MHz",
            "仅覆盖0～约5MHz",
        )
    ):
        issues.append(
            {
                "location": "AD2208 noise",
                "issue": "10 MSPS data are associated with a 10–25 MHz requirement without a Nyquist/coverage qualification",
            }
        )

    if "DC～30MHz" in compact and not any(
        token in compact for token in ("DC未测试", "DC未单独测试", "DC未单独设置测试点", "DC未单独设点")
    ):
        issues.append({"location": "AD2208 frequency response", "issue": "DC–30 MHz claim lacks an explicit DC-not-tested qualification"})

    if "Codepp" in all_text and "Codepk" in all_text:
        if not any(token in compact for token in ("近似", "初步", "快速检查")):
            issues.append({"location": "AD2208 frequency response", "issue": "Codepp≈2|Codepk| is not identified as an approximate or preliminary method"})
        if "正弦拟合" not in compact:
            issues.append({"location": "AD2208 frequency response", "issue": "screenshot-peak method lacks a formal sine-fit retest method"})

    if "JG18" in compact and "1Hz" in compact and "未测试" not in compact:
        for term in ("采集时长", "频率分辨率"):
            if term not in compact:
                issues.append({"location": "JG18 1 Hz noise", "issue": f"missing {term}"})

    if "−149.5dBc/Hz" in compact or "-149.5dBc/Hz" in compact:
        if "底噪" not in compact and "本底" not in compact and "测量下限" not in compact:
            issues.append({"location": "DA9726 phase noise", "issue": "identical far-offset result lacks analyzer-floor qualification"})

    if "DA9726" in compact and "0V" in compact and "±1V" not in compact:
        issues.append({"location": "DA9726 DC noise", "issue": "0 V result is present but ±1 V operating conditions are not separated"})

    if any(token in compact for token in ("−3～+3V", "-3～+3V", "−3~+3V")):
        if "200Ω" not in compact:
            issues.append({"location": "DA9726 output range", "issue": "output-range discussion lacks the 200 Ω load condition"})
        if "15mA" not in compact:
            issues.append({"location": "DA9726 output range", "issue": "output-range discussion lacks the 15 mA drive condition"})

    if all(token in compact for token in ("100000000", "10MSPS", "1s")) and "冲突" not in compact:
        issues.append({"location": "DA9726 DC noise", "issue": "100000000 points at 10 MSPS are reported as 1 s without reconciling duration"})

    if "5.25dBm" in compact and "5.35dBm" in compact and "冲突" not in compact:
        issues.append({"location": "JG12 input power", "issue": "5.25 dBm and 5.35 dBm coexist without conflict handling"})
    if "1.2Vpeak" in compact and "1.2Vpp" in compact and "冲突" not in compact:
        issues.append({"location": "JG18 input amplitude", "issue": "1.2 Vpeak and 1.2 Vpp coexist without conflict handling"})
    if "5.45dBm" in compact and not any(token in compact for token in ("不采用", "不能确认", "待确认", "未确认")):
        issues.append({"location": "JG15 input power", "issue": "5.45 dBm is used without an explicit verification limitation"})
    return issues


def is_heading(paragraph: Paragraph) -> bool:
    return heading_level(paragraph) is not None


def dds_profile_issues(doc: Document) -> list[dict[str, str]]:
    issues: list[dict[str, str]] = []
    paragraphs = list(doc.paragraphs)
    device_sections: list[tuple[str, int]] = []
    for index, paragraph in enumerate(paragraphs):
        if not is_heading(paragraph):
            continue
        match = re.fullmatch(r"(DDS\d+).*测试结果", paragraph.text.strip())
        if match:
            device_sections.append((match.group(1), index))
    if not device_sections:
        issues.append({"location": "document", "issue": "no DDS device result chapter found"})

    for section_index, (device, start) in enumerate(device_sections):
        end = device_sections[section_index + 1][1] if section_index + 1 < len(device_sections) else len(paragraphs)
        section_text = normalized_text("\n".join(paragraph.text for paragraph in paragraphs[start:end]))
        categories = [
            ("control word", ("控制字理论", "理论控制字", "理论频率步进", "理论步进", "数字设定分辨率", "理论设定分辨率", "FTW")),
            ("power", ("输出功率",)),
            ("stability", ("秒稳定度", "频点秒稳定度", "稳定度结果", "频率稳定度", "数据段稳定度", "ADEV")),
            ("distinguishability", ("实测记录比较", "实测可分辨性", "实测分辨", "频率分辨", "微小频差", "置信区间与可分辨性")),
        ]
        for name, tokens in categories:
            if not any(normalized_text(token) in section_text for token in tokens):
                issues.append({"location": device, "issue": f"missing DDS result subsection: {name}"})

    all_text = normalized_text("\n".join(paragraph.text for paragraph in paragraphs))
    if "DDS2" in all_text and "5×10⁻¹²" in all_text:
        if not any(token in all_text for token in ("指标来源", "指标未确认", "正式指标")):
            issues.append(
                {
                    "location": "DDS2 stability",
                    "issue": "DDS2 uses a 5×10⁻¹² requirement without confirming its source",
                }
            )
    return issues


def office_math_metrics(path: Path) -> tuple[int, int, list[str]]:
    with zipfile.ZipFile(path) as z:
        xml = z.read("word/document.xml").decode("utf-8", errors="ignore")
    display_count = xml.count("<m:oMathPara")
    inline_count = len(re.findall(r"<m:oMath\b", xml))
    office_math_count = display_count if display_count else inline_count
    structure_tags = ("m:sSub", "m:sSup", "m:sSubSup", "m:f", "m:rad", "m:nary")
    structured_count = sum(len(re.findall(fr"<{tag}(?:\s|>)", xml)) for tag in structure_tags)
    warnings: list[str] = []
    if office_math_count and structured_count == 0:
        warnings.append("Office Math objects contain no professional subscript/fraction/radical structures")
    if "[[OMATH]]" in xml:
        warnings.append("Unconverted Word equation markers remain in document.xml")
    math_text = re.findall(r"<m:t[^>]*>(.*?)</m:t>", xml, flags=re.DOTALL)
    suspicious = [text for text in math_text if "_" in text or "sqrt" in text or "integral" in text]
    if suspicious:
        warnings.append(f"Linear formula text remains unbuilt: {suspicious[:3]}")
    fractions = re.findall(r"<m:f(?:\s[^>]*)?>.*?</m:f>", xml, flags=re.DOTALL)
    for index, fraction in enumerate(fractions, 1):
        for part_name in ("num", "den"):
            part = re.search(
                fr"<m:{part_name}(?:\s[^>]*)?>(.*?)</m:{part_name}>",
                fraction,
                flags=re.DOTALL,
            )
            if part is None:
                continue
            part_text = "".join(re.findall(r"<m:t[^>]*>(.*?)</m:t>", part.group(1), flags=re.DOTALL))
            if any(symbol in part_text for symbol in ("=", "≈", "≤", "≥")):
                warnings.append(
                    f"Fraction {index} {part_name} contains a relation sign; check fraction grouping: {part_text}"
                )
    return office_math_count, structured_count, warnings


def audit(path: Path, terms: list[str], profile: str | None = None, mode: str = "reconstruction") -> AuditResult:
    doc = Document(str(path))
    office_math_count, office_math_structured_count, formula_structure_warnings = office_math_metrics(path)
    all_text = "\n".join([para.text for para in doc.paragraphs] + ["\t".join(cell.text for cell in row.cells) for table in doc.tables for row in table.rows])
    calibration_warnings: list[str] = []
    if "AD 输入端等效噪声" in all_text:
        for required in ["K_out", "L_ADC"]:
            if required not in all_text:
                calibration_warnings.append(f"AD 输入端等效噪声 section missing {required}")
        if "H_FPGA" not in all_text and "G/H_FPGA" not in all_text and "FPGA 增益" not in all_text:
            calibration_warnings.append("AD 输入端等效噪声 section missing FPGA gain/transfer term")
    return AuditResult(
        path=str(path),
        mode=mode,
        paragraph_count=len(doc.paragraphs),
        table_count=len(doc.tables),
        image_count=len(doc.inline_shapes),
        office_math_count=office_math_count,
        office_math_structured_count=office_math_structured_count,
        problem_hits=collect_problem_hits(doc, terms),
        language_warnings=collect_problem_hits(doc, SOFT_LANGUAGE_TERMS),
        formal_status_issues=formal_status_issues(doc),
        unit_format_issues=unit_format_issues(doc),
        scientific_notation_issues=scientific_notation_issues(doc),
        revision_artifact_issues=revision_artifact_issues(path),
        parenthesis_issues=collect_parenthesis_issues(doc),
        table_indent_issues=table_indent_issues(doc),
        table_alignment_issues=table_alignment_issues(doc),
        table_object_issues=table_object_issues(doc),
        table_pagination_issues=table_pagination_issues(doc),
        table_header_issues=table_header_issues(doc),
        caption_issues=caption_issues(doc),
        structure_issues=structure_issues(doc),
        engineering_profile_issues=(
            engineering_profile_issues(doc)
            if mode != "surgical" and profile in {"engineering", "ad2208-da9726", "dds"}
            else []
        ),
        ad2208_da9726_profile_issues=(
            ad2208_da9726_profile_issues(doc) if profile == "ad2208-da9726" else []
        ),
        dds_profile_issues=dds_profile_issues(doc) if mode != "surgical" and profile == "dds" else [],
        formula_structure_warnings=formula_structure_warnings,
        calibration_warnings=calibration_warnings,
        has_next_plan_heading="下一步计划" in all_text or "后续测试" in all_text,
    )


def main() -> int:
    for stream in (sys.stdout, sys.stderr):
        reconfigure = getattr(stream, "reconfigure", None)
        if reconfigure is not None:
            try:
                reconfigure(encoding="utf-8", errors="replace")
            except (OSError, ValueError):
                pass
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("docx", help="DOCX report path")
    parser.add_argument("--json", action="store_true", help="Print JSON")
    parser.add_argument(
        "--terms",
        nargs="*",
        default=DEFAULT_PROBLEM_TERMS,
        help="Problem terms to scan",
    )
    parser.add_argument("--warn-only", action="store_true", help="Always exit 0")
    parser.add_argument(
        "--profile",
        choices=["engineering", "ad2208-da9726", "dds"],
        help="Enable profile-specific content and architecture checks",
    )
    parser.add_argument(
        "--mode",
        choices=["new", "reconstruction", "surgical"],
        default="reconstruction",
        help="Select strict new/reconstruction audit or format-preserving surgical audit",
    )
    args = parser.parse_args()

    path = Path(args.docx).expanduser().resolve()
    if not path.exists():
        raise SystemExit(f"DOCX not found: {path}")

    result = audit(path, args.terms, args.profile, args.mode)

    if args.json:
        print(json.dumps(asdict(result), ensure_ascii=False, indent=2))
    else:
        print(f"path: {result.path}")
        print(f"mode: {result.mode}")
        print(f"paragraphs: {result.paragraph_count}")
        print(f"tables: {result.table_count}")
        print(f"images: {result.image_count}")
        print(f"office_math: {result.office_math_count}")
        print(f"office_math_structured: {result.office_math_structured_count}")
        print(f"has_next_plan_heading: {result.has_next_plan_heading}")
        print(f"problem_hits: {len(result.problem_hits)}")
        for hit in result.problem_hits[:20]:
            print(f"  - {hit['term']} @ {hit['location']}: {hit['context']}")
        print(f"language_warnings: {len(result.language_warnings)}")
        for hit in result.language_warnings[:20]:
            print(f"  - {hit['term']} @ {hit['location']}: {hit['context']}")
        print(f"formal_status_issues: {len(result.formal_status_issues)}")
        for issue in result.formal_status_issues[:20]:
            print(f"  - {issue['location']}: {issue['issue']}: {issue.get('text', '')}")
        print(f"unit_format_issues: {len(result.unit_format_issues)}")
        for issue in result.unit_format_issues[:20]:
            print(f"  - {issue['location']}: {issue['issue']}: {issue.get('text', '')}")
        print(f"scientific_notation_issues: {len(result.scientific_notation_issues)}")
        for issue in result.scientific_notation_issues[:20]:
            print(f"  - {issue['location']}: {issue['issue']}: {issue.get('text', '')}")
        print(f"revision_artifact_issues: {len(result.revision_artifact_issues)}")
        for issue in result.revision_artifact_issues[:20]:
            print(f"  - {issue['location']}: {issue['issue']}: {issue.get('text', '')}")
        print(f"parenthesis_issues: {len(result.parenthesis_issues)}")
        for issue in result.parenthesis_issues[:20]:
            print(f"  - {issue['location']}: {issue['text']}: {issue['context']}")
        print(f"table_indent_issues: {len(result.table_indent_issues)}")
        for issue in result.table_indent_issues[:20]:
            print(
                "  - {location}: firstLine={firstLine}, left={left}, right={right}, text={text}".format(
                    **issue
                )
            )
        print(f"table_alignment_issues: {len(result.table_alignment_issues)}")
        for issue in result.table_alignment_issues[:20]:
            print(
                "  - {location}: horizontal={horizontal}, vertical={vertical}, text={text}".format(
                    **issue
                )
            )
        print(f"table_object_issues: {len(result.table_object_issues)}")
        for issue in result.table_object_issues[:20]:
            print(f"  - {issue['location']}: {issue['issue']}")
        print(f"table_pagination_issues: {len(result.table_pagination_issues)}")
        for issue in result.table_pagination_issues[:20]:
            print(f"  - {issue['location']}: {issue['issue']}")
        print(f"table_header_issues: {len(result.table_header_issues)}")
        for issue in result.table_header_issues[:20]:
            print(f"  - {issue['location']}: {issue['issue']}: {issue.get('text', '')}")
        print(f"caption_issues: {len(result.caption_issues)}")
        for issue in result.caption_issues[:20]:
            print(f"  - {issue['location']}: {issue['issue']}: {issue.get('text', '')}")
        print(f"structure_issues: {len(result.structure_issues)}")
        for issue in result.structure_issues[:20]:
            print(f"  - {issue['location']}: {issue['issue']}: {issue.get('text', '')}")
        print(f"engineering_profile_issues: {len(result.engineering_profile_issues)}")
        for issue in result.engineering_profile_issues[:30]:
            print(f"  - {issue['location']}: {issue['issue']}: {issue.get('text', '')}")
        print(f"ad2208_da9726_profile_issues: {len(result.ad2208_da9726_profile_issues)}")
        for issue in result.ad2208_da9726_profile_issues[:30]:
            print(f"  - {issue['location']}: {issue['issue']}: {issue.get('text', '')}")
        print(f"dds_profile_issues: {len(result.dds_profile_issues)}")
        for issue in result.dds_profile_issues[:20]:
            print(f"  - {issue['location']}: {issue['issue']}")
        print(f"formula_structure_warnings: {len(result.formula_structure_warnings)}")
        for warning in result.formula_structure_warnings:
            print(f"  - {warning}")
        print(f"calibration_warnings: {len(result.calibration_warnings)}")
        for warning in result.calibration_warnings:
            print(f"  - {warning}")

    common_hard_failures = bool(
        result.problem_hits
        or result.formal_status_issues
        or result.unit_format_issues
        or result.scientific_notation_issues
        or result.revision_artifact_issues
        or result.parenthesis_issues
        or result.table_header_issues
        or result.caption_issues
        or result.structure_issues
        or result.formula_structure_warnings
    )
    if result.mode == "surgical":
        failed = bool(common_hard_failures or result.ad2208_da9726_profile_issues)
    else:
        failed = bool(
            common_hard_failures
            or result.table_indent_issues
            or result.table_alignment_issues
            or result.table_object_issues
            or result.table_pagination_issues
            or result.engineering_profile_issues
            or result.ad2208_da9726_profile_issues
            or result.dds_profile_issues
        )
    return 0 if args.warn_only or not failed else 1


if __name__ == "__main__":
    raise SystemExit(main())

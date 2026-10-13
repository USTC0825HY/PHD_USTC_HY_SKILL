"""Source identity checks for new/changed figures; never infer identity from filenames."""
from __future__ import annotations
import hashlib
import io
import posixpath
import re
import zipfile
from pathlib import Path
from lxml import etree
from docx import Document

NS={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main','a':'http://schemas.openxmlformats.org/drawingml/2006/main','r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}

def figures(path):
    from audit_figure_mapping import image_caption_pairs
    with zipfile.ZipFile(path) as z:
        root=etree.fromstring(z.read('word/document.xml'))
        rels={x.get('Id'):x for x in etree.fromstring(z.read('word/_rels/document.xml.rels'))}
        doc=Document(path)
        styles={s.style_id for s in doc.styles if s.name in ('图表题注','Caption','题注')}
        _,caps,_=image_caption_pairs(root,styles)
        result={}
        for cap in caps:
            media=[]
            for rid in cap.get('relationship_ids',[]):
                rel=rels.get(rid)
                if rel is None or rel.get('TargetMode')=='External': raise ValueError('unresolved/external image relationship')
                name=posixpath.normpath(posixpath.join('word',rel.get('Target',''))).lstrip('/')
                media.append((name,z.read(name)))
            if cap['number'] in result: raise ValueError('duplicate caption number')
            result[cap['number']]=(cap['description'],media)
        return result

def identity_checks(docx,manifest,manifest_path,mode,before=None):
    issues=[]; notes=[]
    def fail(loc,msg): issues.append(dict(location=loc,issue=msg,text=''))
    try:
        actual=figures(docx); old=figures(before) if before else {}
    except Exception as exc: return [dict(location='document',issue='image relationship unreadable',text=str(exc))],notes
    if mode=='surgical' and not before:
        fail('document','strict surgical figure check requires --before'); return issues,notes
    doc=Document(docx)
    appendix=[]
    for t in doc.tables:
        if not t.rows: continue
        headers=[c.text.strip() for c in t.rows[0].cells]
        if headers[:6]==['器件','指标','接口','改焊状态','源数据路径','分析结果路径'] and len(headers)==7:
            appendix.extend([[c.text.strip() for c in r.cells] for r in t.rows[1:]])
    for row in manifest:
        loc='figure '+row['sequence']
        try: no=int(row['sequence'])
        except ValueError: continue
        pair=actual.get(no)
        if not pair: fail(loc,'no embedded image'); continue
        # Only an unchanged caption and identical embedded bytes may use legacy exemption.
        if mode=='surgical' and pair==old.get(no):
            notes.append(dict(location=loc,issue='unchanged legacy figure; identity not re-certified',text=''))
            continue
        media=pair[1]
        if len(media)!=1: fail(loc,'strict schema requires one embedded image per caption'); continue
        name,data=media[0]
        if not row.get('source_path') or not re.fullmatch('[0-9a-fA-F]{64}',row.get('source_sha256','')):
            fail(loc,'new/changed figure requires source_path and valid source_sha256'); continue
        src=Path(row['source_path']); src=src if src.is_absolute() else Path(manifest_path).parent/src
        try: raw=src.read_bytes()
        except OSError as exc: fail(loc,'source inaccessible: '+str(exc)); continue
        if hashlib.sha256(raw).hexdigest()!=row['source_sha256'].lower(): fail(loc,'source hash mismatch')
        if src.name!=row['source_filename']: fail(loc,'source_filename differs from source_path')
        if row.get('embedded_part') and row['embedded_part']!=name: fail(loc,'embedded_part mismatch')
        if row.get('embedded_sha256') and hashlib.sha256(data).hexdigest()!=row['embedded_sha256'].lower(): fail(loc,'embedded hash mismatch')
        try:
            from PIL import Image
            with Image.open(io.BytesIO(raw)) as a,Image.open(io.BytesIO(data)) as b:
                a.load(); b.load()
                if raw!=data and (a.size!=b.size or a.convert('RGBA').tobytes()!=b.convert('RGBA').tobytes()):
                    fail(loc,'pixel identity unverified (lossy re-encoding requires traceable approved derivative)')
        except Exception as exc: fail(loc,'image decode/identity unverified: '+str(exc))
        try:
            index=int(row.get('traceability_row',''))
            if index<1: raise ValueError()
            entry=appendix[index-1]
            if entry[0]!=row['object'] or entry[2]!=row['interface']: fail(loc,'appendix object/interface mismatch')
            if not entry[4] or not entry[5]: fail(loc,'appendix source/results paths missing')
            target=Path(entry[5]); target=target if target.is_absolute() else Path(manifest_path).parent/target
            if src.resolve()!=target.resolve() and not src.resolve().is_relative_to(target.resolve()): fail(loc,'figure source is outside declared analysis result path')
        except (ValueError,IndexError): fail(loc,'new/changed figure needs valid traceability_row in seven-column data appendix')
    return issues,notes

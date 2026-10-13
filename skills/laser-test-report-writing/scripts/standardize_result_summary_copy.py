#!/usr/bin/env python3
"""Create a format-standardized copy of an existing result-summary DOCX."""

from __future__ import annotations

import argparse
import hashlib
import json
import shutil
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

from create_result_summary_template import bind_heading_numbering, configure_page, configure_styles, format_table
from result_summary_style_contract import HEADING_STYLES, SPECIAL_STYLES


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def package_fingerprint(path: Path) -> dict:
    doc = Document(path)
    paragraphs = [p.text for p in doc.paragraphs]
    tables = [[[cell.text for cell in row.cells] for row in table.rows] for table in doc.tables]
    fields = []
    for paragraph in doc.paragraphs:
        fields.extend(node.text or "" for node in paragraph._p.xpath(".//w:instrText"))
    with zipfile.ZipFile(path) as archive:
        media = {
            name: hashlib.sha256(archive.read(name)).hexdigest()
            for name in archive.namelist()
            if name.startswith("word/media/")
        }
        rels = sorted(
            (rel.reltype, rel.target_ref)
            for rel in doc.part.rels.values()
            if "image" in rel.reltype
        )
    return {
        "paragraphs": paragraphs,
        "tables": tables,
        "fields": fields,
        "media": media,
        "image_relationships": rels,
        "inline_shapes": len(doc.inline_shapes),
    }


def _has_drawing(paragraph) -> bool:
    return bool(paragraph._p.xpath(".//w:drawing | .//w:pict"))


def _apply_paragraph_roles(doc: Document) -> None:
    heading_map = {
        "Heading 1": HEADING_STYLES[1]["name"],
        "Heading 2": HEADING_STYLES[2]["name"],
        "Heading 3": HEADING_STYLES[3]["name"],
        "标题 1": HEADING_STYLES[1]["name"],
        "标题 2": HEADING_STYLES[2]["name"],
        "标题 3": HEADING_STYLES[3]["name"],
    }
    for paragraph in doc.paragraphs:
        if paragraph.style.name in heading_map:
            paragraph.style = heading_map[paragraph.style.name]
        if _has_drawing(paragraph):
            paragraph.style = "图片"
        if paragraph.style.name in SPECIAL_STYLES:
            ppr = paragraph._p.get_or_add_pPr()
            ind = ppr.find(qn("w:ind"))
            if ind is not None:
                for key in ("w:firstLine", "w:firstLineChars", "w:left", "w:leftChars", "w:right", "w:rightChars"):
                    ind.set(qn(key), "0")


def standardize(source: Path, output: Path) -> dict:
    source = source.resolve()
    output = output.resolve()
    if source == output:
        raise ValueError("Input and output paths must differ.")
    if not source.is_file():
        raise FileNotFoundError(source)
    lock = source.with_name("~$" + source.name)
    if lock.exists():
        raise RuntimeError(f"Word lock file exists: {lock}")

    source_meta = {
        "sha256": sha256(source),
        "size": source.stat().st_size,
        "mtime_ns": source.stat().st_mtime_ns,
    }
    before = package_fingerprint(source)
    output.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source, output)
    try:
        doc = Document(output)
        configure_page(doc)
        configure_styles(doc)
        bind_heading_numbering(doc)
        _apply_paragraph_roles(doc)
        for table in doc.tables:
            format_table(table)
        doc.save(output)
        after = package_fingerprint(output)
        if before != after:
            raise RuntimeError("Non-format content changed; output rejected.")
    except Exception:
        output.unlink(missing_ok=True)
        raise
    return {
        "source": str(source),
        "output": str(output),
        "source_metadata": source_meta,
        "output_sha256": sha256(output),
        "images": len(after["media"]),
        "inline_shapes": after["inline_shapes"],
        "tables": len(after["tables"]),
        "content_preserved": True,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    print(json.dumps(standardize(args.source, args.output), ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()

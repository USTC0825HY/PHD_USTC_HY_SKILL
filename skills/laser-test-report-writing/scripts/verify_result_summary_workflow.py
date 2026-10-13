"""Exercise compact DOCX creation, append, and authorized edits on synthetic evidence."""
import argparse, copy, csv, json, io
from pathlib import Path
from lxml import etree
from PIL import Image, ImageDraw
from docx import Document
from docx.shared import Cm
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from create_result_summary_template import configure_page, configure_styles, bind_heading_numbering, format_table, format_rework_table, add_seq_figure_caption, _set_table_widths
from audit_result_summary_docx import audit
from audit_figure_mapping import audit_document
from docx_safe_edit import set_cell_text, surgical_diff, table_checks

def main(out):
    out.mkdir(parents=True,exist_ok=True)
    image=Image.new("RGB",(900,420),"white");draw=ImageDraw.Draw(image)
    draw.text((25,20),"SYNTHETIC LAYOUT TEST - NOT MEASURED",fill="black")
    draw.line([(60,350),(840,350)],fill="black",width=3);draw.line([(60,350),(60,70)],fill="black",width=3)
    draw.line([(60,310),(250,260),(500,190),(800,110)],fill="blue",width=4)
    image.save(out/"layout.png")
    doc=Document();configure_page(doc);configure_styles(doc);bind_heading_numbering(doc)
    doc.add_paragraph("技能排版验证样稿",style="报告题名")
    doc.add_paragraph("测试对象：合成排版样例。所有数值仅验证格式，不是实测数据。")
    doc.add_paragraph("示例器件",style="一级标题")
    manifests=[]
    def indicator(number):
        doc.add_paragraph("示例指标"+str(number),style="二级标题")
        doc.add_paragraph("测试/分析方法：用合成数据检查表格、改焊配对与图片的对应关系。")
        doc.add_paragraph("指标要求：未提供。")
        table=doc.add_table(rows=4,cols=4)
        values=[["接口","状态","测试条件","测试结果"],["TEST1","改焊后","合成条件","1.000 V"],["","改焊前","合成条件","0.900 V"],["TEST2","—","未复测","未记录"]]
        for row,items in zip(table.rows,values):
            for cell,value in zip(row.cells,items):cell.text=value
        _set_table_widths(table,[1560,1000,2700,4100]);format_rework_table(table)
        doc.add_paragraph("结果说明：本页仅验证显示。后行蓝底在前，前行红底灰字在后；不作器件判定。")
        doc.add_paragraph(style="图片").add_run().add_picture(str(out/"layout.png"),width=Cm(13.5))
        caption=f"TEST{number}：合成排版验证曲线"
        add_seq_figure_caption(doc.add_paragraph(),caption,number)
        manifests.append(dict(sequence=number,evidence_id=f"TEST-CH{number}-LAYOUT-01",object="TEST",interface=f"TEST{number}",content_type="排版",caption=caption,source_filename="layout.png",word_section=f"示例指标{number}"))
    def appendix(copy_doc):
        p=copy_doc.add_paragraph("附录：数据溯源表",style="一级标题")
        num=OxmlElement("w:numPr");nid=OxmlElement("w:numId");nid.set(qn("w:val"),"0");num.append(nid);p._p.get_or_add_pPr().append(num)
        copy_doc.add_paragraph("以下路径为合成样例；正式报告按选定主库相对路径或库外绝对路径填写。").paragraph_format.keep_with_next = True
        table=copy_doc.add_table(rows=2,cols=7)
        fields=["器件","指标","接口","改焊状态","源数据路径","分析结果路径","汇总文件/备注"]
        vals=["示例","排版","TEST1","待确认","synthetic/raw","synthetic/results","非实测"]
        for row,items in zip(table.rows,[fields,vals]):
            for c,v in zip(row.cells,items):c.text=v
        _set_table_widths(table,[800,1000,1000,1000,2000,2000,1500]);format_table(table)
    def save_check(name):
        buffer=io.BytesIO();doc.save(buffer);buffer.seek(0);clone=Document(buffer);appendix(clone)
        path=out/(name+".docx");clone.save(path)
        manifest=out/(name+".csv")
        with manifest.open("w",encoding="utf-8-sig",newline="") as stream:
            w=csv.DictWriter(stream,fieldnames=list(manifests[0]));w.writeheader();w.writerows(manifests)
        result=audit(path,"new",True)
        assert result["errors"]==0,result
        figures=audit_document(path,manifest)
        assert figures.passed,figures
        return path,manifest
    indicator(1);before,manifest=save_check("01_new")
    prefix=[etree.tostring(x) for x in doc.element.body if not x.tag.endswith("sectPr")]
    doc.add_page_break();indicator(2);save_check("02_append")
    assert [etree.tostring(x) for x in list(doc.element.body)[:len(prefix)]]==prefix
    revised=Document(before)
    table=revised.tables[0];before_cell=etree.tostring(table.cell(1,3)._tc)
    set_cell_text(table.cell(1,3),"1.001 V")
    after=out/"03_surgical.docx";revised.save(after)
    body=list(revised.element.body);index=body.index(table._tbl)+1
    scope={"before_body":[index],"after_body":[index],"allowed_parts":[]}
    assert not surgical_diff(before,after,scope)
    restored=copy.deepcopy(table._tbl)

    target=restored.tr_lst[1].tc_lst[3]
    target.getparent().replace(target,etree.fromstring(before_cell))
    assert etree.tostring(restored,method="c14n",exclusive=True)==etree.tostring(Document(before).tables[0]._tbl,method="c14n",exclusive=True),"untouched cells changed"
    assert audit_document(after,manifest).passed
    (out/"scope.json").write_text(json.dumps(scope,indent=2),encoding="utf-8")
    # Two-row header fixture with complete grid, repeat headers, percentage/dxa widths.
    double=Document(before);t=double.add_table(rows=3,cols=4)
    _set_table_widths(t,[2000,2000,2000,2000])
    for row,items in zip(t.rows,[["接口","隔离度","完整性","最差隔离度"],["接口","驱动/安静","完整性","最差隔离度"],["TEST1","合成配对","1/1","40 dB"]]):
        for c,v in zip(row.cells,items):c.text=v
    format_table(t)
    for col in (0,2,3):
        t.cell(1,col).text = ""
        t.cell(0,col).merge(t.cell(1,col))
    header=OxmlElement("w:tblHeader");t._tbl.tr_lst[1].get_or_add_trPr().append(header)
    for tc in [cell for row in t._tbl.tr_lst[:2] for cell in row.tc_lst]:
        shd=tc.get_or_add_tcPr().find(qn("w:shd"))
        shd.set(qn("w:fill"),"1F4E79")
        for color in tc.xpath(".//w:color"):color.set(qn("w:val"),"FFFFFF")
    # Add the demonstration before the appendix, preserving the final appendix placement.
    t._tbl.getparent().remove(t._tbl)
    appendix_index=next(i for i,n in enumerate(double.element.body) if "".join(x.text or "" for x in n.iter(qn("w:t")))=="附录：数据溯源表")
    double.element.body.insert(appendix_index,t._tbl)
    double.save(out/"04_double_header.docx")
    assert not table_checks(double),table_checks(double)
    result={"new":"passed","append_preserves_existing":"passed","pure_text_captions":"passed","required_appendix":"passed","surgical_untouched_cells_and_parts":"passed","double_header":"passed","rework_pair":"passed"}
    (out/"validation.json").write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding="utf-8")
    print(json.dumps(result))

if __name__=="__main__":
    parser=argparse.ArgumentParser();parser.add_argument("--output-dir",required=True,type=Path)
    main(parser.parse_args().output_dir)


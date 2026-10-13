from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from audit_result_summary_docx import audit
from create_result_summary_template import create_template
from standardize_result_summary_copy import package_fingerprint, standardize


def test_template_passes_and_placeholders_block_final(tmp_path):
    target = tmp_path / "template.docx"
    create_template(target)
    assert audit(target, "new")["errors"] == 0
    assert audit(target, "new", final_content=True)["errors"] > 0


def test_negative_contract_violations_are_detected(tmp_path):
    target = tmp_path / "bad.docx"
    create_template(target)
    doc = Document(target)
    doc.sections[0].page_width = Cm(21.0)
    heading = next(p for p in doc.paragraphs if p.style.name == "一级标题")
    heading.text = "1. 手工编号"
    color = doc.styles["一级标题"].element.rPr.find(qn("w:color"))
    color.set(qn("w:val"), "0000FF")
    tblpr = doc.tables[0]._tbl.tblPr
    ind = tblpr.find(qn("w:tblInd"))
    ind.set(qn("w:w"), "240")
    trpr = doc.tables[0].rows[0]._tr.trPr
    trpr.remove(trpr.find(qn("w:tblHeader")))
    run = doc.tables[0].cell(0, 0).paragraphs[0].runs[0]
    run._r.get_or_add_rPr().find(qn("w:color")).set(qn("w:val"), "000000")
    jump = doc.add_paragraph("跳级标题", style="三级标题")
    numpr = jump._p.get_or_add_pPr().find(qn("w:numPr"))
    if numpr is not None:
        jump._p.pPr.remove(numpr)
    doc.save(target)
    codes = {f["code"] for f in audit(target, "new")["findings"]}
    assert {"PAGE.GEOMETRY", "STYLE.COLOR", "HEADING.MANUAL_NUMBER", "HEADING.LEVEL_JUMP", "TABLE.INDENT", "TABLE.REPEAT_HEADER", "TABLE.HEADER_FONT"} <= codes


def test_standardizer_preserves_content(tmp_path):
    source = tmp_path / "source.docx"
    output = tmp_path / "output.docx"
    create_template(source)
    before = package_fingerprint(source)
    result = standardize(source, output)
    assert result["content_preserved"] is True
    assert package_fingerprint(output) == before

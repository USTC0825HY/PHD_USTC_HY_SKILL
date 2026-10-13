#!/usr/bin/env python3
"""Audit a DOCX against the ultra-stable laser electrical-test report standard."""

from __future__ import annotations

import argparse
import json
import re
import sys
import zipfile
from dataclasses import asdict, dataclass, field
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

from create_standard_template import document_body_width_twips
from style_contract import (
    BODY_STYLE,
    HEADING_STYLES,
    NUMBERING_LEVELS,
    PAGE_STYLE,
    REPORT_SUBTITLE_STYLE,
    REPORT_TITLE_STYLE,
    TABLE_STYLE,
)


STATUS_VALUES = ("满足", "不满足", "暂不能判定", "未测试")
METADATA_PATTERN = re.compile(r"^测试(时间|人员|地点|目的|内容)\s*[：:]")
MANUAL_HEADING_PATTERN = re.compile(r"^(\d+(?:\.\d+){0,2}|[一二三四五六七八九十]+、)\s*\S+")
MEASUREMENT_HEADING_PATTERN = re.compile(
    r"^\d+(?:\.\d+){0,2}\s*(?:s|ms|us|ns|Hz|kHz|MHz|GHz|V|mV|uV|dB|dBm|dBc)\b",
    re.IGNORECASE,
)
CAPTION_PATTERN = re.compile(r"^(图|表)\s*\d+(?:[-.]\d+)?\s*\S+")
UNIT_METRICS = ("ASD", "PSD", "噪声", "频率", "电压", "电流", "温度", "功率", "稳定度")
BODY_FONT_EAST_ASIA = "宋体"
BODY_FONT_LATIN = "Times New Roman"
TOC_TITLE_FONT_EAST_ASIA = "微软雅黑"
TOC_TITLE_FONT_LATIN = "Microsoft YaHei"
TOC_LEFT_INDENTS_PT = {1: 0.0, 2: 21.0, 3: 42.0}
TOC_RIGHT_TAB_TWIPS = 8787
FULLWIDTH_TECHNICAL_PARENS = re.compile(r"（[^）\n]*[A-Za-z][^）\n]*）")
FINAL_PLACEHOLDERS = ("待填写", "待补充", "在此插入", "右键更新域", "板卡或测试对象简称")
LANGUAGE_REVIEW_TERMS = (
    "本流程用于",
    "本节列出",
    "下表用于",
    "以下三张图用于",
    "每幅图仅",
    "图像不在其他章节重复",
    "结果读法",
    "证据状态",
    "证据链",
    "正式判定",
    "可形成正式通过结论",
    "历史复算",
    "当前方案",
    "分析边界",
)
MATH_STRUCTURE_TAGS = (
    "m:sSub",
    "m:sSup",
    "m:sSubSup",
    "m:f",
    "m:rad",
    "m:nary",
    "m:d",
    "m:func",
)
KNOWN_INLINE_LITERAL = re.compile(
    r"(?<![A-Za-z0-9_])(?:"
    r"f_(?:s|0|alias|-3dB)|T_(?:s|alias)|K_out|L_ADC|H_FPGA(?:\(f\))?|"
    r"G/H_FPGA\(f\)|A_v\(f\)|S_v\(f\)|V_rms"
    r")(?![A-Za-z0-9_])"
)


@dataclass
class Finding:
    severity: str
    check: str
    location: str
    message: str


@dataclass
class AuditResult:
    path: str
    paragraph_count: int
    table_count: int
    image_count: int
    office_math_count: int
    heading_levels: list[int] = field(default_factory=list)
    findings: list[Finding] = field(default_factory=list)

    @property
    def errors(self) -> int:
        return sum(item.severity == "error" for item in self.findings)

    @property
    def warnings(self) -> int:
        return sum(item.severity == "warning" for item in self.findings)


def add(result: AuditResult, severity: str, check: str, location: str, message: str) -> None:
    result.findings.append(Finding(severity, check, location, message))


def clean(text: str) -> str:
    return " ".join(text.split())


def close(value: float | None, expected: float, tolerance: float = 0.2) -> bool:
    return value is not None and abs(value - expected) <= tolerance


def length_pt(value) -> float | None:
    return None if value is None else value.pt


def style_color(style) -> str:
    current = style
    while current is not None:
        value = current.font.color.rgb
        if value is not None:
            return str(value).upper()
        current = current.base_style
    return "000000"


def direct_style_rfonts(style) -> dict[str, str | None]:
    r_pr = style.element.rPr
    r_fonts = None if r_pr is None else r_pr.rFonts
    names = ("ascii", "hAnsi", "eastAsia", "cs", "asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme")
    return {name: None if r_fonts is None else r_fonts.get(qn(f"w:{name}")) for name in names}


def effective_style_rfonts(style) -> dict[str, str | None]:
    names = ("ascii", "hAnsi", "eastAsia", "cs", "asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme")
    result: dict[str, str | None] = {name: None for name in names}
    current = style
    while current is not None:
        direct = direct_style_rfonts(current)
        for name in names:
            if result[name] is None and direct[name] is not None:
                result[name] = direct[name]
        current = current.base_style
    return result


def effective_style_value(style, getter, default=None):
    current = style
    while current is not None:
        value = getter(current)
        if value is not None:
            return value
        current = current.base_style
    return default


def effective_style_alignment(style):
    return effective_style_value(style, lambda item: item.paragraph_format.alignment, WD_ALIGN_PARAGRAPH.LEFT)


def effective_style_indent_pt(style, attribute: str) -> float:
    value = effective_style_value(style, lambda item: getattr(item.paragraph_format, attribute), None)
    return 0.0 if value is None else value.pt


def effective_style_size_pt(style) -> float | None:
    value = effective_style_value(style, lambda item: item.font.size, None)
    return None if value is None else value.pt


def effective_style_space_pt(style, attribute: str) -> float:
    value = effective_style_value(style, lambda item: getattr(item.paragraph_format, attribute), None)
    return 0.0 if value is None else value.pt


def document_default_size_pt(doc: Document) -> float | None:
    defaults = doc.styles.element.find(qn("w:docDefaults"))
    if defaults is None:
        return None
    size = defaults.find("./" + qn("w:rPrDefault") + "/" + qn("w:rPr") + "/" + qn("w:sz"))
    if size is None or size.get(qn("w:val")) is None:
        return None
    return int(size.get(qn("w:val"))) / 2


def effective_style_bold(style) -> bool:
    return bool(effective_style_value(style, lambda item: item.font.bold, False))


def effective_alignment(paragraph):
    return paragraph.alignment if paragraph.alignment is not None else effective_style_alignment(paragraph.style)


def contains_image(paragraph) -> bool:
    return "<w:drawing" in paragraph._p.xml or "<w:pict" in paragraph._p.xml


def paragraph_has_explicit_page_break(paragraph) -> bool:
    return bool(
        paragraph._p.xpath(".//w:br[@w:type='page']")
        or paragraph._p.xpath("./w:pPr/w:pageBreakBefore")
    )


def paragraph_has_next_page_section_break(paragraph) -> bool:
    sect_pr = paragraph._p.find("./" + qn("w:pPr") + "/" + qn("w:sectPr"))
    if sect_pr is None:
        return False
    section_type = sect_pr.find(qn("w:type"))
    return section_type is None or section_type.get(qn("w:val"), "nextPage") == "nextPage"


def style_zero_indent(style) -> bool:
    return all(
        close(effective_style_indent_pt(style, attribute), 0.0)
        for attribute in ("first_line_indent", "left_indent", "right_indent")
    )


def heading_level(paragraph) -> int | None:
    style_id = (paragraph.style.style_id or "").lower().replace(" ", "")
    style_name = paragraph.style.name.lower().replace(" ", "")
    for level in (1, 2, 3):
        if style_id == f"heading{level}" or style_name in {f"heading{level}", f"标题{level}"}:
            return level
    if "一级标题" in paragraph.style.name:
        return 1
    if "二级标题" in paragraph.style.name:
        return 2
    if "三级标题" in paragraph.style.name:
        return 3
    return None


def has_direct_numbering(paragraph) -> bool:
    p_pr = paragraph._p.pPr
    return p_pr is not None and p_pr.numPr is not None and p_pr.numPr.numId is not None


def direct_num_id(paragraph) -> str | None:
    p_pr = paragraph._p.pPr
    if p_pr is None or p_pr.numPr is None or p_pr.numPr.numId is None:
        return None
    return str(p_pr.numPr.numId.val)


def style_numbering(style) -> tuple[str | None, str | None]:
    p_pr = style.element.pPr
    if p_pr is None or p_pr.numPr is None or p_pr.numPr.numId is None:
        return None, None
    ilvl = "0" if p_pr.numPr.ilvl is None else str(p_pr.numPr.ilvl.val)
    return str(p_pr.numPr.numId.val), ilvl


def numbering_level_details(doc: Document, num_id: str, ilvl: str) -> dict[str, str | int | None]:
    numbering = doc.part.numbering_part.element
    num = next((item for item in numbering.findall(qn("w:num")) if item.get(qn("w:numId")) == num_id), None)
    if num is None:
        return {}
    abstract_ref = num.find(qn("w:abstractNumId"))
    if abstract_ref is None:
        return {}
    abstract_id = abstract_ref.get(qn("w:val"))
    abstract = next(
        (item for item in numbering.findall(qn("w:abstractNum")) if item.get(qn("w:abstractNumId")) == abstract_id),
        None,
    )
    if abstract is None:
        return {}
    level = next((item for item in abstract.findall(qn("w:lvl")) if item.get(qn("w:ilvl")) == ilvl), None)
    if level is None:
        return {}
    lvl_text = level.find(qn("w:lvlText"))
    suffix = level.find(qn("w:suff"))
    p_style = level.find(qn("w:pStyle"))
    p_level = level.find(qn("w:pPr"))
    indent = None if p_level is None else p_level.find(qn("w:ind"))
    left = None if indent is None else indent.get(qn("w:left"))
    hanging = None if indent is None else indent.get(qn("w:hanging"))
    return {
        "text": None if lvl_text is None else lvl_text.get(qn("w:val")),
        "suffix": None if suffix is None else suffix.get(qn("w:val")),
        "p_style": None if p_style is None else p_style.get(qn("w:val")),
        "left": int(left) if left is not None else None,
        "hanging": int(hanging) if hanging is not None else None,
    }


def has_seq_field(paragraph, label: str) -> bool:
    # Word can serialize fields either as one field instruction or as several
    # w:instrText runs after updating and saving the document.
    instruction = "".join(node.text or "" for node in paragraph._p.xpath(".//w:instrText"))
    return bool(re.search(rf"\bSEQ\s*{re.escape(label)}\b", instruction))


def has_zero_indent(paragraph) -> bool:
    fmt = paragraph.paragraph_format
    for attribute in ("first_line_indent", "left_indent", "right_indent"):
        direct = getattr(fmt, attribute)
        actual = effective_style_indent_pt(paragraph.style, attribute) if direct is None else direct.pt
        if not close(actual, 0.0):
            return False
    return True


def table_header_repeats(table) -> bool:
    tr_pr = table.rows[0]._tr.trPr
    return tr_pr is not None and tr_pr.find(qn("w:tblHeader")) is not None


def count_office_math(path: Path) -> int:
    with zipfile.ZipFile(path) as archive:
        xml = archive.read("word/document.xml").decode("utf-8", errors="ignore")
    display_count = xml.count("<m:oMathPara")
    inline_count = len(re.findall(r"<m:oMath(?:\s|>)", xml))
    return max(display_count, inline_count)


def get_style(doc: Document, result: AuditResult, name: str):
    try:
        return doc.styles[name]
    except KeyError:
        normalized = name.lower().replace(" ", "")
        for style in doc.styles:
            if style.name.lower().replace(" ", "") == normalized or style.style_id.lower().replace(" ", "") == normalized:
                return style
    add(result, "error", "required_style", "styles", f"Missing required style: {name}.")
    return None


def audit_style_font(
    result: AuditResult,
    style,
    name: str,
    east_asia: str,
    latin: str,
    size_pt: float,
    bold: bool,
    alignment,
    zero_indent: bool = True,
    reject_theme: bool = False,
    default_size_pt: float | None = None,
) -> None:
    fonts = effective_style_rfonts(style)
    accepted_latin = {latin}
    if latin == TOC_TITLE_FONT_LATIN:
        accepted_latin.add(TOC_TITLE_FONT_EAST_ASIA)
    if fonts["eastAsia"] != east_asia or fonts["ascii"] not in accepted_latin or fonts["hAnsi"] not in accepted_latin:
        add(
            result,
            "error",
            "style_font",
            f"style {name}",
            f"Expected {east_asia}/{latin}; found eastAsia={fonts['eastAsia']}, ascii={fonts['ascii']}, hAnsi={fonts['hAnsi']}.",
        )
    direct_fonts = direct_style_rfonts(style)
    if reject_theme and any(direct_fonts[key] for key in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme")):
        direct_latin = {direct_fonts["ascii"], direct_fonts["hAnsi"]}
        if direct_fonts["eastAsia"] != east_asia or not direct_latin.issubset(accepted_latin):
            add(
                result,
                "error",
                "style_font_theme",
                f"style {name}",
                "Theme-font attributes remain without complete explicit font overrides.",
            )
    actual_size = effective_style_size_pt(style)
    if actual_size is None:
        actual_size = default_size_pt
    if not close(actual_size, size_pt):
        add(result, "error", "style_size", f"style {name}", f"Expected {size_pt} pt; found {actual_size}.")
    actual_bold = effective_style_bold(style)
    if actual_bold != bold:
        add(result, "error", "style_bold", f"style {name}", f"Expected bold={bold}; found {actual_bold}.")
    actual_alignment = effective_style_alignment(style)
    if actual_alignment != alignment:
        add(
            result,
            "error",
            "style_alignment",
            f"style {name}",
            f"Expected alignment={alignment}; found {actual_alignment}.",
        )
    if zero_indent and not style_zero_indent(style):
        add(result, "error", "style_indent", f"style {name}", "Style must have zero first-line, left, and right indent.")


def audit_style_contract(doc: Document, result: AuditResult) -> None:
    default_size = document_default_size_pt(doc)
    normal = get_style(doc, result, "Normal")
    if normal is not None:
        audit_style_font(
            result,
            normal,
            "Normal",
            BODY_FONT_EAST_ASIA,
            BODY_FONT_LATIN,
            12,
            False,
            WD_ALIGN_PARAGRAPH.JUSTIFY,
            zero_indent=False,
            default_size_pt=default_size,
        )
        if not close(effective_style_indent_pt(normal, "first_line_indent"), 24):
            add(result, "error", "body_indent", "style Normal", "Body first-line indent must be 24 pt.")
        if not close(length_pt(normal.paragraph_format.line_spacing), BODY_STYLE["line_spacing_pt"]):
            add(
                result,
                "error",
                "body_line_spacing",
                "style Normal",
                f"Body line spacing must be exactly {BODY_STYLE['line_spacing_pt']} pt.",
            )
        if not close(effective_style_space_pt(normal, "space_before"), 0) or not close(
            effective_style_space_pt(normal, "space_after"), 0
        ):
            add(result, "error", "body_spacing", "style Normal", "Body paragraph spacing must be 0 pt before and after.")

    style_specs = (
        (
            "报告题名",
            REPORT_TITLE_STYLE["east_asia"],
            REPORT_TITLE_STYLE["latin"],
            REPORT_TITLE_STYLE["size_pt"],
            True,
            WD_ALIGN_PARAGRAPH.CENTER,
            True,
        ),
        ("目录标题", TOC_TITLE_FONT_EAST_ASIA, TOC_TITLE_FONT_LATIN, 16, True, WD_ALIGN_PARAGRAPH.CENTER, True),
        (
            "报告短题名",
            REPORT_SUBTITLE_STYLE["east_asia"],
            REPORT_SUBTITLE_STYLE["latin"],
            REPORT_SUBTITLE_STYLE["size_pt"],
            True,
            WD_ALIGN_PARAGRAPH.CENTER,
            True,
        ),
        ("报告信息", BODY_FONT_EAST_ASIA, BODY_FONT_LATIN, 12, False, WD_ALIGN_PARAGRAPH.LEFT, False),
        ("图表题注", BODY_FONT_EAST_ASIA, BODY_FONT_LATIN, 10.5, False, WD_ALIGN_PARAGRAPH.CENTER, False),
        ("表格文字", BODY_FONT_EAST_ASIA, BODY_FONT_LATIN, 9.5, False, WD_ALIGN_PARAGRAPH.CENTER, False),
        ("表格表头", BODY_FONT_EAST_ASIA, BODY_FONT_LATIN, 9.5, True, WD_ALIGN_PARAGRAPH.CENTER, False),
        ("图片", BODY_FONT_EAST_ASIA, BODY_FONT_LATIN, 12, False, WD_ALIGN_PARAGRAPH.CENTER, False),
        ("公式", BODY_FONT_EAST_ASIA, BODY_FONT_LATIN, 12, False, WD_ALIGN_PARAGRAPH.CENTER, False),
    )
    for name, east_asia, latin, size, bold, alignment, reject_theme in style_specs:
        style = get_style(doc, result, name)
        if style is not None:
            audit_style_font(
                result,
                style,
                name,
                east_asia,
                latin,
                size,
                bold,
                alignment,
                zero_indent=True,
                reject_theme=reject_theme,
                default_size_pt=default_size,
            )
    title_style = doc.styles["报告题名"]
    if style_color(title_style) != REPORT_TITLE_STYLE["color"]:
        add(result, "error", "style_color", "style 报告题名", f"Expected #{REPORT_TITLE_STYLE['color']} title color.")
    subtitle_style = doc.styles["报告短题名"]
    if style_color(subtitle_style) != REPORT_SUBTITLE_STYLE["color"]:
        add(result, "error", "style_color", "style 报告短题名", f"Expected #{REPORT_SUBTITLE_STYLE['color']} subtitle color.")
    for level in (1, 2, 3):
        contract = HEADING_STYLES[level]
        style = get_style(doc, result, contract["name"])
        if style is None:
            continue
        audit_style_font(
            result,
            style,
            contract["name"],
            contract["east_asia"],
            contract["latin"],
            contract["size_pt"],
            contract["bold"],
            WD_ALIGN_PARAGRAPH.JUSTIFY,
            zero_indent=False,
            reject_theme=True,
            default_size_pt=default_size,
        )
        if (
            not close(effective_style_indent_pt(style, "left_indent"), 0)
            or not close(effective_style_indent_pt(style, "right_indent"), 0)
            or not close(effective_style_indent_pt(style, "first_line_indent"), contract["first_line_pt"])
        ):
            add(result, "error", "heading_style_indent", f"style {contract['name']}", "Heading style indent differs from the original reference.")
        if not close(effective_style_space_pt(style, "space_before"), contract["space_before_pt"]) or not close(
            effective_style_space_pt(style, "space_after"), contract["space_after_pt"]
        ):
            add(result, "error", "heading_spacing", f"style {contract['name']}", "Heading paragraph spacing does not match the original reference.")
        effective_line_spacing = effective_style_value(style, lambda item: item.paragraph_format.line_spacing, None)
        if effective_line_spacing != contract["line_spacing"]:
            add(
                result,
                "error",
                "heading_line_spacing",
                f"style {contract['name']}",
                f"Heading line spacing must be {contract['line_spacing']}.",
            )
        if bool(style.paragraph_format.keep_with_next) != contract["keep_with_next"] or bool(
            style.paragraph_format.keep_together
        ) != contract["keep_together"]:
            add(result, "error", "heading_pagination", f"style {contract['name']}", "Heading pagination controls differ from the active contract.")
        if style_color(style) != contract["color"]:
            add(result, "error", "style_color", f"style {contract['name']}", f"Expected #{contract['color']} heading color.")
        outline = style.element.get_or_add_pPr().find(qn("w:outlineLvl"))
        if outline is None or outline.get(qn("w:val")) != str(contract["outline_level"]):
            add(result, "error", "heading_outline", f"style {contract['name']}", "Heading outline level is incorrect.")
        if style.hidden or not style.quick_style:
            add(result, "error", "heading_quick_style", f"style {contract['name']}", "Heading must be a visible quick style.")

    for level, text, left, hanging, suffix in NUMBERING_LEVELS:
        style = doc.styles[HEADING_STYLES[level + 1]["name"]]
        num_id, ilvl = style_numbering(style)
        if num_id is None or ilvl != str(level):
            add(result, "error", "heading_style_numbering", f"style {style.name}", "Heading style is not bound to the expected list level.")
            continue
        details = numbering_level_details(doc, num_id, ilvl)
        expected = {"text": text, "suffix": suffix, "p_style": style.style_id, "left": left, "hanging": hanging}
        if details != expected:
            add(result, "error", "heading_numbering_definition", f"style {style.name}", f"Expected {expected}; found {details}.")


def audit_toc_styles(doc: Document, result: AuditResult) -> None:
    default_size = document_default_size_pt(doc)
    for level, expected_left in TOC_LEFT_INDENTS_PT.items():
        name = f"TOC {level}"
        style = get_style(doc, result, name)
        if style is None:
            continue
        audit_style_font(
            result,
            style,
            name,
            BODY_FONT_EAST_ASIA,
            BODY_FONT_LATIN,
            12,
            False,
            WD_ALIGN_PARAGRAPH.LEFT,
            zero_indent=False,
            default_size_pt=default_size,
        )
        if not close(effective_style_indent_pt(style, "first_line_indent"), 0) or not close(
            effective_style_indent_pt(style, "left_indent"), expected_left
        ):
            add(
                result,
                "error",
                "toc_indent",
                f"style {name}",
                f"Expected first-line 0 pt and left {expected_left} pt.",
            )
        tabs = style.element.get_or_add_pPr().find(qn("w:tabs"))
        valid_tab = False
        if tabs is not None:
            for tab in tabs.findall(qn("w:tab")):
                if (
                    tab.get(qn("w:val")) == "right"
                    and tab.get(qn("w:leader")) == "dot"
                    and abs(int(tab.get(qn("w:pos"), "-99999")) - TOC_RIGHT_TAB_TWIPS) <= 2
                ):
                    valid_tab = True
                    break
        if not valid_tab:
            add(result, "error", "toc_page_tab", f"style {name}", "Missing 15.5 cm right-aligned dotted page-number tab.")


def audit_headings(doc: Document, result: AuditResult) -> None:
    heading_num_id, _ = style_numbering(doc.styles[HEADING_STYLES[1]["name"]])
    all_paragraphs = list(doc.paragraphs)
    seen_cells: set[int] = set()
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                marker = id(cell._tc)
                if marker in seen_cells:
                    continue
                seen_cells.add(marker)
                all_paragraphs.extend(cell.paragraphs)
    for index, paragraph in enumerate(all_paragraphs, 1):
        if heading_level(paragraph) is None and heading_num_id is not None and direct_num_id(paragraph) == heading_num_id:
            add(result, "error", "heading_numbering_collision", f"paragraph {index}", "A non-heading paragraph shares the heading list instance.")

    previous = 0
    first_heading_seen = False
    for index, paragraph in enumerate(doc.paragraphs, 1):
        text = clean(paragraph.text)
        level = heading_level(paragraph)
        if not text or level is None:
            continue
        result.heading_levels.append(level)
        location = f"paragraph {index}"
        expected_style = HEADING_STYLES[level]["name"]
        if paragraph.style.name != expected_style:
            add(result, "error", "heading_style", location, f"Expected {expected_style}; found {paragraph.style.name}.")
        p_pr = paragraph._p.pPr
        forbidden = ("w:numPr", "w:ind", "w:spacing", "w:jc", "w:keepNext", "w:keepLines", "w:pageBreakBefore")
        direct = [tag for tag in forbidden if p_pr is not None and p_pr.find(qn(tag)) is not None]
        if direct:
            add(result, "error", "heading_direct_format", location, f"Heading retains direct paragraph formatting: {', '.join(direct)}.")
        if paragraph._p.xpath(".//w:rPr/w:rFonts | .//w:rPr/w:sz | .//w:rPr/w:szCs"):
            add(result, "error", "heading_direct_font", location, "Heading retains direct font or size formatting.")
        if MANUAL_HEADING_PATTERN.match(text) and not MEASUREMENT_HEADING_PATTERN.match(text):
            add(result, "error", "manual_heading_number", location, "Heading text contains a manual sequence number.")
        if not first_heading_seen:
            first_heading_seen = True
            if level != 1:
                add(result, "error", "heading_hierarchy", location, "First body heading must use 一级标题.")
        elif level > previous + 1:
            add(result, "error", "heading_hierarchy", location, f"Heading level jumps from {previous} to {level}.")
        previous = level

    if not result.heading_levels:
        add(result, "error", "heading_hierarchy", "document", "No custom body headings were found.")


def audit_document_title(doc: Document, result: AuditResult) -> None:
    first = next((paragraph for paragraph in doc.paragraphs if clean(paragraph.text)), None)
    if first is None:
        add(result, "error", "report_title", "document", "Document has no report title.")
        return
    if first.style.name != "报告题名":
        add(result, "error", "report_title", "paragraph 1", f"Report title must use 报告题名; found {first.style.name}.")
    if has_direct_numbering(first):
        add(result, "error", "report_title_numbering", "paragraph 1", "Report title must not retain list numbering.")


def audit_metadata(doc: Document, result: AuditResult) -> None:
    for index, paragraph in enumerate(doc.paragraphs, 1):
        text = clean(paragraph.text)
        if not METADATA_PATTERN.match(text):
            continue
        if heading_level(paragraph) is not None:
            add(result, "error", "metadata_style", f"paragraph {index}", "Report-information paragraph uses a heading style.")
        if paragraph.style.name != "报告信息":
            add(result, "warning", "metadata_style", f"paragraph {index}", f"Expected 报告信息 style, found {paragraph.style.name}.")
        if has_direct_numbering(paragraph):
            add(result, "error", "metadata_numbering", f"paragraph {index}", "Report information must not retain list numbering.")

    cover_text = "\n".join(clean(cell.text) for table in doc.tables[:2] for row in table.rows for cell in row.cells)
    required_labels = ("版本", "测试时间", "报告日期", "测试地点", "测试人员", "编写", "审核", "批准", "修订记录")
    for label in required_labels:
        if label not in cover_text and not any(label in clean(paragraph.text) for paragraph in doc.paragraphs[:12]):
            add(result, "error", "cover_metadata", "front matter", f"Missing cover field: {label}.")


def audit_front_matter_page_break(doc: Document, result: AuditResult) -> None:
    paragraphs = list(doc.paragraphs)
    toc_index = next(
        (
            index
            for index, paragraph in enumerate(paragraphs)
            if paragraph.style.name == "目录标题" or clean(paragraph.text) == "目录"
        ),
        None,
    )
    if toc_index is None:
        add(result, "error", "toc_title", "document", "Document has no directory title.")
        return
    report_info_indices = [
        index
        for index, paragraph in enumerate(paragraphs[:toc_index])
        if paragraph.style.name == "报告信息" or METADATA_PATTERN.match(clean(paragraph.text))
    ]
    title_index = next((index for index, paragraph in enumerate(paragraphs[:toc_index]) if clean(paragraph.text)), 0)
    anchor_index = max(report_info_indices, default=title_index)
    between = paragraphs[anchor_index + 1 : toc_index]
    if not any(
        paragraph_has_explicit_page_break(paragraph) or paragraph_has_next_page_section_break(paragraph)
        for paragraph in between
    ) and not (
        paragraph_has_explicit_page_break(paragraphs[toc_index])
        or paragraph_has_next_page_section_break(paragraphs[toc_index])
    ):
        add(
            result,
            "error",
            "front_matter_page_break",
            f"paragraph {toc_index + 1}",
            "Directory must start on a new page after the report-information block.",
        )


def story_instructions(story) -> str:
    return " ".join(
        node.text or ""
        for paragraph in story.paragraphs
        for node in paragraph._p.xpath(".//w:instrText")
    )


def page_number_settings(section) -> tuple[str | None, str | None]:
    node = section._sectPr.find(qn("w:pgNumType"))
    if node is None:
        return None, None
    return node.get(qn("w:fmt")), node.get(qn("w:start"))


def audit_page_system(doc: Document, result: AuditResult) -> None:
    expected_geometry = {
        "page_width": PAGE_STYLE["page_width_cm"],
        "page_height": PAGE_STYLE["page_height_cm"],
        "top_margin": PAGE_STYLE["top_margin_cm"],
        "bottom_margin": PAGE_STYLE["bottom_margin_cm"],
        "left_margin": PAGE_STYLE["left_margin_cm"],
        "right_margin": PAGE_STYLE["right_margin_cm"],
        "header_distance": PAGE_STYLE["header_distance_cm"],
        "footer_distance": PAGE_STYLE["footer_distance_cm"],
    }
    for index, section in enumerate(doc.sections, 1):
        for attribute, expected in expected_geometry.items():
            actual = getattr(section, attribute)
            actual_cm = None if actual is None else actual.cm
            if actual_cm is None or abs(actual_cm - expected) > 0.03:
                add(
                    result,
                    "error",
                    "page_geometry",
                    f"section {index} {attribute}",
                    f"Expected {expected} cm; found {actual_cm}.",
                )
    if len(doc.sections) < 3:
        add(result, "error", "page_sections", "document", "Cover, directory, and body must use three new-page sections.")
        return

    cover, toc, body = doc.sections[:3]
    cover_fields = f"{story_instructions(cover.header)} {story_instructions(cover.footer)}"
    if re.search(r"\b(?:PAGE|STYLEREF)\b", cover_fields, re.IGNORECASE):
        add(result, "error", "cover_page_furniture", "section 1", "Cover must not show a running header or page number.")

    toc_fmt, toc_start = page_number_settings(toc)
    if toc_fmt != "lowerRoman" or toc_start != "1" or not re.search(
        r"\bPAGE\b", story_instructions(toc.footer), re.IGNORECASE
    ):
        add(result, "error", "toc_page_number", "section 2", "Directory must use lower-Roman PAGE numbering starting at i.")

    body_fmt, body_start = page_number_settings(body)
    if body_fmt not in (None, "decimal") or body_start != "1" or not re.search(
        r"\bPAGE\b", story_instructions(body.footer), re.IGNORECASE
    ):
        add(result, "error", "body_page_number", "section 3", "Body must use decimal PAGE numbering starting at 1.")
    if not re.search(r"\bSTYLEREF\b", story_instructions(body.header), re.IGNORECASE):
        add(result, "error", "body_running_header", "section 3", "Body header must use the report short-title STYLEREF field.")


def audit_captions(doc: Document, result: AuditResult) -> None:
    for index, paragraph in enumerate(doc.paragraphs, 1):
        text = clean(paragraph.text)
        is_standard_caption = paragraph.style.name == "图表题注"
        match = CAPTION_PATTERN.match(text)
        if not is_standard_caption and not match:
            continue
        if match:
            label = match.group(1)
        else:
            instruction = "".join(node.text or "" for node in paragraph._p.xpath(".//w:instrText"))
            label = "图" if re.search(r"\bSEQ\s*图\b", instruction) else "表"
        location = f"paragraph {index}"
        if paragraph.style.name != "图表题注":
            add(result, "error", "caption_style", location, "Caption must use 图表题注 style.")
        if has_direct_numbering(paragraph):
            add(result, "error", "caption_numbering", location, "Caption must not retain list numbering.")
        if not has_seq_field(paragraph, label):
            add(result, "error", "caption_sequence", location, f"{label} caption lacks an automatic SEQ field.")
        if effective_alignment(paragraph) != WD_ALIGN_PARAGRAPH.CENTER:
            add(result, "error", "caption_alignment", location, "Caption must remain centered.")
        if not has_zero_indent(paragraph):
            add(result, "error", "caption_indent", location, "Caption must explicitly use zero first-line, left, and right indent.")


def _cell_fill(cell) -> str | None:
    shading = cell._tc.get_or_add_tcPr().find(qn("w:shd"))
    return shading.get(qn("w:fill")) if shading is not None else None


def _header_cell_has_white_text(cell) -> bool:
    visible_runs = [run for paragraph in cell.paragraphs for run in paragraph.runs if clean(run.text)]
    if not visible_runs:
        return True
    for run in visible_runs:
        r_pr = run._r.get_or_add_rPr()
        color = r_pr.find(qn("w:color"))
        if color is None or (color.get(qn("w:val")) or "").upper() != TABLE_STYLE["header_font_color"]:
            return False
    return True


def _cell_margin(cell, side: str) -> str | None:
    margins = cell._tc.get_or_add_tcPr().find(qn("w:tcMar"))
    if margins is None:
        return None
    node = margins.find(qn(f"w:{side}"))
    return None if node is None else node.get(qn("w:w"))


def _row_cannot_split(row) -> bool:
    tr_pr = row._tr.get_or_add_trPr()
    return tr_pr.find(qn("w:cantSplit")) is not None


def _table_border_matches(table, edge: str) -> bool:
    borders = table._tbl.tblPr.find(qn("w:tblBorders"))
    if borders is None:
        return False
    border = borders.find(qn(f"w:{edge}"))
    if border is None:
        return False
    return (
        border.get(qn("w:val")) == "single"
        and border.get(qn("w:sz")) == TABLE_STYLE["border_size_eighth_pt"]
        and border.get(qn("w:color"), "").upper() == TABLE_STYLE["border_color"]
    )


def audit_tables(doc: Document, result: AuditResult, strict_visual_style: bool = False) -> None:
    max_width = document_body_width_twips(doc)
    for table_index, table in enumerate(doc.tables, 1):
        if strict_visual_style:
            if table.style is None or table.style.name != TABLE_STYLE["name"]:
                add(result, "error", "table_visual_style", f"table {table_index}", f"Expected table style {TABLE_STYLE['name']}.")
            for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
                if not _table_border_matches(table, edge):
                    add(result, "error", "table_border_style", f"table {table_index} {edge}", "Expected 0.5 pt #808080 single border.")
            for cell_index, cell in enumerate(table.rows[0].cells, 1):
                if (_cell_fill(cell) or "").upper() != TABLE_STYLE["header_fill"]:
                    add(result, "error", "table_header_fill", f"table {table_index} header {cell_index}", f"Expected #{TABLE_STYLE['header_fill']} fill.")
                if not _header_cell_has_white_text(cell):
                    add(result, "error", "table_header_font_color", f"table {table_index} header {cell_index}", f"Expected bold white #{TABLE_STYLE['header_font_color']} header text.")
            conclusion_columns = [
                index for index, cell in enumerate(table.rows[0].cells) if clean(cell.text) == "结论"
            ]
            for row_index, row in enumerate(table.rows[1:], 2):
                for cell_index, cell in enumerate(row.cells, 1):
                    expected_fill = TABLE_STYLE["body_fill"]
                    if cell_index - 1 in conclusion_columns:
                        expected_fill = TABLE_STYLE["status_fills"].get(clean(cell.text), expected_fill)
                    if (_cell_fill(cell) or "").upper() != expected_fill:
                        add(
                            result,
                            "error",
                            "table_body_fill",
                            f"table {table_index} row {row_index} col {cell_index}",
                            f"Expected #{expected_fill} fill.",
                        )
        if table.alignment != WD_TABLE_ALIGNMENT.CENTER:
            add(result, "error", "table_object_alignment", f"table {table_index}", "Table object must be centered.")
        table_indent = table._tbl.tblPr.find(qn("w:tblInd"))
        if table_indent is not None and table_indent.get(qn("w:w"), "0") != "0":
            add(result, "error", "table_object_indent", f"table {table_index}", "Centered table must not retain a nonzero table indent.")
        grid_width = sum(int(item.get(qn("w:w"), "0")) for item in table._tbl.tblGrid.findall(qn("w:gridCol")))
        if grid_width > max_width + 2:
            add(
                result,
                "error",
                "table_width",
                f"table {table_index}",
                f"Table width {grid_width} twips exceeds body width {max_width} twips.",
            )
        if not table_header_repeats(table):
            add(result, "warning", "table_header_repeat", f"table {table_index}", "First row is not marked as a repeated header row.")
        headers = [clean(cell.text) for cell in table.rows[0].cells]
        for header_index, header in enumerate(headers, 1):
            is_reference_column = any(token in header for token in ("图", "图像", "路径", "索引", "文件"))
            if (
                not is_reference_column
                and any(metric in header for metric in UNIT_METRICS)
                and not re.search(r"[（(].+[）)]", header)
            ):
                add(result, "warning", "table_units", f"table {table_index} header {header_index}", "Metric-like header has no explicit unit.")
        for row_index, row in enumerate(table.rows, 1):
            if strict_visual_style and not _row_cannot_split(row):
                add(result, "error", "table_row_split", f"table {table_index} row {row_index}", "Table row must not split across pages.")
            for cell_index, cell in enumerate(row.cells, 1):
                if strict_visual_style:
                    expected_margins = {
                        "top": TABLE_STYLE["cell_margin_top_twips"],
                        "bottom": TABLE_STYLE["cell_margin_bottom_twips"],
                        "left": TABLE_STYLE["cell_margin_left_twips"],
                        "right": TABLE_STYLE["cell_margin_right_twips"],
                    }
                    for side, expected in expected_margins.items():
                        if _cell_margin(cell, side) != expected:
                            add(
                                result,
                                "error",
                                "table_cell_margin",
                                f"table {table_index} row {row_index} col {cell_index} {side}",
                                f"Expected {expected} twips.",
                            )
                if cell.vertical_alignment != WD_CELL_VERTICAL_ALIGNMENT.CENTER:
                    add(
                        result,
                        "error",
                        "table_vertical_alignment",
                        f"table {table_index} row {row_index} col {cell_index}",
                        "Table cell must be vertically centered.",
                    )
                for paragraph_index, paragraph in enumerate(cell.paragraphs, 1):
                    location = f"table {table_index} row {row_index} col {cell_index} paragraph {paragraph_index}"
                    if has_direct_numbering(paragraph):
                        add(result, "error", "table_paragraph_numbering", location, "Table content must not retain list numbering.")
                    if not clean(paragraph.text) and not contains_image(paragraph):
                        continue
                    expected_style = "图片" if contains_image(paragraph) else ("表格表头" if row_index == 1 else "表格文字")
                    if paragraph.style.name != expected_style:
                        add(
                            result,
                            "error",
                            "table_paragraph_style",
                            location,
                            f"Expected {expected_style}; found {paragraph.style.name}.",
                        )
                    if effective_alignment(paragraph) != WD_ALIGN_PARAGRAPH.CENTER:
                        add(result, "error", "table_alignment", location, "All table content must remain centered.")
                    if not has_zero_indent(paragraph):
                        add(
                            result,
                            "error",
                            "table_zero_indent",
                            location,
                            "Table-cell paragraph must explicitly set firstLine, left, and right indent to zero.",
                        )


def audit_math(doc: Document, result: AuditResult) -> None:
    paragraphs = list(doc.paragraphs)
    seen_cells: set[int] = set()
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                marker = id(cell._tc)
                if marker in seen_cells:
                    continue
                seen_cells.add(marker)
                paragraphs.extend(cell.paragraphs)
    for index, paragraph in enumerate(paragraphs, 1):
        for math in paragraph._p.xpath(".//m:oMath"):
            text = "".join(node.text or "" for node in math.iter(qn("m:t")))
            structured = any(math.find(".//" + qn(tag)) is not None for tag in MATH_STRUCTURE_TAGS)
            if not structured and re.search(r"_|sqrt|integral|∫\[|\\frac|\\sqrt", text):
                add(result, "error", "pseudo_omml", f"math paragraph {index}", f"Unbuilt math text remains: {text}")
        if heading_level(paragraph) is not None or paragraph.style.name.lower().startswith("toc"):
            continue
        if paragraph.style.name in {"报告题名", "目录标题", "图表题注", "图片"}:
            continue
        match = KNOWN_INLINE_LITERAL.search(paragraph.text)
        if match:
            add(
                result,
                "error",
                "inline_math_literal",
                f"paragraph {index}",
                f"Known mathematical quantity remains as plain text: {match.group(0)}",
            )


def audit_images(doc: Document, result: AuditResult) -> None:
    for index, paragraph in enumerate(doc.paragraphs, 1):
        if not contains_image(paragraph):
            continue
        location = f"paragraph {index}"
        if paragraph.style.name != "图片":
            add(result, "error", "image_style", location, f"Image paragraph must use 图片; found {paragraph.style.name}.")
        if has_direct_numbering(paragraph):
            add(result, "error", "image_numbering", location, "Image paragraph must not retain list numbering.")
        if effective_alignment(paragraph) != WD_ALIGN_PARAGRAPH.CENTER:
            add(result, "error", "image_alignment", location, "Image paragraph must remain centered.")
        if not has_zero_indent(paragraph):
            add(result, "error", "image_indent", location, "Image paragraph must explicitly use zero first-line, left, and right indent.")


def audit_language_and_required_sections(
    doc: Document,
    result: AuditResult,
    method_note: bool = False,
    final_content: bool = False,
) -> None:
    text = "\n".join(clean(p.text) for p in doc.paragraphs)
    for table in doc.tables:
        text += "\n" + "\n".join(clean(cell.text) for row in table.rows for cell in row.cells)
    for term in LANGUAGE_REVIEW_TERMS:
        if term in text:
            add(result, "warning", "language_review", "document", f"Review legacy wording: {term}.")
    for term in ("PASS", "FAIL"):
        if re.search(rf"\b{term}\b", text, re.IGNORECASE):
            add(result, "error", "formal_status", "document", f"Use Chinese formal status instead of {term}.")
    for term in ("通过", "未通过", "待补测"):
        if term in text:
            add(result, "error", "formal_status", "document", f"Use 满足/不满足/暂不能判定/未测试 instead of {term}.")
    for term in ("基本通过", "大致满足", "应该通过"):
        if term in text:
            add(result, "warning", "formal_status", "document", f"Replace ambiguous status wording: {term}.")
    for match in FULLWIDTH_TECHNICAL_PARENS.finditer(text):
        add(result, "error", "technical_parentheses", "document", f"Use ASCII parentheses around English content: {match.group(0)}")
    if final_content:
        for placeholder in FINAL_PLACEHOLDERS:
            if placeholder in text:
                add(result, "error", "final_placeholder", "document", f"Final report still contains template text: {placeholder}.")
    if not method_note:
        required_groups = (
            ("follow-up tests", ("总结与后续测试", "后续测试", "下一步计划")),
            ("data index", ("数据来源", "数据与脚本索引", "原始记录与数据索引")),
        )
        for label, alternatives in required_groups:
            if not any(heading in text for heading in alternatives):
                add(result, "error", "required_section", "document", f"Missing required section: {label}.")
    if method_note and "数据来源" not in text and "引用来源" not in text:
        add(result, "error", "required_section", "document", "Method note must contain 数据来源 or 引用来源.")
    formula_tokens = ("A_v(", "S_v(", "V_rms", "sqrt(", "√", "∫", "\\frac")
    if any(token in text for token in formula_tokens) and result.office_math_count == 0:
        add(result, "warning", "office_math", "document", "Noise or calibration report has no Office Math object.")
    if "AD 输入端等效噪声" in text:
        for required in ("K_out", "L_ADC"):
            if required not in text:
                add(result, "warning", "adc_equivalent_noise", "document", f"AD input-equivalent noise section lacks {required}.")
        if not any(term in text for term in ("H_FPGA", "G/H_FPGA", "FPGA 增益")):
            add(result, "warning", "adc_equivalent_noise", "document", "AD input-equivalent noise section lacks FPGA gain or transfer term.")


def audit(
    path: Path,
    method_note: bool = False,
    strict_table_style: bool = False,
    final_content: bool = False,
) -> AuditResult:
    doc = Document(path)
    result = AuditResult(
        path=str(path),
        paragraph_count=len(doc.paragraphs),
        table_count=len(doc.tables),
        image_count=len(doc.inline_shapes),
        office_math_count=count_office_math(path),
    )
    audit_style_contract(doc, result)
    audit_toc_styles(doc, result)
    audit_document_title(doc, result)
    audit_headings(doc, result)
    audit_metadata(doc, result)
    audit_front_matter_page_break(doc, result)
    audit_page_system(doc, result)
    audit_captions(doc, result)
    audit_tables(doc, result, strict_table_style)
    audit_math(doc, result)
    audit_images(doc, result)
    audit_language_and_required_sections(doc, result, method_note, final_content)
    return result


def main() -> int:
    sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("docx", type=Path)
    parser.add_argument("--json", action="store_true")
    parser.add_argument("--warn-only", action="store_true")
    parser.add_argument("--method-note", action="store_true", help="Allow a same-content experiment-method document without result-report-only sections.")
    parser.add_argument("--strict-table-style", action="store_true", help="Require the project blue-header table visual contract on every table.")
    parser.add_argument("--final-content", action="store_true", help="Reject template placeholders and other unfinished final-report content.")
    args = parser.parse_args()
    path = args.docx.resolve()
    if not path.is_file():
        parser.error(f"DOCX not found: {path}")
    result = audit(path, args.method_note, args.strict_table_style, args.final_content)
    if args.json:
        print(json.dumps(asdict(result), ensure_ascii=False, indent=2))
    else:
        print(f"path: {result.path}")
        print(f"paragraphs: {result.paragraph_count}")
        print(f"tables: {result.table_count}")
        print(f"images: {result.image_count}")
        print(f"office_math: {result.office_math_count}")
        print(f"errors: {result.errors}")
        print(f"warnings: {result.warnings}")
        for finding in result.findings:
            print(f"[{finding.severity}] {finding.check} @ {finding.location}: {finding.message}")
    return 0 if args.warn_only or result.errors == 0 else 1


if __name__ == "__main__":
    raise SystemExit(main())

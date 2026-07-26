#!/usr/bin/env python3
"""Create the standard DOCX template for ultra-stable laser electrical tests."""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.text.paragraph import Paragraph

from latex_to_omml import latex_to_display_omml
from style_contract import (
    BODY_STYLE,
    HEADING_STYLES,
    NUMBERING_LEVELS,
    PAGE_STYLE,
    REPORT_SUBTITLE_STYLE,
    REPORT_TITLE_STYLE,
    TABLE_STYLE,
)


SKILL_DIR = Path(__file__).resolve().parents[1]
DEFAULT_OUTPUT = SKILL_DIR / "assets" / "超稳激光电控箱测试报告标准模板.docx"
REPORT_ROOT = Path(
    r"C:\01_obsidian\ustc_hy\01_科研主系统\04_实验测试与证据"
    r"\03_超稳激光电控箱指标测试"
    r"\LASER 测试报告\202607\01_正式报告"
)
BODY_FONT_EAST_ASIA = "宋体"
BODY_FONT_LATIN = "Times New Roman"
TOC_TITLE_FONT_EAST_ASIA = "微软雅黑"
TOC_TITLE_FONT_LATIN = "Microsoft YaHei"
TOC_RIGHT_TAB_CM = 15.5
ACCENT_BLUE = "1F4E79"
MUTED_GRAY = "7F7F7F"


def ensure_child(parent, tag: str):
    child = parent.find(qn(tag))
    if child is None:
        child = OxmlElement(tag)
        parent.append(child)
    return child


def set_table_visual_contract(table) -> None:
    """Apply the project blue-header/gray-grid visual contract."""
    try:
        table.style = TABLE_STYLE["name"]
    except KeyError:
        table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    indent = tbl_pr.find(qn("w:tblInd"))
    if indent is not None:
        tbl_pr.remove(indent)
    layout = ensure_child(tbl_pr, "w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    borders = ensure_child(tbl_pr, "w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = ensure_child(borders, f"w:{edge}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), TABLE_STYLE["border_size_eighth_pt"])
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), TABLE_STYLE["border_color"])
    for row_index, row in enumerate(table.rows):
        fill = TABLE_STYLE["header_fill"] if row_index == 0 else TABLE_STYLE["body_fill"]
        for cell in row.cells:
            tc_pr = cell._tc.get_or_add_tcPr()
            shading = ensure_child(tc_pr, "w:shd")
            shading.set(qn("w:val"), "clear")
            shading.set(qn("w:color"), "auto")
            shading.set(qn("w:fill"), fill)
            margins = ensure_child(tc_pr, "w:tcMar")
            margin_values = {
                "top": TABLE_STYLE["cell_margin_top_twips"],
                "left": TABLE_STYLE["cell_margin_left_twips"],
                "bottom": TABLE_STYLE["cell_margin_bottom_twips"],
                "right": TABLE_STYLE["cell_margin_right_twips"],
            }
            for side, value in margin_values.items():
                margin = ensure_child(margins, f"w:{side}")
                margin.set(qn("w:w"), value)
                margin.set(qn("w:type"), "dxa")


def set_font(
    style,
    east_asia: str,
    size_pt: float,
    bold: bool = False,
    latin: str = BODY_FONT_LATIN,
    color: str = "000000",
) -> None:
    style.font.name = latin
    style.font.size = Pt(size_pt)
    style.font.bold = bold
    style.font.color.rgb = RGBColor.from_string(color)
    r_pr = style.element.get_or_add_rPr()
    r_fonts = ensure_child(r_pr, "w:rFonts")
    r_fonts.set(qn("w:ascii"), latin)
    r_fonts.set(qn("w:hAnsi"), latin)
    r_fonts.set(qn("w:eastAsia"), east_asia)
    r_fonts.set(qn("w:cs"), latin)
    for attribute in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        qualified = qn(attribute)
        if qualified in r_fonts.attrib:
            del r_fonts.attrib[qualified]


def set_run_font(
    run,
    east_asia: str,
    size_pt: float,
    bold: bool | None = None,
    latin: str = BODY_FONT_LATIN,
    color: str = "000000",
) -> None:
    run.font.name = latin
    run.font.size = Pt(size_pt)
    if bold is not None:
        run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    r_pr = run._r.get_or_add_rPr()
    r_fonts = ensure_child(r_pr, "w:rFonts")
    r_fonts.set(qn("w:ascii"), latin)
    r_fonts.set(qn("w:hAnsi"), latin)
    r_fonts.set(qn("w:eastAsia"), east_asia)
    r_fonts.set(qn("w:cs"), latin)
    for attribute in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        qualified = qn(attribute)
        if qualified in r_fonts.attrib:
            del r_fonts.attrib[qualified]


def set_outline_level(style, level: int) -> None:
    p_pr = style.element.get_or_add_pPr()
    outline = ensure_child(p_pr, "w:outlineLvl")
    outline.set(qn("w:val"), str(level))


def set_heading_style_indent(style, level: int) -> None:
    """Canonicalize point and character-unit indents for Chinese Word files."""
    first_line_pt = HEADING_STYLES[level]["first_line_pt"]
    style.paragraph_format.left_indent = None
    style.paragraph_format.right_indent = None
    style.paragraph_format.first_line_indent = Pt(first_line_pt)
    p_pr = style.element.get_or_add_pPr()
    tabs = p_pr.find(qn("w:tabs"))
    if tabs is not None:
        p_pr.remove(tabs)
    ind = p_pr.find(qn("w:ind"))
    if ind is None:
        ind = OxmlElement("w:ind")
        p_pr.append(ind)
    for attribute in (
        "w:left",
        "w:right",
        "w:hanging",
        "w:leftChars",
        "w:rightChars",
        "w:hangingChars",
        "w:firstLineChars",
    ):
        qualified = qn(attribute)
        if qualified in ind.attrib:
            del ind.attrib[qualified]
    ind.set(qn("w:firstLine"), str(int(round(first_line_pt * 20))))
    if level in (1, 2, 3):
        ind.set(qn("w:firstLineChars"), "0")


def set_zero_indent(paragraph) -> None:
    paragraph.paragraph_format.first_line_indent = Pt(0)
    paragraph.paragraph_format.left_indent = Pt(0)
    paragraph.paragraph_format.right_indent = Pt(0)
    p_pr = paragraph._p.get_or_add_pPr()
    for tag in ("w:numPr", "w:tabs"):
        child = p_pr.find(qn(tag))
        if child is not None:
            p_pr.remove(child)
    ind = p_pr.find(qn("w:ind"))
    if ind is None:
        ind = OxmlElement("w:ind")
        p_pr.append(ind)
    ind.set(qn("w:firstLine"), "0")
    ind.set(qn("w:left"), "0")
    ind.set(qn("w:right"), "0")
    ind.set(qn("w:firstLineChars"), "0")
    ind.set(qn("w:leftChars"), "0")
    ind.set(qn("w:rightChars"), "0")
    for attribute in ("w:hanging", "w:hangingChars"):
        qualified = qn(attribute)
        if qualified in ind.attrib:
            del ind.attrib[qualified]


def set_style_zero_indent(style) -> None:
    style.paragraph_format.first_line_indent = Pt(0)
    style.paragraph_format.left_indent = Pt(0)
    style.paragraph_format.right_indent = Pt(0)
    p_pr = style.element.get_or_add_pPr()
    for tag in ("w:numPr", "w:tabs"):
        child = p_pr.find(qn(tag))
        if child is not None:
            p_pr.remove(child)
    ind = ensure_child(p_pr, "w:ind")
    ind.set(qn("w:firstLine"), "0")
    ind.set(qn("w:left"), "0")
    ind.set(qn("w:right"), "0")
    ind.set(qn("w:firstLineChars"), "0")
    ind.set(qn("w:leftChars"), "0")
    ind.set(qn("w:rightChars"), "0")
    for attribute in ("w:hanging", "w:hangingChars"):
        qualified = qn(attribute)
        if qualified in ind.attrib:
            del ind.attrib[qualified]


def configure_toc_style(style, level: int) -> None:
    set_font(style, BODY_FONT_EAST_ASIA, 12)
    style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    style.paragraph_format.first_line_indent = Pt(0)
    style.paragraph_format.left_indent = Pt((level - 1) * 21)
    style.paragraph_format.right_indent = Pt(0)
    style.paragraph_format.line_spacing = 1.5
    style.paragraph_format.space_before = Pt(0)
    style.paragraph_format.space_after = Pt(0)
    tab_stops = style.paragraph_format.tab_stops
    tab_stops.clear_all()
    tab_stops.add_tab_stop(
        Cm(TOC_RIGHT_TAB_CM),
        WD_TAB_ALIGNMENT.RIGHT,
        WD_TAB_LEADER.DOTS,
    )
    p_pr = style.element.get_or_add_pPr()
    outline = ensure_child(p_pr, "w:outlineLvl")
    outline.set(qn("w:val"), "9")


def set_update_fields(doc: Document) -> None:
    update = ensure_child(doc.settings.element, "w:updateFields")
    update.set(qn("w:val"), "true")


def paragraph_has_explicit_page_break(paragraph) -> bool:
    return bool(
        paragraph._p.xpath(".//w:br[@w:type='page']")
        or paragraph._p.xpath("./w:pPr/w:pageBreakBefore")
    )


def ensure_page_break_before(paragraph) -> None:
    previous = paragraph._p.getprevious()
    if previous is not None and previous.tag == qn("w:p"):
        previous_paragraph = Paragraph(previous, paragraph._parent)
        if paragraph_has_explicit_page_break(previous_paragraph):
            return
    page_break = paragraph.insert_paragraph_before()
    page_break.add_run().add_break(WD_BREAK.PAGE)


def ensure_front_matter_page_break(doc: Document) -> bool:
    toc_title = next(
        (
            paragraph
            for paragraph in doc.paragraphs
            if paragraph.style.name == "目录标题" or " ".join(paragraph.text.split()) == "目录"
        ),
        None,
    )
    if toc_title is None:
        return False
    ensure_page_break_before(toc_title)
    return True


def next_numbering_id(doc: Document) -> int:
    numbering = doc.part.numbering_part.element
    values = []
    for item in numbering.findall(qn("w:num")):
        value = item.get(qn("w:numId"))
        if value and value.isdigit():
            values.append(int(value))
    for element in doc.element.body.xpath(".//w:numPr/w:numId"):
        value = element.get(qn("w:val"))
        if value and value.isdigit():
            values.append(int(value))
    for element in doc.styles.element.xpath(".//w:numPr/w:numId"):
        value = element.get(qn("w:val"))
        if value and value.isdigit():
            values.append(int(value))
    return max(values, default=0) + 1


def next_abstract_numbering_id(numbering) -> int:
    values = []
    for item in numbering.findall(qn("w:abstractNum")):
        value = item.get(qn("w:abstractNumId"))
        if value and value.isdigit():
            values.append(int(value))
    return max(values, default=0) + 1


def materialize_legacy_heading_linked_lists(doc: Document) -> None:
    """Preserve legacy list labels as text when Word would merge them with headings."""
    heading_style_ids = {doc.styles[item["name"]].style_id for item in HEADING_STYLES.values()}
    numbering = doc.part.numbering_part.element
    num_to_abstract: dict[str, object] = {}
    for num in numbering.findall(qn("w:num")):
        abstract_ref = num.find(qn("w:abstractNumId"))
        if abstract_ref is None:
            continue
        abstract_id = abstract_ref.get(qn("w:val"))
        abstract = next(
            (item for item in numbering.findall(qn("w:abstractNum")) if item.get(qn("w:abstractNumId")) == abstract_id),
            None,
        )
        if abstract is not None:
            num_to_abstract[num.get(qn("w:numId"))] = abstract

    counters: dict[str, list[int]] = {}
    for paragraph in doc.element.body.findall(qn("w:p")):
        p_pr = paragraph.find(qn("w:pPr"))
        if p_pr is None:
            continue
        p_style = p_pr.find(qn("w:pStyle"))
        if p_style is not None and p_style.get(qn("w:val")) in heading_style_ids:
            continue
        num_pr = p_pr.find(qn("w:numPr"))
        if num_pr is None:
            continue
        num_id_node = num_pr.find(qn("w:numId"))
        if num_id_node is None:
            continue
        num_id = num_id_node.get(qn("w:val"))
        abstract = num_to_abstract.get(num_id)
        if abstract is None:
            continue
        ilvl_node = num_pr.find(qn("w:ilvl"))
        level_number = 0 if ilvl_node is None else int(ilvl_node.get(qn("w:val"), "0"))
        level = next(
            (item for item in abstract.findall(qn("w:lvl")) if item.get(qn("w:ilvl")) == str(level_number)),
            None,
        )
        if level is None:
            continue
        num_fmt = level.find(qn("w:numFmt"))
        lvl_text = level.find(qn("w:lvlText"))
        display = None if lvl_text is None else lvl_text.get(qn("w:val"))
        if num_fmt is None or num_fmt.get(qn("w:val")) != "decimal" or display not in {"%1.", "%1.%2", "%1.%2.%3"}:
            continue

        values = counters.setdefault(num_id, [0, 0, 0])
        if level_number == 0:
            values[0] += 1
            values[1] = 0
            values[2] = 0
        elif level_number == 1:
            if values[0] == 0:
                values[0] = 1
            values[1] += 1
            values[2] = 0
        else:
            if values[0] == 0:
                values[0] = 1
            if values[1] == 0:
                values[1] = 1
            values[2] += 1
        label = display.replace("%1", str(values[0])).replace("%2", str(values[1])).replace("%3", str(values[2]))
        text_nodes = paragraph.xpath(".//w:t")
        if not text_nodes:
            continue
        text_nodes[0].text = f"{label} {text_nodes[0].text or ''}"
        text_nodes[0].set(qn("xml:space"), "preserve")
        p_pr.remove(num_pr)
        source_p_pr = level.find(qn("w:pPr"))
        source_ind = None if source_p_pr is None else source_p_pr.find(qn("w:ind"))
        target_ind = p_pr.find(qn("w:ind"))
        if target_ind is None:
            target_ind = OxmlElement("w:ind")
            p_pr.append(target_ind)
        for attribute in ("w:firstLine", "w:firstLineChars", "w:leftChars", "w:rightChars", "w:hangingChars"):
            qualified = qn(attribute)
            if qualified in target_ind.attrib:
                del target_ind.attrib[qualified]
        if source_ind is not None:
            for attribute in ("w:left", "w:hanging"):
                value = source_ind.get(qn(attribute))
                if value is not None:
                    target_ind.set(qn(attribute), value)


def detach_nonheading_lists_from_heading_styles(doc: Document) -> None:
    """Leave exactly one active heading-linked list definition in the file."""
    heading_style_ids = {doc.styles[item["name"]].style_id for item in HEADING_STYLES.values()}
    numbering = doc.part.numbering_part.element
    ordinary_num_ids: set[str] = set()
    for paragraph in doc.element.body.xpath(".//w:p"):
        p_pr = paragraph.find(qn("w:pPr"))
        if p_pr is None:
            continue
        p_style = p_pr.find(qn("w:pStyle"))
        if p_style is not None and p_style.get(qn("w:val")) in heading_style_ids:
            continue
        num_id = p_pr.find("./" + qn("w:numPr") + "/" + qn("w:numId"))
        if num_id is not None and num_id.get(qn("w:val")):
            ordinary_num_ids.add(num_id.get(qn("w:val")))
    ordinary_abstract_ids: set[str] = set()
    for num in numbering.findall(qn("w:num")):
        if num.get(qn("w:numId")) not in ordinary_num_ids:
            continue
        abstract_ref = num.find(qn("w:abstractNumId"))
        if abstract_ref is not None and abstract_ref.get(qn("w:val")):
            ordinary_abstract_ids.add(abstract_ref.get(qn("w:val")))
    for abstract in numbering.findall(qn("w:abstractNum")):
        abstract_id = abstract.get(qn("w:abstractNumId"))
        for level in abstract.findall(qn("w:lvl")):
            p_style = level.find(qn("w:pStyle"))
            if p_style is not None and p_style.get(qn("w:val")) in heading_style_ids:
                level.remove(p_style)
            num_fmt = level.find(qn("w:numFmt"))
            lvl_text = level.find(qn("w:lvlText"))
            if (
                abstract_id in ordinary_abstract_ids
                and num_fmt is not None
                and num_fmt.get(qn("w:val")) == "decimal"
                and lvl_text is not None
                and lvl_text.get(qn("w:val")) in {"%1.", "%1.%2", "%1.%2.%3"}
            ):
                suffix = level.find(qn("w:suff"))
                if suffix is None:
                    suffix = OxmlElement("w:suff")
                    level.insert(level.index(lvl_text), suffix)
                suffix.set(qn("w:val"), "space")


def add_multilevel_numbering(doc: Document) -> int:
    numbering = doc.part.numbering_part.element
    abstract_id = next_abstract_numbering_id(numbering)
    num_id = next_numbering_id(doc)

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    nsid = OxmlElement("w:nsid")
    nsid.set(qn("w:val"), f"{(0xA17E0000 + abstract_id) & 0xFFFFFFFF:08X}")
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "multilevel")
    template = OxmlElement("w:tmpl")
    template.set(qn("w:val"), f"{(0xC0DE0000 + abstract_id) & 0xFFFFFFFF:08X}")
    abstract.extend([nsid, multi, template])

    for level, text, left, hanging, suffix in NUMBERING_LEVELS:
        item = OxmlElement("w:lvl")
        item.set(qn("w:ilvl"), str(level))
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        num_fmt = OxmlElement("w:numFmt")
        num_fmt.set(qn("w:val"), "decimal")
        lvl_text = OxmlElement("w:lvlText")
        lvl_text.set(qn("w:val"), text)
        suff = OxmlElement("w:suff")
        suff.set(qn("w:val"), suffix)
        lvl_jc = OxmlElement("w:lvlJc")
        lvl_jc.set(qn("w:val"), "left")
        p_style = OxmlElement("w:pStyle")
        p_style.set(qn("w:val"), doc.styles[HEADING_STYLES[level + 1]["name"]].style_id)
        p_pr = OxmlElement("w:pPr")
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), str(left))
        ind.set(qn("w:hanging"), str(hanging))
        p_pr.append(ind)
        # Preserve the OOXML schema order so Word does not rewrite or drop
        # the style linkage when fields are updated and the file is saved.
        item.extend([start, num_fmt, p_style, suff, lvl_text, lvl_jc, p_pr])
        abstract.append(item)

    numbering.append(abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    for level in (0, 1):
        override = OxmlElement("w:lvlOverride")
        override.set(qn("w:ilvl"), str(level))
        start_override = OxmlElement("w:startOverride")
        start_override.set(qn("w:val"), "1")
        override.append(start_override)
        num.append(override)
    numbering.append(num)
    return num_id


def apply_numbering(paragraph, num_id: int, level: int) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    existing = p_pr.find(qn("w:numPr"))
    if existing is not None:
        p_pr.remove(existing)
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), str(level))
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num])
    p_pr.insert(0, num_pr)


def bind_style_numbering(style, num_id: int, level: int | None) -> None:
    """Bind the custom heading style to numbering without paragraph overrides."""
    p_pr = style.element.get_or_add_pPr()
    existing = p_pr.find(qn("w:numPr"))
    if existing is not None:
        p_pr.remove(existing)
    if level is None:
        return
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), str(level))
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num])
    p_pr.insert(0, num_pr)


def clear_direct_run_font(paragraph, remove_bold: bool = False) -> None:
    """Remove direct font/size overrides while preserving semantic run formatting."""
    removable = {qn("w:rFonts"), qn("w:sz"), qn("w:szCs")}
    if remove_bold:
        removable.update({qn("w:b"), qn("w:bCs")})
    for r_pr in paragraph._p.xpath(".//w:rPr"):
        for child in list(r_pr):
            if child.tag in removable:
                r_pr.remove(child)


def clear_direct_heading_format(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    for tag in (
        "w:numPr",
        "w:ind",
        "w:spacing",
        "w:jc",
        "w:keepNext",
        "w:keepLines",
        "w:pageBreakBefore",
        "w:outlineLvl",
    ):
        child = p_pr.find(qn(tag))
        if child is not None:
            p_pr.remove(child)


def apply_reference_heading_style(paragraph, level: int) -> None:
    """Apply the original report's custom quick style with no direct layout."""
    paragraph.style = HEADING_STYLES[level]["name"]
    clear_direct_heading_format(paragraph)
    clear_direct_run_font(paragraph, remove_bold=True)


def apply_heading_layout(paragraph, level: int) -> None:
    """Compatibility wrapper for callers that previously wrote direct indents."""
    apply_reference_heading_style(paragraph, level)


def add_field(paragraph, instruction: str, placeholder: str = ""):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(separate)
    if placeholder:
        run.add_text(placeholder)
    run._r.append(end)
    return run


def clear_paragraph(paragraph) -> None:
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def configure_section_geometry(section) -> None:
    section.page_width = Cm(PAGE_STYLE["page_width_cm"])
    section.page_height = Cm(PAGE_STYLE["page_height_cm"])
    section.left_margin = Cm(PAGE_STYLE["left_margin_cm"])
    section.right_margin = Cm(PAGE_STYLE["right_margin_cm"])
    section.top_margin = Cm(PAGE_STYLE["top_margin_cm"])
    section.bottom_margin = Cm(PAGE_STYLE["bottom_margin_cm"])
    section.header_distance = Cm(PAGE_STYLE["header_distance_cm"])
    section.footer_distance = Cm(PAGE_STYLE["footer_distance_cm"])


def set_page_number_format(section, fmt: str, start: int = 1) -> None:
    sect_pr = section._sectPr
    pg_num_type = ensure_child(sect_pr, "w:pgNumType")
    pg_num_type.set(qn("w:fmt"), fmt)
    pg_num_type.set(qn("w:start"), str(start))


def configure_story_paragraph(paragraph, alignment, size_pt: float = 9.0, color: str = MUTED_GRAY) -> None:
    paragraph.alignment = alignment
    set_zero_indent(paragraph)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    for run in paragraph.runs:
        set_run_font(run, BODY_FONT_EAST_ASIA, size_pt, latin=BODY_FONT_LATIN, color=color)


def configure_cover_section(section) -> None:
    configure_section_geometry(section)
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    for paragraph in section.header.paragraphs + section.footer.paragraphs:
        clear_paragraph(paragraph)


def configure_toc_section(section) -> None:
    configure_section_geometry(section)
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    for paragraph in section.header.paragraphs:
        clear_paragraph(paragraph)
    footer = section.footer.paragraphs[0]
    clear_paragraph(footer)
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    page_run = add_field(footer, " PAGE ", "i")
    set_run_font(page_run, BODY_FONT_EAST_ASIA, 9, latin=BODY_FONT_LATIN, color=MUTED_GRAY)
    configure_story_paragraph(footer, WD_ALIGN_PARAGRAPH.CENTER)
    set_page_number_format(section, "lowerRoman", 1)


def configure_body_section(section) -> None:
    configure_section_geometry(section)
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    header = section.header.paragraphs[0]
    clear_paragraph(header)
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    title_run = add_field(header, ' STYLEREF "报告短题名" ', "报告短题名")
    set_run_font(title_run, BODY_FONT_EAST_ASIA, 9, latin=BODY_FONT_LATIN, color=MUTED_GRAY)
    configure_story_paragraph(header, WD_ALIGN_PARAGRAPH.RIGHT)
    footer = section.footer.paragraphs[0]
    clear_paragraph(footer)
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    page_run = add_field(footer, " PAGE ", "1")
    set_run_font(page_run, BODY_FONT_EAST_ASIA, 9, latin=BODY_FONT_LATIN, color=MUTED_GRAY)
    configure_story_paragraph(footer, WD_ALIGN_PARAGRAPH.CENTER)
    set_page_number_format(section, "decimal", 1)


def add_math(paragraph, latex: str) -> None:
    paragraph._p.append(latex_to_display_omml(latex))


def add_heading(doc: Document, text: str, level: int, num_id: int | None = None):
    paragraph = doc.add_paragraph(text, style=HEADING_STYLES[level]["name"])
    apply_reference_heading_style(paragraph, level)
    return paragraph


def add_caption(doc: Document, label: str, text: str):
    paragraph = doc.add_paragraph(style="图表题注")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_zero_indent(paragraph)
    paragraph.paragraph_format.keep_with_next = label == "表"
    paragraph.add_run(f"{label} ")
    add_field(paragraph, f" SEQ {label} \\* ARABIC ", "1")
    paragraph.add_run(f" {text.rstrip().rstrip('。.')}")
    return paragraph


def mark_header_row(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:tblHeader")) is not None:
        return
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def prevent_row_split(row) -> None:
    """Keep both header and data rows intact across page breaks."""
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is not None:
        return
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_cell_fill(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = ensure_child(tc_pr, "w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)


def apply_status_fills(table) -> None:
    """Color only formal status cells in a column explicitly named 结论."""
    if not table.rows:
        return
    conclusion_columns = [index for index, cell in enumerate(table.rows[0].cells) if clean_cell_text(cell) == "结论"]
    for row in table.rows[1:]:
        for index in conclusion_columns:
            value = clean_cell_text(row.cells[index])
            fill = TABLE_STYLE["status_fills"].get(value)
            if fill:
                set_cell_fill(row.cells[index], fill)


def clean_cell_text(cell) -> str:
    return " ".join(cell.text.split())


def document_body_width_twips(doc: Document) -> int:
    section = doc.sections[0]
    width_emu = section.page_width - section.left_margin - section.right_margin
    return max(1, int(round(width_emu / 635)))


def fit_table_to_width(table, max_width_twips: int | None) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_pr = table._tbl.tblPr
    indent = tbl_pr.find(qn("w:tblInd"))
    if indent is not None:
        tbl_pr.remove(indent)
    if not max_width_twips:
        return
    grid_columns = table._tbl.tblGrid.findall(qn("w:gridCol"))
    widths = [int(item.get(qn("w:w"), "0")) for item in grid_columns]
    total = sum(widths)
    if total <= 0 or total <= max_width_twips:
        return
    scale = max_width_twips / total
    scaled_widths = [max(1, int(round(width * scale))) for width in widths]
    scaled_widths[-1] = max(1, scaled_widths[-1] + max_width_twips - sum(scaled_widths))
    for item, width in zip(grid_columns, scaled_widths):
        item.set(qn("w:w"), str(width))
    seen_cells: set[int] = set()
    for row in table.rows:
        for cell in row.cells:
            marker = id(cell._tc)
            if marker in seen_cells:
                continue
            seen_cells.add(marker)
            tc_width = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tc_width is not None and tc_width.get(qn("w:type"), "dxa") == "dxa":
                width = int(tc_width.get(qn("w:w"), "0"))
                if width > 0:
                    tc_width.set(qn("w:w"), str(max(1, int(round(width * scale)))))
    table.autofit = False
    tbl_width = ensure_child(tbl_pr, "w:tblW")
    tbl_width.set(qn("w:type"), "dxa")
    tbl_width.set(qn("w:w"), str(max_width_twips))


def set_table_column_weights(table, total_width_twips: int, column_weights: list[float] | None) -> None:
    if not column_weights or len(column_weights) != len(table.columns):
        return
    weight_sum = sum(column_weights)
    if weight_sum <= 0:
        return
    widths = [max(1, int(round(total_width_twips * weight / weight_sum))) for weight in column_weights]
    widths[-1] = max(1, widths[-1] + total_width_twips - sum(widths))
    grid_columns = table._tbl.tblGrid.findall(qn("w:gridCol"))
    for grid, width in zip(grid_columns, widths):
        grid.set(qn("w:w"), str(width))
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            tc_width = ensure_child(cell._tc.get_or_add_tcPr(), "w:tcW")
            tc_width.set(qn("w:type"), "dxa")
            tc_width.set(qn("w:w"), str(widths[index]))
    tbl_width = ensure_child(table._tbl.tblPr, "w:tblW")
    tbl_width.set(qn("w:type"), "dxa")
    tbl_width.set(qn("w:w"), str(total_width_twips))


def format_table(
    table,
    long_note_columns: set[int] | None = None,
    max_width_twips: int | None = None,
) -> None:
    # Retain the argument for compatibility with existing callers. The report
    # standard now requires every table cell, including long notes, to center.
    set_table_visual_contract(table)
    fit_table_to_width(table, max_width_twips)
    mark_header_row(table.rows[0])
    for row_index, row in enumerate(table.rows):
        prevent_row_split(row)
        for col_index, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                contains_image = "<w:drawing" in paragraph._p.xml or "<w:pict" in paragraph._p.xml
                paragraph.style = "图片" if contains_image else ("表格表头" if row_index == 0 else "表格文字")
                set_zero_indent(paragraph)
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.0
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    if not contains_image:
                        set_run_font(
                            run,
                            BODY_FONT_EAST_ASIA,
                            9.5,
                            bold=True if row_index == 0 else None,
                        )
    apply_status_fills(table)


def add_table(
    doc: Document,
    headers: list[str],
    rows: list[list[str]],
    long_note_columns: set[int] | None = None,
    column_weights: list[float] | None = None,
):
    table = doc.add_table(rows=1, cols=len(headers))
    for cell, text in zip(table.rows[0].cells, headers):
        cell.text = text
    for row in rows:
        cells = table.add_row().cells
        for cell, text in zip(cells, row):
            cell.text = text
    body_width = document_body_width_twips(doc)
    set_table_column_weights(table, body_width, column_weights)
    format_table(table, long_note_columns, body_width)
    return table


def ensure_style(doc: Document, name: str, base_name: str = "Normal"):
    try:
        return doc.styles[name]
    except KeyError:
        return doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH, builtin=False)


def ensure_builtin_style(doc: Document, name: str):
    try:
        return doc.styles[name]
    except KeyError:
        return doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH, builtin=True)


def ensure_heading_style(doc: Document, level: int):
    style = ensure_style(doc, HEADING_STYLES[level]["name"])
    style.hidden = False
    style.quick_style = True
    style.priority = 10 + level
    return style


def configure_document(doc: Document, add_numbering: bool = True) -> int | None:
    section = doc.sections[0]
    configure_section_geometry(section)

    normal = doc.styles["Normal"]
    set_font(normal, BODY_STYLE["east_asia"], BODY_STYLE["size_pt"], latin=BODY_STYLE["latin"])
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing = Pt(BODY_STYLE["line_spacing_pt"])
    normal.paragraph_format.first_line_indent = Pt(BODY_STYLE["first_line_pt"])
    normal.paragraph_format.space_before = Pt(BODY_STYLE["space_before_pt"])
    normal.paragraph_format.space_after = Pt(BODY_STYLE["space_after_pt"])

    title = ensure_style(doc, "报告题名")
    set_font(
        title,
        REPORT_TITLE_STYLE["east_asia"],
        REPORT_TITLE_STYLE["size_pt"],
        REPORT_TITLE_STYLE["bold"],
        REPORT_TITLE_STYLE["latin"],
        REPORT_TITLE_STYLE["color"],
    )
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_style_zero_indent(title)
    title.paragraph_format.space_before = Pt(76)
    title.paragraph_format.space_after = Pt(10)
    title.paragraph_format.keep_with_next = True

    subtitle = ensure_style(doc, "报告短题名")
    set_font(
        subtitle,
        REPORT_SUBTITLE_STYLE["east_asia"],
        REPORT_SUBTITLE_STYLE["size_pt"],
        REPORT_SUBTITLE_STYLE["bold"],
        REPORT_SUBTITLE_STYLE["latin"],
        REPORT_SUBTITLE_STYLE["color"],
    )
    subtitle.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_style_zero_indent(subtitle)
    subtitle.paragraph_format.space_before = Pt(0)
    subtitle.paragraph_format.space_after = Pt(26)
    subtitle.paragraph_format.keep_with_next = True

    info = ensure_style(doc, "报告信息")
    set_font(info, BODY_FONT_EAST_ASIA, 12)
    info.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_style_zero_indent(info)
    info.paragraph_format.line_spacing = 1.5
    info.paragraph_format.space_after = Pt(0)

    cover_label = ensure_style(doc, "封面小标题")
    set_font(cover_label, "微软雅黑", 10.5, True, "Microsoft YaHei", ACCENT_BLUE)
    cover_label.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_style_zero_indent(cover_label)
    cover_label.paragraph_format.space_before = Pt(10)
    cover_label.paragraph_format.space_after = Pt(4)

    prompt = ensure_style(doc, "模板提示")
    set_font(prompt, BODY_FONT_EAST_ASIA, 10.5, False, BODY_FONT_LATIN, MUTED_GRAY)
    prompt.font.italic = True
    prompt.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_style_zero_indent(prompt)
    prompt.paragraph_format.space_before = Pt(4)
    prompt.paragraph_format.space_after = Pt(4)

    toc_title = ensure_style(doc, "目录标题")
    set_font(toc_title, TOC_TITLE_FONT_EAST_ASIA, 16, True, TOC_TITLE_FONT_LATIN)
    toc_title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_style_zero_indent(toc_title)
    toc_title.paragraph_format.space_after = Pt(10)

    for level in (1, 2, 3):
        configure_toc_style(ensure_builtin_style(doc, f"TOC {level}"), level)

    caption = ensure_style(doc, "图表题注")
    set_font(caption, BODY_FONT_EAST_ASIA, 10.5, False, BODY_FONT_LATIN, ACCENT_BLUE)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_style_zero_indent(caption)
    caption.paragraph_format.space_before = Pt(8)
    caption.paragraph_format.space_after = Pt(4)

    table_text = ensure_style(doc, "表格文字")
    set_font(table_text, BODY_FONT_EAST_ASIA, 9.5)
    table_text.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_style_zero_indent(table_text)
    table_text.paragraph_format.line_spacing = 1.0
    table_text.paragraph_format.space_before = Pt(0)
    table_text.paragraph_format.space_after = Pt(0)

    table_header = ensure_style(doc, "表格表头")
    table_header.base_style = table_text
    set_font(table_header, BODY_FONT_EAST_ASIA, 9.5, True)
    table_header.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_style_zero_indent(table_header)
    table_header.paragraph_format.line_spacing = 1.0
    table_header.paragraph_format.space_before = Pt(0)
    table_header.paragraph_format.space_after = Pt(0)

    try:
        report_table_style = doc.styles[TABLE_STYLE["name"]]
    except KeyError:
        report_table_style = doc.styles.add_style(TABLE_STYLE["name"], WD_STYLE_TYPE.TABLE, builtin=False)
    set_font(
        report_table_style,
        TABLE_STYLE["east_asia"],
        TABLE_STYLE["font_size_pt"],
        latin=TABLE_STYLE["latin"],
    )

    image = ensure_style(doc, "图片")
    set_font(image, BODY_FONT_EAST_ASIA, 12)
    image.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_style_zero_indent(image)
    image.paragraph_format.space_before = Pt(0)
    image.paragraph_format.space_after = Pt(0)
    image.paragraph_format.keep_with_next = True

    formula = ensure_style(doc, "公式")
    set_font(formula, BODY_FONT_EAST_ASIA, 12)
    formula.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_style_zero_indent(formula)
    formula.paragraph_format.space_before = Pt(6)
    formula.paragraph_format.space_after = Pt(6)

    for level in (1, 2, 3):
        contract = HEADING_STYLES[level]
        style = ensure_heading_style(doc, level)
        style.base_style = normal
        set_font(
            style,
            contract["east_asia"],
            contract["size_pt"],
            contract["bold"],
            contract["latin"],
            contract["color"],
        )
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        # Preserve the original style behavior: list levels provide the
        # 25.5 pt effective left indent for levels 1-2.
        set_heading_style_indent(style, level)
        style.paragraph_format.space_before = Pt(contract["space_before_pt"])
        style.paragraph_format.space_after = Pt(contract["space_after_pt"])
        style.paragraph_format.line_spacing = contract["line_spacing"]
        style.paragraph_format.keep_with_next = contract["keep_with_next"]
        style.paragraph_format.keep_together = contract["keep_together"]
        set_outline_level(style, level - 1)

    set_update_fields(doc)
    if not add_numbering:
        return None
    materialize_legacy_heading_linked_lists(doc)
    detach_nonheading_lists_from_heading_styles(doc)
    num_id = add_multilevel_numbering(doc)
    bind_style_numbering(doc.styles[HEADING_STYLES[1]["name"]], num_id, 0)
    bind_style_numbering(doc.styles[HEADING_STYLES[2]["name"]], num_id, 1)
    bind_style_numbering(doc.styles[HEADING_STYLES[3]["name"]], num_id, 2)
    return num_id


def build_template(output: Path) -> None:
    doc = Document()
    num_id = configure_document(doc)
    assert num_id is not None

    cover_section = doc.sections[0]
    configure_cover_section(cover_section)
    doc.add_paragraph("超稳激光电控箱测试报告", style="报告题名")
    doc.add_paragraph("板卡或测试对象简称", style="报告短题名")
    add_table(
        doc,
        ["项目", "内容"],
        [
            ["版本", "V1.0"],
            ["测试时间", "待填写"],
            ["报告日期", "待填写"],
            ["测试地点", "待填写"],
            ["测试人员", "待填写"],
            ["编写 / 审核 / 批准", "待填写 / 未记录 / 未记录"],
        ],
        {1},
        [0.30, 0.70],
    )
    doc.add_paragraph("修订记录", style="封面小标题")
    add_table(
        doc,
        ["版本", "日期", "修订内容"],
        [["V1.0", "待填写", "首次编制"]],
        {2},
        [0.18, 0.24, 0.58],
    )

    toc_section = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_toc_section(toc_section)
    doc.add_paragraph("目录", style="目录标题")
    toc = doc.add_paragraph()
    set_zero_indent(toc)
    add_field(toc, ' TOC \\o "1-3" \\h \\z \\u ', "右键更新域或按 F9 更新目录")

    body_section = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_body_section(body_section)

    add_heading(doc, "测试目的与结论摘要", 1, num_id)
    doc.add_paragraph("说明测试对象、测试范围、已完成项目、主要结果、未完成项目、总体结论和决定性限制。")
    add_caption(doc, "表", "测试完成度与结论摘要")
    add_table(
        doc,
        ["对象", "已完成项目", "主要异常或未测试项", "总体结论", "关键限制"],
        [["待填写", "待填写", "待填写", "未测试", "待填写"]],
        {1, 2, 4},
        [0.15, 0.23, 0.25, 0.15, 0.22],
    )

    add_heading(doc, "测试对象、接口与公共条件", 1, num_id)
    add_heading(doc, "被测对象及接口映射", 2, num_id)
    add_table(
        doc,
        ["对象", "器件或功能", "接口", "说明"],
        [["待填写", "待填写", "待填写", "待填写"]],
        {1, 2, 3},
        [0.18, 0.26, 0.20, 0.36],
    )
    add_heading(doc, "端接、负载、时钟与软件状态", 2, num_id)
    doc.add_paragraph("记录实际端接、负载、时钟、固件、软件版本和关键配置；未记录信息明确写为“未记录”。")
    add_heading(doc, "公共采集条件", 2, num_id)
    add_table(
        doc,
        ["项目", "配置或条件", "适用范围", "记录状态"],
        [["采样率 / 带宽 / 时长", "待填写", "待填写", "未记录"]],
        {1, 2, 3},
        [0.23, 0.34, 0.25, 0.18],
    )

    add_heading(doc, "测试仪器、系统连接与数据处理方法", 1, num_id)
    add_heading(doc, "测试仪器及主要配置", 2, num_id)
    add_table(
        doc,
        ["仪器或工具", "型号或编号", "主要设置", "校准状态"],
        [["待填写", "未记录", "待填写", "未记录"]],
        {1, 2, 3},
        [0.22, 0.24, 0.36, 0.18],
    )
    add_heading(doc, "测试系统与实际连接", 2, num_id)
    doc.add_paragraph("给出与本次测试一致的信号路径、参考面、端接和测量节点；未经过的器件不得画入测试链路。")
    doc.add_paragraph("[在此插入清晰、可编辑或高分辨率的测试系统图]", style="模板提示")
    add_caption(doc, "图", "测试系统与实际连接")
    add_heading(doc, "公共数据处理方法", 2, num_id)
    doc.add_paragraph("说明符号位解析、去均值、拟合或频谱估计方法、窗函数、重叠率、频率分辨率、单位和排除规则。")

    add_heading(doc, "分器件或分链路测试结果", 1, num_id)
    add_heading(doc, "被测器件或链路", 2, num_id)
    add_heading(doc, "测试指标与条件", 3, num_id)
    doc.add_paragraph("明确指标值、规定工况、实际连接、输入、负载、时钟、采样率、时长和仪器设置。")
    add_heading(doc, "测试方法", 3, num_id)
    doc.add_paragraph("按实际操作顺序说明测试步骤和工况变化，不用通用理论替代实际流程。")
    add_heading(doc, "数据处理", 3, num_id)
    doc.add_paragraph("给出公式、参考面、刻度来源、频率分辨率、筛点依据、量纲和可复算参数。")
    add_heading(doc, "结果、分析与判定", 3, num_id)
    add_caption(doc, "表", "被测器件或链路结果")
    add_table(
        doc,
        ["通道", "测试条件", "实测或计算值", "指标要求", "结论", "数据编号"],
        [["待填写", "待填写", "待填写", "待填写", "未测试", "待填写"]],
        {1, 2, 3, 5},
        [0.12, 0.22, 0.25, 0.19, 0.12, 0.10],
    )
    doc.add_paragraph("[在此插入能够直接支撑本项结论的结果图]", style="模板提示")
    add_caption(doc, "图", "被测对象在规定条件下的测试结果")
    doc.add_paragraph("定量分析：填写精确数值、趋势、异常点、杂散或残差，以及测试覆盖限制。")
    doc.add_paragraph("本项结论：使用“满足”“不满足”“暂不能判定”或“未测试”，并说明必要限定条件。")

    add_heading(doc, "异常现象、证据限制与未测试项", 1, num_id)
    add_table(
        doc,
        ["项目或现象", "已有数据", "对结论的影响", "可能原因", "确认或补测方法"],
        [["待填写", "待填写", "待填写", "待填写", "待填写"]],
        {0, 1, 2, 3, 4},
        [0.18, 0.22, 0.20, 0.18, 0.22],
    )

    add_heading(doc, "指标实测对比总表与后续计划", 1, num_id)
    add_heading(doc, "指标实测对比", 2, num_id)
    add_caption(doc, "表", "全部适用指标与实测结果对比")
    add_table(
        doc,
        ["对象", "测试项目", "指标要求", "规定工况", "实测或计算值", "数据覆盖", "结论", "主要依据"],
        [["待填写", "待填写", "待填写", "待填写", "待填写", "无数据", "未测试", "待填写"]],
        set(range(8)),
        [0.10, 0.15, 0.15, 0.14, 0.17, 0.10, 0.09, 0.10],
    )
    add_heading(doc, "后续测试计划", 2, num_id)
    add_table(
        doc,
        ["项目", "当前缺口", "后续测试", "必须保存的数据", "优先级"],
        [["待填写", "待填写", "待填写", "待填写", "待填写"]],
        {0, 1, 2, 3},
        [0.16, 0.22, 0.28, 0.24, 0.10],
    )

    add_heading(doc, "附录", 1, num_id)
    add_heading(doc, "计算公式与处理参数", 2, num_id)
    doc.add_paragraph("列出正文首次使用公式的完整定义、量纲、计算区间和处理参数。")
    formula = doc.add_paragraph(style="公式")
    add_math(formula, r"A_v(f)=\sqrt{S_v(f)}")
    add_heading(doc, "仪器设置与原始记录", 2, num_id)
    doc.add_paragraph("保存仪器截图、原始截图和必要的配置记录；低清晰度拼图只作为原始证据，不作为正文定量结果图。")
    add_heading(doc, "数据与脚本索引", 2, num_id)
    add_table(
        doc,
        ["数据编号", "内容", "测试日期", "原始文件", "处理脚本", "版本"],
        [["待填写", "待填写", "待填写", "待填写", "待填写", "待填写"]],
        {1, 3, 4, 5},
        [0.12, 0.19, 0.14, 0.25, 0.20, 0.10],
    )

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def main() -> int:
    reconfigure = getattr(sys.stdout, "reconfigure", None)
    if reconfigure is not None:
        try:
            reconfigure(encoding="utf-8", errors="replace")
        except (OSError, ValueError):
            pass
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    build_template(args.output.resolve())
    print(args.output.resolve())
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

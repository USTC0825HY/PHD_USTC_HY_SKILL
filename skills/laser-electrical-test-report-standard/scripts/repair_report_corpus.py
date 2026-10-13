#!/usr/bin/env python3
"""Repair professional math, reference headings, and centered tables in formal reports."""

from __future__ import annotations

import argparse
import os
import re
import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

from _portable_profile import profile_path
from create_standard_template import (
    apply_reference_heading_style,
    clear_direct_run_font,
    configure_document,
    document_body_width_twips,
    ensure_front_matter_page_break,
    format_table,
)
from latex_to_omml import latex_to_display_omml, latex_to_omml
from style_contract import HEADING_STYLES


DEFAULT_ROOT = (
    profile_path("obsidian_vault_root", Path.home() / "Obsidian")
    / "01_科研主系统"
    / "04_实验测试与证据"
    / "03_超稳激光电控箱指标测试"
    / "LASER 测试报告"
    / "202607"
    / "01_正式报告"
)
RF_REPORT = "04_频率综合板/04_射频参考_100MHz时钟_DDS输出测试.docx"


def normalized(text: str) -> str:
    return " ".join(text.split())


LEGACY_DISPLAY_MATH = {
    normalized("A = sqrt(a²+b²),  xpp = 2A"): r"A=\sqrt{a^2+b^2},\qquad x_{\mathrm{pp}}=2A",
    normalized("ASD(f) = sqrt(PSD(f))"): r"\operatorname{ASD}(f)=\sqrt{\operatorname{PSD}(f)}",
    normalized("Vrms(f1,f2) = sqrt(integral PSD(f) df from f1 to f2)"): (
        r"V_{\mathrm{rms}}(f_1,f_2)=\sqrt{\int_{f_1}^{f_2}\operatorname{PSD}(f)\,\mathrm{d}f}"
    ),
    normalized("S_ADC,in(f) = L_ADC²/(|H_FPGA(f)|² K_out²) · max[S_Vout,total(f)-S_Vout,DAC(f), 0]"): (
        r"S_{\mathrm{ADC,in}}(f)="
        r"\frac{L_{\mathrm{ADC}}^2}{\lvert H_{\mathrm{FPGA}}(f)\rvert^2K_{\mathrm{out}}^2}"
        r"\max\left[S_{V_{\mathrm{out,total}}}(f)-S_{V_{\mathrm{out,DAC}}}(f),0\right]"
    ),
    normalized("sigma_y²(tau) = 1/2 · mean[(ybar[k+1]-ybar[k])²]"): (
        r"\sigma_y^2(\tau)=\frac{1}{2}\operatorname{mean}"
        r"\left[\left(\bar y_{k+1}-\bar y_k\right)^2\right]"
    ),
    normalized("tau_d = -(1/2π) · dphi(f)/df"): r"\tau_d=-\frac{1}{2\pi}\frac{\mathrm{d}\phi(f)}{\mathrm{d}f}",
    normalized("D0 = Verr,amp/Dnu,amp = 3.379 Code/Hz × 0.0428 mV/Code ≈ 1.45×10^-4 V/Hz"): (
        r"D_0=\frac{V_{\mathrm{err,amp}}}{\Delta\nu_{\mathrm{amp}}}="
        r"3.379\ \mathrm{Code/Hz}\times0.0428\ \mathrm{mV/Code}"
        r"\approx1.45\times10^{-4}\ \mathrm{V/Hz}"
    ),
    normalized("L(f)=D(f)G(f)K(f)exp(-j2πf tau_d),  T(f)=L(f)/(1+L(f))"): (
        r"L(f)=D(f)G(f)K(f)\exp(-j2\pi f\tau_d),\qquad T(f)=\frac{L(f)}{1+L(f)}"
    ),
    normalized("sigma_y²(tau)=integral Sy(f) · 2sin^4(πf tau)/(πf tau)² df"): (
        r"\sigma_y^2(\tau)=\int S_y(f)\frac{2\sin^4(\pi f\tau)}{(\pi f\tau)^2}\,\mathrm{d}f"
    ),
    normalized("f_alias = | f_0 - k f_s |,    k = 0, ±1, ±2, …"): (
        r"f_{\mathrm{alias}}=\lvert f_0-kf_s\rvert,\qquad k=0,\pm1,\pm2,\ldots"
    ),
    normalized("f_alias = | Δf | = | f_0 - f_s |"): (
        r"f_{\mathrm{alias}}=\lvert\Delta f\rvert=\lvert f_0-f_s\rvert"
    ),
    normalized("T_alias = 1 / | f_0 - f_s |"): r"T_{\mathrm{alias}}=\frac{1}{\lvert f_0-f_s\rvert}",
    normalized("f_0 = f_s ± Δf = 100000.003 ± 0.0412371  Hz"): (
        r"f_0=f_s\pm\Delta f=100000.003\pm0.0412371\ \mathrm{Hz}"
    ),
    normalized("f_alias,k = | k f_0 - k f_s | = k | f_0 - f_s | = k Δf"): (
        r"f_{\mathrm{alias},k}=\lvert kf_0-kf_s\rvert=k\lvert f_0-f_s\rvert=k\Delta f"
    ),
    normalized("A_v(f) = √( S_v(f) )"): r"A_v(f)=\sqrt{S_v(f)}",
    normalized("V_rms = √( ∫[f_1 to f_2] S_v(f) df )"): (
        r"V_{\mathrm{rms}}=\sqrt{\int_{f_1}^{f_2}S_v(f)\,\mathrm{d}f}"
    ),
    normalized("A_v(f)=sqrt(S_v(f))"): r"A_v(f)=\sqrt{S_v(f)}",
    normalized("V_rms=sqrt(integral S_v(f) df)"): r"V_{\mathrm{rms}}=\sqrt{\int S_v(f)\,\mathrm{d}f}",
}


INLINE_MATH = {
    "f_s=100000.003 S/s": r"f_s=100000.003\ \mathrm{S/s}",
    "T_alias=24.25 s": r"T_{\mathrm{alias}}=24.25\ \mathrm{s}",
    "f_0=f_s+1/T_alias": r"f_0=f_s+\frac{1}{T_{\mathrm{alias}}}",
    "f_0=f_s+0.0412371 Hz": r"f_0=f_s+0.0412371\ \mathrm{Hz}",
    "T=1/|f_0-kf_s|": r"T=\frac{1}{\lvert f_0-kf_s\rvert}",
    "f_0=f_s+Δf": r"f_0=f_s+\Delta f",
    "f_0=f_s-Δf": r"f_0=f_s-\Delta f",
    "x[n]=x(nT_s)": r"x[n]=x(nT_s)",
    "t=nT_s": r"t=nT_s",
    "G/H_FPGA(f)": r"G/H_{\mathrm{FPGA}}(f)",
    "H_FPGA(f)": r"H_{\mathrm{FPGA}}(f)",
    "1/T_alias": r"\frac{1}{T_{\mathrm{alias}}}",
    "L_ADC²": r"L_{\mathrm{ADC}}^2",
    "K_out²": r"K_{\mathrm{out}}^2",
    "f_-3dB": r"f_{-3\mathrm{dB}}",
    "T_alias": r"T_{\mathrm{alias}}",
    "f_alias": r"f_{\mathrm{alias}}",
    "K_out": r"K_{\mathrm{out}}",
    "L_ADC": r"L_{\mathrm{ADC}}",
    "V_rms": r"V_{\mathrm{rms}}",
    "A_v(f)": r"A_v(f)",
    "S_v(f)": r"S_v(f)",
    "T_s": r"T_s",
    "f_0": r"f_0",
    "f_s": r"f_s",
}


INLINE_PATTERN = re.compile(
    "|".join(
        rf"(?<![A-Za-z0-9_]){re.escape(item)}(?![A-Za-z0-9_])"
        for item in sorted(INLINE_MATH, key=len, reverse=True)
    )
)
STRUCTURAL_MATH_TAGS = (
    "m:sSub",
    "m:sSup",
    "m:sSubSup",
    "m:f",
    "m:rad",
    "m:nary",
    "m:d",
    "m:func",
)
RUN_TEXT_REPLACEMENTS = {
    "伺服噪声没有办法直接测量到": "伺服噪声无法直接测量",
    "，而是采用环外误差信号进行计算。": "，采用环外误差信号计算。",
    "统一指标口径": "统一指标判定依据",
    "证据与状态口径": "证据与状态说明",
    "已发现的口径冲突": "已发现的数据处理方法冲突",
    "该结果不能替代1 MHz载波相噪、带载范围和驱动电流测试。": (
        "1 MHz载波相噪、带载范围和驱动电流仍需专项测试。"
    ),
    "约400 kHz的伺服峰值仅是闭环谱峰位置，不能替代交越频率。": (
        "约400 kHz的伺服峰值仅表示闭环谱峰位置，交越频率需依据开环增益0 dB条件确定。"
    ),
    "不能替代数字锁定板<250 ns整链路指标": "不作为数字锁定板<250 ns整链路指标的判定结果",
    "判定边界": "判定条件",
    "处于限值边界": "接近限值",
    "数据与图片来源": "数据来源",
    "数据和图片来源": "数据来源",
}


def heading_level(paragraph) -> int | None:
    style_id = (paragraph.style.style_id or "").lower().replace(" ", "")
    style_name = paragraph.style.name.lower().replace(" ", "")
    for level in (1, 2, 3):
        if (
            style_id in {f"heading{level}", str(level)}
            or style_name in {f"heading{level}", f"标题{level}", HEADING_STYLES[level]["name"]}
        ):
            return level
    return None


def is_plain_body_paragraph(paragraph) -> bool:
    style_id = (paragraph.style.style_id or "").lower()
    if style_id != "normal" and paragraph.style.name not in {"Normal", "正文"}:
        return False
    return not paragraph._p.xpath(
        ".//w:fldChar | .//w:instrText | .//w:drawing | .//w:pict | .//w:sdt | .//m:oMath"
    )


def formula_text(paragraph) -> str:
    return normalized("".join(node.text or "" for node in paragraph._p.xpath(".//m:t")))


def replace_legacy_display_math(paragraph) -> bool:
    if not paragraph._p.xpath(".//m:oMath"):
        return False
    latex = LEGACY_DISPLAY_MATH.get(formula_text(paragraph))
    if latex is None:
        return False
    paragraph._p.clear_content()
    paragraph.style = "公式"
    paragraph._p.append(latex_to_display_omml(latex))
    return True


def unsafe_inline_paragraph(paragraph) -> bool:
    if heading_level(paragraph) is not None:
        return True
    style = paragraph.style.name.lower()
    if style.startswith("toc") or paragraph.style.name in {"报告题名", "目录标题", "图表题注", "图片", "公式"}:
        return True
    return bool(
        paragraph._p.xpath(
            ".//w:fldChar | .//w:instrText | .//w:hyperlink | .//w:drawing | .//w:pict | .//w:sdt | .//m:oMath"
        )
    )


def replace_inline_math(paragraph) -> int:
    if unsafe_inline_paragraph(paragraph):
        return 0
    text = paragraph.text
    matches = list(INLINE_PATTERN.finditer(text))
    if not matches:
        return 0
    paragraph._p.clear_content()
    cursor = 0
    for match in matches:
        if match.start() > cursor:
            paragraph.add_run(text[cursor : match.start()])
        source = match.group(0)
        paragraph._p.append(latex_to_omml(INLINE_MATH[source]))
        cursor = match.end()
    if cursor < len(text):
        paragraph.add_run(text[cursor:])
    return len(matches)


def iter_table_paragraphs(doc: Document):
    seen: set[int] = set()
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                marker = id(cell._tc)
                if marker in seen:
                    continue
                seen.add(marker)
                yield from cell.paragraphs


def replace_text_across_runs(paragraph, source: str, replacement: str) -> int:
    """Replace reviewed prose even when Word split it across several runs."""
    nodes = [node for node in paragraph._p.xpath(".//w:t") if node.text]
    if not nodes:
        return 0
    combined = "".join(node.text or "" for node in nodes)
    starts: list[int] = []
    cursor = 0
    while True:
        found = combined.find(source, cursor)
        if found < 0:
            break
        starts.append(found)
        cursor = found + len(source)
    if not starts:
        return 0

    def locate(absolute_offset: int) -> tuple[int, int]:
        running = 0
        for index, node in enumerate(nodes):
            length = len(node.text or "")
            if running <= absolute_offset < running + length:
                return index, absolute_offset - running
            running += length
        raise IndexError(absolute_offset)

    for start in reversed(starts):
        end = start + len(source)
        start_index, start_offset = locate(start)
        end_index, end_offset = locate(end - 1)
        if start_index == end_index:
            text = nodes[start_index].text or ""
            nodes[start_index].text = text[:start_offset] + replacement + text[end_offset + 1 :]
            continue
        start_text = nodes[start_index].text or ""
        end_text = nodes[end_index].text or ""
        nodes[start_index].text = start_text[:start_offset] + replacement
        for index in range(start_index + 1, end_index):
            nodes[index].text = ""
        nodes[end_index].text = end_text[end_offset + 1 :]
    return len(starts)


def assert_no_unmapped_pseudo_math(doc: Document, path: Path) -> None:
    for paragraph in list(doc.paragraphs) + list(iter_table_paragraphs(doc)):
        for math in paragraph._p.xpath(".//m:oMath"):
            text = "".join(node.text or "" for node in math.iter(qn("m:t")))
            structured = any(math.find(".//" + qn(tag)) is not None for tag in STRUCTURAL_MATH_TAGS)
            if not structured and re.search(r"_|sqrt|integral|∫\[|\\frac|\\sqrt", text):
                raise ValueError(f"Unmapped pseudo-math remains in {path}: {text}")


def repair(path: Path) -> tuple[int, int]:
    doc = Document(path)
    configure_document(doc, add_numbering=True)
    ensure_front_matter_page_break(doc)
    first_nonempty = next((paragraph for paragraph in doc.paragraphs if normalized(paragraph.text)), None)
    if first_nonempty is not None:
        first_nonempty.style = "报告题名"
        clear_direct_run_font(first_nonempty, remove_bold=True)
    for paragraph in doc.paragraphs:
        level = heading_level(paragraph)
        if level is not None:
            if re.match(r"^\.\s+", paragraph.text):
                paragraph.text = re.sub(r"^\.\s+", "", paragraph.text, count=1)
            apply_reference_heading_style(paragraph, level)
            continue
        if paragraph is not first_nonempty and is_plain_body_paragraph(paragraph):
            clear_direct_run_font(paragraph, remove_bold=False)

    display_count = 0
    inline_count = 0
    all_paragraphs = list(doc.paragraphs) + list(iter_table_paragraphs(doc))
    for paragraph in all_paragraphs:
        for source, replacement in RUN_TEXT_REPLACEMENTS.items():
            replace_text_across_runs(paragraph, source, replacement)
        if replace_legacy_display_math(paragraph):
            display_count += 1
        else:
            inline_count += replace_inline_math(paragraph)

    max_width = document_body_width_twips(doc)
    for table in doc.tables:
        format_table(table, max_width_twips=max_width)
    assert_no_unmapped_pseudo_math(doc, path)

    temporary = path.with_name(path.stem + ".repairing.docx")
    doc.save(temporary)
    os.replace(temporary, path)
    return display_count, inline_count


def corpus_files(root: Path) -> list[Path]:
    reports = sorted(path for path in root.rglob("*_规范修订版.docx") if not path.name.startswith("~$"))
    rf = root / Path(RF_REPORT)
    if rf.is_file():
        reports.append(rf)
    return reports


def main() -> int:
    sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", type=Path, default=DEFAULT_ROOT)
    parser.add_argument("docx", nargs="*", type=Path)
    args = parser.parse_args()
    paths = [path.resolve() for path in args.docx] or corpus_files(args.root.resolve())
    if not paths:
        parser.error("No report DOCX files were found.")
    for path in paths:
        if not path.is_file():
            parser.error(f"DOCX not found: {path}")
        display_count, inline_count = repair(path)
        print(f"{path}\tdisplay_math={display_count}\tinline_math={inline_count}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

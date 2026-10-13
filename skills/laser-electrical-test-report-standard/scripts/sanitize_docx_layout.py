#!/usr/bin/env python3
"""Reapply table, image, caption, and font layout after Word field updates."""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt

from create_standard_template import (
    apply_status_fills,
    BODY_FONT_EAST_ASIA,
    document_body_width_twips,
    ensure_child,
    fit_table_to_width,
    format_table,
    mark_header_row,
    prevent_row_split,
    set_run_font,
    set_table_visual_contract,
    set_zero_indent,
)
from style_contract import TABLE_STYLE


def contains_image(paragraph) -> bool:
    return "<w:drawing" in paragraph._p.xml or "<w:pict" in paragraph._p.xml


def contains_figure_seq(paragraph) -> bool:
    instruction = " ".join(node.text or "" for node in paragraph._p.xpath(".//w:instrText"))
    return "SEQ 图" in " ".join(instruction.split())


def has_style(doc, name: str) -> bool:
    try:
        doc.styles[name]
        return True
    except KeyError:
        return False


def set_run_script_fonts(run, east_asia: str) -> None:
    r_pr = run._r.get_or_add_rPr()
    fonts = ensure_child(r_pr, "w:rFonts")
    fonts.set(qn("w:ascii"), "Times New Roman")
    fonts.set(qn("w:hAnsi"), "Times New Roman")
    fonts.set(qn("w:cs"), "Times New Roman")
    fonts.set(qn("w:eastAsia"), east_asia)


def iter_story_tables(story):
    def visit(table):
        yield table
        for row in table.rows:
            for cell in row.cells:
                for nested in cell.tables:
                    yield from visit(nested)

    for table in story.tables:
        yield from visit(table)


def iter_story_paragraphs(story):
    for paragraph in story.paragraphs:
        yield paragraph
    for table in iter_story_tables(story):
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    # Do not de-duplicate with Python object ids. python-docx can
                    # create short-lived wrappers whose ids are reused, causing
                    # unrelated cell paragraphs to be skipped. Formatting the
                    # same merged-cell paragraph twice is harmless.
                    yield paragraph


def story_parts(doc):
    yield doc
    seen: set[str] = set()
    for section in doc.sections:
        for story in (section.header, section.first_page_header, section.even_page_header, section.footer, section.first_page_footer, section.even_page_footer):
            part_name = str(story.part.partname)
            if part_name not in seen:
                seen.add(part_name)
                yield story


def is_layout_table(table) -> bool:
    xml = table._tbl.xml
    return "<w:drawing" in xml or "<w:pict" in xml or "SEQ 图" in xml


def sanitize_layout_table(table, doc) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    indent = table._tbl.tblPr.find(qn("w:tblInd"))
    if indent is not None:
        indent.set(qn("w:w"), "0")
        indent.set(qn("w:type"), "dxa")
    for row in table.rows:
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                set_zero_indent(paragraph)
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                if contains_image(paragraph):
                    if has_style(doc, "图片"):
                        paragraph.style = "图片"
                    paragraph.paragraph_format.keep_with_next = True
                elif contains_figure_seq(paragraph) or paragraph.style.name == "图表题注":
                    if has_style(doc, "图表题注"):
                        paragraph.style = "图表题注"


def cell_fill(cell) -> str | None:
    shading = cell._tc.get_or_add_tcPr().find(qn("w:shd"))
    return None if shading is None else shading.get(qn("w:fill"))


def set_run_color(run, color: str) -> None:
    r_pr = run._r.get_or_add_rPr()
    color_node = ensure_child(r_pr, "w:color")
    color_node.set(qn("w:val"), color)


def preserve_user_table(table) -> None:
    """Repair refresh-sensitive indents without restyling a hand-edited table."""
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                set_zero_indent(paragraph)
    for cell in table.rows[0].cells:
        if (cell_fill(cell) or "").upper() != TABLE_STYLE["header_fill"]:
            continue
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                if run.text.strip():
                    run.bold = True
                    set_run_color(run, TABLE_STYLE["header_font_color"])


def format_legacy_table(table, max_width_twips: int) -> None:
    """Apply the project table contract when project quick styles are absent."""
    set_table_visual_contract(table)
    fit_table_to_width(table, max_width_twips)
    mark_header_row(table.rows[0])
    for row_index, row in enumerate(table.rows):
        prevent_row_split(row)
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                set_zero_indent(paragraph)
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.0
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    if contains_image(paragraph):
                        continue
                    set_run_font(
                        run,
                        BODY_FONT_EAST_ASIA,
                        9.5,
                        bold=True if row_index == 0 else None,
                        color=TABLE_STYLE["header_font_color"] if row_index == 0 else "000000",
                    )
    apply_status_fills(table)


def sanitize(path: Path, *, restyle_tables: bool = False, normalize_fonts: bool = False) -> None:
    doc = Document(path)
    max_table_width = document_body_width_twips(doc)
    has_standard_table_styles = all(has_style(doc, name) for name in ("表格表头", "表格文字", "图片", "图表题注"))
    for story in story_parts(doc):
        for table in iter_story_tables(story):
            if is_layout_table(table):
                sanitize_layout_table(table, doc)
            elif restyle_tables:
                if has_standard_table_styles:
                    format_table(table, max_width_twips=max_table_width)
                else:
                    format_legacy_table(table, max_table_width)
            else:
                preserve_user_table(table)
        for paragraph in iter_story_paragraphs(story):
            if contains_image(paragraph):
                if has_style(doc, "图片"):
                    paragraph.style = "图片"
                set_zero_indent(paragraph)
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.keep_with_next = True
            elif contains_figure_seq(paragraph) or paragraph.style.name == "图表题注":
                if has_style(doc, "图表题注"):
                    paragraph.style = "图表题注"
                set_zero_indent(paragraph)
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if normalize_fonts:
                for run in paragraph.runs:
                    if "<m:oMath" in run._r.xml:
                        continue
                    east_asia = "微软雅黑" if paragraph.style.name in {"报告题名", "报告短题名", "目录标题", "一级标题", "二级标题", "三级标题"} else "宋体"
                    set_run_script_fonts(run, east_asia)
    doc.save(path)


def main() -> int:
    sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("docx", type=Path)
    parser.add_argument("--restyle-tables", action="store_true", help="Apply the complete project table contract. Never use for a user-edited report unless explicitly authorized.")
    parser.add_argument("--normalize-fonts", action="store_true", help="Normalize all non-math run fonts. Default preserves user direct formatting.")
    args = parser.parse_args()
    path = args.docx.resolve()
    if not path.is_file():
        parser.error(f"DOCX not found: {path}")
    sanitize(path, restyle_tables=args.restyle_tables, normalize_fonts=args.normalize_fonts)
    print(path)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

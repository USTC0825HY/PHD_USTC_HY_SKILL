#!/usr/bin/env python3
"""Repair a CW operation manual's title, heading ladder, and SOP list linkage.

The input document is never overwritten.  The resulting revision has:
* `手册题名` for the unnumbered document title;
* `一级标题` / `二级标题` / `三级标题` linked to 第1章 / 1.1 / 1.1.1;
* `操作步骤` linked to independent 1. lists that restart per contiguous group.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import tempfile
import zipfile
from copy import deepcopy
from pathlib import Path
from typing import Iterable, Sequence

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


HEADING_STYLE_NAMES = (
    "Heading 1", "Heading 2", "Heading 3", "一级标题", "二级标题", "三级标题",
)
HEADING_TARGETS = (
    ("一级标题", "ManualHeading1", 0, 15.5, 12, 6),
    ("二级标题", "ManualHeading2", 1, 12.5, 9, 3),
    ("三级标题", "ManualHeading3", 2, 11.5, 6, 3),
)


def child(parent, tag: str, **attrs: str):
    element = OxmlElement(tag)
    for key, value in attrs.items():
        element.set(qn(key), str(value))
    parent.append(element)
    return element


def replace_child(parent, tag: str, replacement):
    existing = parent.find(qn(tag))
    if existing is not None:
        parent.remove(existing)
    parent.append(replacement)


def set_rfonts(rpr):
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr, value in {
        "w:ascii": "Times New Roman",
        "w:hAnsi": "Times New Roman",
        "w:cs": "Times New Roman",
    }.items():
        rfonts.set(qn(attr), value)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:csTheme"):
        rfonts.attrib.pop(qn(attr), None)


def set_style_fonts(style, size_pt: float, bold: bool):
    style.font.name = "Times New Roman"
    style.font.size = Pt(size_pt)
    style.font.bold = bold
    style.font.color.rgb = None
    rpr = style._element.get_or_add_rPr()
    set_rfonts(rpr)
    if bold:
        child(rpr, "w:b") if rpr.find(qn("w:b")) is None else None
        child(rpr, "w:bCs") if rpr.find(qn("w:bCs")) is None else None


def ensure_style(doc, name: str, style_id: str):
    try:
        style = doc.styles[name]
    except KeyError:
        style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style._element.set(qn("w:styleId"), style_id)
    style.quick_style = True
    style.hidden = False
    return style


def clear_numpr(paragraph):
    ppr = paragraph._p.pPr
    if ppr is not None:
        numpr = ppr.find(qn("w:numPr"))
        if numpr is not None:
            ppr.remove(numpr)


def has_direct_numpr(paragraph) -> bool:
    ppr = paragraph._p.pPr
    return ppr is not None and ppr.find(qn("w:numPr")) is not None


def set_numpr(ppr, num_id: int, level: int):
    current = ppr.find(qn("w:numPr"))
    if current is not None:
        ppr.remove(current)
    numpr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), str(level))
    numid = OxmlElement("w:numId")
    numid.set(qn("w:val"), str(num_id))
    numpr.append(ilvl)
    numpr.append(numid)
    ppr.insert(0, numpr)


def clear_style_numpr(style):
    ppr = style._element.get_or_add_pPr()
    current = ppr.find(qn("w:numPr"))
    if current is not None:
        ppr.remove(current)


def set_style_numpr(style, num_id: int, level: int):
    """Link a paragraph style to a Word list level.

    A style-level numPr is the important part of a true heading list.  Direct
    paragraph numPr values are deliberately removed from headings below so a
    pasted or previously broken list cannot override this linkage.
    """
    ppr = style._element.get_or_add_pPr()
    current = ppr.find(qn("w:numPr"))
    if current is not None:
        ppr.remove(current)
    numpr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), str(level))
    numid = OxmlElement("w:numId")
    numid.set(qn("w:val"), str(num_id))
    numpr.append(ilvl)
    numpr.append(numid)
    ppr.insert(0, numpr)


def set_outline_level(ppr, level: int | None):
    current = ppr.find(qn("w:outlineLvl"))
    if current is not None:
        ppr.remove(current)
    if level is not None:
        outline = OxmlElement("w:outlineLvl")
        outline.set(qn("w:val"), str(level))
        ppr.append(outline)


def set_heading_border(ppr):
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "6")
    bottom.set(qn("w:color"), "1F4E79")
    border.append(bottom)
    replace_child(ppr, "w:pBdr", border)


def next_numbering_id(numbering, tag: str, attr: str) -> int:
    values = []
    for element in numbering.findall(qn(tag)):
        value = element.get(qn(attr))
        if value is not None and value.isdigit():
            values.append(int(value))
    return max(values, default=0) + 1


def add_number_instance(numbering, abstract_id: int) -> int:
    num_id = next_numbering_id(numbering, "w:num", "w:numId")
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract = OxmlElement("w:abstractNumId")
    abstract.set(qn("w:val"), str(abstract_id))
    num.append(abstract)
    numbering.append(num)
    return num_id


def insert_abstract(numbering, abstract):
    first_num = numbering.find(qn("w:num"))
    if first_num is None:
        numbering.append(abstract)
    else:
        numbering.insert(numbering.index(first_num), abstract)


def add_level(abstract, level: int, text: str, style_id: str | None, size_hp: int):
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), str(level))
    child(lvl, "w:start", **{"w:val": "1"})
    child(lvl, "w:numFmt", **{"w:val": "decimal"})
    if style_id:
        child(lvl, "w:pStyle", **{"w:val": style_id})
    child(lvl, "w:suff", **{"w:val": "space"})
    child(lvl, "w:lvlText", **{"w:val": text})
    child(lvl, "w:lvlJc", **{"w:val": "left"})
    ppr = OxmlElement("w:pPr")
    tabs = child(ppr, "w:tabs")
    child(tabs, "w:tab", **{"w:val": "num", "w:pos": "510"})
    child(ppr, "w:ind", **{"w:left": "510", "w:hanging": "226"})
    lvl.append(ppr)
    rpr = OxmlElement("w:rPr")
    set_rfonts(rpr)
    child(rpr, "w:b")
    child(rpr, "w:bCs")
    child(rpr, "w:color", **{"w:val": "000000"})
    child(rpr, "w:sz", **{"w:val": str(size_hp)})
    child(rpr, "w:szCs", **{"w:val": str(size_hp)})
    lvl.append(rpr)
    abstract.append(lvl)


def add_heading_numbering(doc, style_ids: Sequence[str]) -> int:
    numbering = doc.part.numbering_part.element
    abstract_id = next_numbering_id(numbering, "w:abstractNum", "w:abstractNumId")
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    child(abstract, "w:multiLevelType", **{"w:val": "multilevel"})
    for level, text, style_id, size_hp in (
        (0, "第%1章", style_ids[0], 31),
        (1, "%1.%2", style_ids[1], 25),
        (2, "%1.%2.%3", style_ids[2], 23),
    ):
        add_level(abstract, level, text, style_id, size_hp)
    insert_abstract(numbering, abstract)
    return add_number_instance(numbering, abstract_id)


def add_step_numbering(doc) -> int:
    numbering = doc.part.numbering_part.element
    abstract_id = next_numbering_id(numbering, "w:abstractNum", "w:abstractNumId")
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    child(abstract, "w:multiLevelType", **{"w:val": "singleLevel"})
    add_level(abstract, 0, "%1.", "ManualStep", 21)
    insert_abstract(numbering, abstract)
    return abstract_id


def set_run_fonts(paragraph, reset_direct: bool = False):
    for run in paragraph.runs:
        run.font.name = "Times New Roman"
        # Keep existing Chinese size/emphasis, even for legacy callers.
        rpr = run._element.get_or_add_rPr()
        set_rfonts(rpr)


def iter_table_paragraphs(table):
    for row in table.rows:
        for cell in row.cells:
            yield from cell.paragraphs
            for nested in cell.tables:
                yield from iter_table_paragraphs(nested)


def iter_all_document_paragraphs(doc):
    """Yield body, table, header, and footer paragraphs for font normalization."""
    yield from doc.paragraphs
    for table in doc.tables:
        yield from iter_table_paragraphs(table)
    for section in doc.sections:
        for container in (
            section.header, section.first_page_header, section.even_page_header,
            section.footer, section.first_page_footer, section.even_page_footer,
        ):
            yield from container.paragraphs
            for table in container.tables:
                yield from iter_table_paragraphs(table)


def configure_styles(doc):
    def inherited(name, sid, source):
        existed = name in doc.styles
        style = ensure_style(doc, name, sid)
        if not existed:
            original = doc.styles[source]
            style.base_style = original.base_style or doc.styles['Normal']
            for tag in ('w:pPr', 'w:rPr'):
                node = original._element.find(qn(tag))
                if node is not None:
                    replace_child(style._element, tag, deepcopy(node))
        set_rfonts(style._element.get_or_add_rPr())
        return style
    title = inherited('手册题名', 'ManualTitle', 'Title')
    meta = inherited('手册元信息', 'ManualMeta', 'Normal')
    for style in (title, meta):
        clear_style_numpr(style)
        set_outline_level(style._element.get_or_add_pPr(), 9)
    styles = []
    for name, style_id, level, size, before, after in HEADING_TARGETS:
        style = inherited(name, style_id, f'Heading {level+1}')
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True
        ppr = style._element.get_or_add_pPr()
        set_outline_level(ppr, level)
        styles.append(style)
    step = inherited('操作步骤', 'ManualStep', 'Normal')
    clear_style_numpr(step)
    set_outline_level(step._element.get_or_add_pPr(), 9)
    return title, meta, styles, step


def heading_level(paragraph) -> int | None:
    name = paragraph.style.name
    if name in ("Heading 1", "一级标题"):
        return 0
    if name in ("Heading 2", "二级标题"):
        return 1
    if name in ("Heading 3", "三级标题"):
        return 2
    return None


def contiguous_groups(paragraphs: Iterable):
    group = []
    for paragraph in paragraphs:
        if group and paragraph._p.getprevious() is not group[-1]._p:
            yield group
            group = []
        group.append(paragraph)
    if group:
        yield group


def repair(input_path: Path, output_path: Path, paragraph_map=None):
    if input_path.resolve() == output_path.resolve():
        raise ValueError("Output must be a new revision; source overwrite is not allowed.")
    if output_path.exists():
        raise FileExistsError(f"Refusing to overwrite existing revision: {output_path}")

    original_bytes = input_path.read_bytes()
    doc = Document(input_path)
    mapping = paragraph_map or {}
    permitted_keys = {'title', 'metadata', 'headings', 'steps'}
    if set(mapping) - permitted_keys:
        raise ValueError('Unknown paragraph-map keys')
    roles = {}
    def role(index, value):
        if type(index) is not int or not 1 <= index <= len(doc.paragraphs) or index in roles:
            raise ValueError('Invalid or conflicting paragraph-map index')
        roles[index] = value
    if mapping.get('title') is not None: role(mapping['title'], 'title')
    for i in mapping.get('metadata', []): role(i, 'meta')
    for i in mapping.get('steps', []): role(i, 'step')
    for i, level in mapping.get('headings', {}).items():
        if type(level) is not int or level not in (1, 2, 3): raise ValueError('Heading level must be 1..3')
        role(int(i), level-1)
    for i, p in enumerate(doc.paragraphs, 1):
        if mapping or i in roles: continue
        level = heading_level(p)
        if level is not None: roles[i] = level
        elif p.style.name in ('Title', '手册题名'): roles[i] = 'title'
        elif p.style.name == '手册元信息': roles[i] = 'meta'
        elif p.style.name == '操作步骤': roles[i] = 'step'
    if not any(type(v) is int for v in roles.values()) and not (mapping and any(p.style.name in ("一级标题","二级标题","三级标题") for p in doc.paragraphs)):
        raise ValueError('No heading identified; provide an explicit paragraph map')
    original_nonparagraph = [node.xml for node in doc._element.body if node.tag != qn("w:p")]
    original_payload = [[__import__("lxml.etree",fromlist=["etree"]).tostring(n) for n in p._p if n.tag != qn("w:pPr")] for p in doc.paragraphs]
    original_text = [(p.text, tuple(p._p.xpath('.//w:instrText/text()'))) for p in doc.paragraphs]
    untouched = {i: p._p.xml for i,p in enumerate(doc.paragraphs,1) if i not in roles}
    existing = all(n in doc.styles for n in ('手册题名','手册元信息','一级标题','二级标题','三级标题','操作步骤'))
    if mapping and existing:
        # Local repair must not mutate shared style definitions used by other paragraphs.
        title, meta = doc.styles['手册题名'],doc.styles['手册元信息']
        headings = [doc.styles[n] for n in ('一级标题','二级标题','三级标题')]
        step_style = doc.styles['操作步骤']
        step_abstract_id = None
        if 'step' in roles.values(): step_abstract_id = add_step_numbering(doc)
    else:
        # Existing named styles are not implicitly redefined by an explicit local map.
        if mapping and any(n in doc.styles for n in ('一级标题','二级标题','三级标题','操作步骤')):
            raise ValueError('Partial existing style set; explicit whole-manual repair required')
        title, meta, headings, step_style = configure_styles(doc)
        heading_num_id = add_heading_numbering(doc, [style.style_id for style in headings])
        for level, style in enumerate(headings): set_style_numpr(style, heading_num_id, level)
        step_abstract_id = add_step_numbering(doc)

    for i, paragraph in enumerate(doc.paragraphs, 1):
        target = roles.get(i)
        if target in ('title', 'meta'):
            paragraph.style = title if target == 'title' else meta
            clear_numpr(paragraph)
        elif type(target) is int:
            paragraph.style = headings[target]
            clear_numpr(paragraph)
    direct_numbered = [p for i,p in enumerate(doc.paragraphs,1) if roles.get(i)=='step']
    for group in contiguous_groups(direct_numbered):
        step_num_id = add_number_instance(doc.part.numbering_part.element, step_abstract_id)
        instance = doc.part.numbering_part.element.xpath(f'./w:num[@w:numId="{step_num_id}"]')[0]
        override = child(instance, 'w:lvlOverride', **{'w:ilvl': '0'})
        child(override, 'w:startOverride', **{'w:val': '1'})
        for paragraph in group:
            paragraph.style = step_style
            set_numpr(paragraph._p.get_or_add_pPr(), step_num_id, 0)

    if original_text != [(p.text, tuple(p._p.xpath('.//w:instrText/text()'))) for p in doc.paragraphs]:
        raise ValueError('Text or field content changed')
    if any(doc.paragraphs[i-1]._p.xml != xml for i,xml in untouched.items()):
        raise ValueError('Unselected paragraph changed')

    if original_nonparagraph != [node.xml for node in doc._element.body if node.tag != qn("w:p")]: raise ValueError("Non-paragraph body content changed")
    if original_payload != [[__import__("lxml.etree",fromlist=["etree"]).tostring(n) for n in p._p if n.tag != qn("w:pPr")] for p in doc.paragraphs]: raise ValueError("Paragraph content payload changed")

    settings = doc.settings.element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    if not mapping: update.set(qn("w:val"), "true")

    fd, temporary = tempfile.mkstemp(suffix='.docx', dir=output_path.parent)
    os.close(fd)
    try:
        doc.save(temporary)
        check(Path(temporary))
        with zipfile.ZipFile(input_path) as a, zipfile.ZipFile(temporary) as b:
            allowed = {'word/document.xml','word/styles.xml','word/numbering.xml','word/settings.xml'}
            for part in set(a.namelist()) | set(b.namelist()):
                if part not in allowed and (part not in a.namelist() or part not in b.namelist() or a.read(part)!=b.read(part)):
                    raise ValueError('Non-target package part changed: '+part)
        if input_path.read_bytes() != original_bytes: raise RuntimeError('Source changed during repair')
        # Exclusive creation protects against a concurrently created revision.
        with output_path.open('xb') as stream: stream.write(Path(temporary).read_bytes())
    finally:
        Path(temporary).unlink(missing_ok=True)
    return len([p for p in doc.paragraphs if heading_level(p) is not None]), len(direct_numbered)


def numpr_values(paragraph):
    ppr = paragraph._p.pPr
    if ppr is None:
        return None
    numpr = ppr.find(qn("w:numPr"))
    if numpr is None:
        return None
    ilvl = numpr.find(qn("w:ilvl"))
    numid = numpr.find(qn("w:numId"))
    return (
        ilvl.get(qn("w:val")) if ilvl is not None else "0",
        numid.get(qn("w:val")) if numid is not None else None,
    )


def check(path: Path):
    doc = Document(path)
    expected = {"手册题名", "手册元信息", "一级标题", "二级标题", "三级标题", "操作步骤"}
    missing = [name for name in expected if name not in [style.name for style in doc.styles]]
    if missing:
        raise ValueError("Missing styles: " + ", ".join(missing))

    heading_styles = [doc.styles[name] for name, *_ in HEADING_TARGETS]
    style_nums = []
    for level, style in enumerate(heading_styles):
        values = numpr_values(type("StyleParagraph", (), {"_p": type("P", (), {"pPr": style._element.pPr})()})())
        if values is None or values[0] != str(level):
            raise ValueError(f"Heading style linkage is missing for {style.name}")
        style_nums.append(values[1])
    if len(set(style_nums)) != 1:
        raise ValueError("Heading styles are not linked to one shared multilevel list")
    if any(numpr_values(type("StyleParagraph", (), {"_p": type("P", (), {"pPr": doc.styles[name]._element.pPr})()})()) is not None for name in ("手册题名", "手册元信息", "操作步骤")):
        raise ValueError("Title/meta/step styles must not carry the heading list linkage")

    headings = 0
    steps = 0
    heading_num_id = style_nums[0]
    numbering = doc.part.numbering_part.element
    nums = numbering.xpath(f'./w:num[@w:numId="{heading_num_id}"]')
    if len(nums) != 1: raise ValueError('Heading numId does not resolve')
    abstract_id = nums[0].find(qn('w:abstractNumId')).get(qn('w:val'))
    abstracts = numbering.xpath(f'./w:abstractNum[@w:abstractNumId="{abstract_id}"]')
    if len(abstracts) != 1: raise ValueError('Heading abstractNumId does not resolve')
    for level, text in enumerate(('第%1章','%1.%2','%1.%2.%3')):
        levels = [n for n in abstracts[0].findall(qn("w:lvl")) if n.get(qn("w:ilvl")) == str(level)]
        if len(levels)!=1 or levels[0].find(qn('w:lvlText')).get(qn('w:val'))!=text:
            raise ValueError('Heading numbering pattern invalid')
        for tag,val in [('w:numFmt','decimal'),('w:start','1'),('w:pStyle',heading_styles[level].style_id)]:
            node=levels[0].find(qn(tag))
            if node is None or node.get(qn('w:val'))!=val: raise ValueError('Heading numbering definition invalid: '+tag)
    for paragraph in doc.paragraphs:
        if paragraph.style.name in ("一级标题", "二级标题", "三级标题"):
            headings += 1
            if numpr_values(paragraph) is not None:
                raise ValueError("Heading paragraph has a direct numPr override: " + paragraph.text)
        if paragraph.style.name == "操作步骤":
            steps += 1
            values = numpr_values(paragraph)
            if values is None or values[0] != "0":
                raise ValueError("Operation step lacks an independent list: " + paragraph.text)
            if values[1] == heading_num_id:
                raise ValueError("Operation step shares the heading list: " + paragraph.text)
            instances = numbering.xpath(f'./w:num[@w:numId="{values[1]}"]')
            if len(instances)!=1 or not any(n.get(qn("w:ilvl")) == "0" and n.find(qn("w:startOverride")) is not None and n.find(qn("w:startOverride")).get(qn("w:val")) == "1" for n in instances[0].findall(qn("w:lvlOverride"))):
                raise ValueError('Step list must explicitly restart at 1')
            abstract=instances[0].find(qn('w:abstractNumId'))
            aid=abstract.get(qn('w:val')) if abstract is not None else ''
            candidates=numbering.xpath(f'./w:abstractNum[@w:abstractNumId="{aid}"]')
            if len(candidates)!=1: raise ValueError('Step abstract list missing')
            levels=[n for n in candidates[0].findall(qn('w:lvl')) if n.get(qn('w:ilvl'))=='0']
            if len(levels)!=1: raise ValueError('Step level missing')
            for tag,value in [('w:numFmt','decimal'),('w:lvlText','%1.')]:
                n=levels[0].find(qn(tag))
                if n is None or n.get(qn('w:val'))!=value: raise ValueError('Step numbering definition invalid')
    if not headings: raise ValueError('No heading paragraphs; cannot validate heading repair')
    print(f"MANUAL_HEADING_NUMBERING_CHECK_PASS headings={headings} steps={steps} list_num={style_nums[0]}")


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("input", type=Path, help="source DOCX for a new revision")
    parser.add_argument("output", type=Path, nargs="?", help="new output DOCX")
    parser.add_argument("--check", action="store_true", help="validate an already repaired DOCX")
    parser.add_argument('--paragraph-map', type=Path, help='Explicit 1-based body paragraph roles JSON')
    args = parser.parse_args()
    if args.check:
        if args.output is not None:
            parser.error("--check accepts exactly one DOCX path")
        check(args.input)
        return
    if args.output is None:
        parser.error("output DOCX is required unless --check is used")
    mapping = json.loads(args.paragraph_map.read_text(encoding='utf-8-sig')) if args.paragraph_map else None
    headings, steps = repair(args.input, args.output, mapping)
    print(f"MANUAL_HEADING_NUMBERING_REPAIR_PASS headings={headings} steps={steps} output={args.output}")


if __name__ == "__main__":
    main()
